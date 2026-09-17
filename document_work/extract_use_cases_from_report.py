from pathlib import Path
import json
import re

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path(r"D:\PFE - BRI Technology\Rapport\Master Report.docx")
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\use_case_report_inventory.json")
IMAGE_DIR = Path(r"D:\PFE M2\Platforme SaaS\document_work\source_use_case_images")

doc = Document(str(SOURCE))

paragraphs = []
for i, p in enumerate(doc.paragraphs):
    text = " ".join(p.text.split())
    if text:
        paragraphs.append({"index": i, "style": p.style.name if p.style else "", "text": text})

tables = []
for idx, table in enumerate(doc.tables, 1):
    rows = []
    for row in table.rows:
        vals = [" ".join(cell.text.split()) for cell in row.cells]
        # python-docx can expose merged cells more than once; retain order but collapse exact repeats.
        compact = []
        for value in vals:
            if not compact or value != compact[-1]:
                compact.append(value)
        rows.append(compact)
    flat = " | ".join(v for row in rows for v in row if v)
    tables.append({"index": idx, "rows": rows, "flat": flat})

payload = {
    "source": str(SOURCE),
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "paragraphs": paragraphs,
    "tables": tables,
}
OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

IMAGE_DIR.mkdir(parents=True, exist_ok=True)
extracted = []
for i, p in enumerate(doc.paragraphs):
    blips = p._p.xpath(".//a:blip")
    for j, blip in enumerate(blips, 1):
        rid = blip.get(qn("r:embed"))
        if not rid or rid not in doc.part.rels:
            continue
        part = doc.part.rels[rid].target_part
        ext = Path(str(part.partname)).suffix or ".bin"
        target = IMAGE_DIR / f"paragraph_{i:04d}_{j}{ext}"
        target.write_bytes(part.blob)
        extracted.append({"paragraph": i, "path": str(target), "caption_context": paragraphs[i]["text"] if i < len(paragraphs) else ""})
payload["extracted_images"] = extracted
OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)}")
print(f"inventory={OUTPUT}")
print(f"extracted_images={len(extracted)} directory={IMAGE_DIR}")
print("\nHEADINGS / USE-CASE LABELS")
for item in paragraphs:
    t = item["text"]
    if (item["style"].startswith("Heading") or
            re.search(r"UC-?\d+", t, re.I) or
            "cas d’utilisation" in t.lower() or
            "cas d'utilisation" in t.lower()):
        print(f'{item["index"]:04d}\t{item["style"]}\t{t}')

print("\nTABLE SUMMARIES")
for table in tables:
    if re.search(r"UC-?\d+|Acteur principal|Scénario nominal", table["flat"], re.I):
        print(f'TABLE {table["index"]}: {table["flat"][:700]}')
