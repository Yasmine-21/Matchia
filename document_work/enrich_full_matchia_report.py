from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement

ROOT = Path(r"D:\PFE M2\Platforme SaaS")
SOURCE = ROOT / "document_work" / "Rapport_PFE_Complet_Matchia.docx"
OUT = ROOT / "document_work" / "Rapport_PFE_Complet_Matchia_Approfondi.docx"

def insert_after(anchor, text, style="Normal"):
    p = OxmlElement("w:p")
    anchor._p.addnext(p)
    new = Paragraph(p, anchor._parent)
    new.style = style
    new.add_run(text)
    return new

def find_heading(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise ValueError(prefix)

def add_block(doc, heading, paragraphs):
    current = find_heading(doc, heading)
    for paragraph in paragraphs:
        current = insert_after(current, paragraph)

def main():
    doc = Document(SOURCE)
    add_block(doc, "1.1", [
        "Dans l’organisation antérieure, une banque qui souhaitait disposer d’une marketplace devait solliciter une équipe de développement afin d’obtenir une version spécifique de l’application. Chaque nouvelle demande impliquait des adaptations manuelles du code, des tests propres à la banque et un déploiement indépendant. Cette organisation crée une dépendance forte à l’équipe technique et ne permet pas d’industrialiser simplement l’ouverture d’une nouvelle marketplace.",
        "Le parcours commercial était également fragmenté. La création d’une marketplace reposait sur des échanges traditionnels entre la banque et le fournisseur de la solution. Le paiement de la prestation n’était pas intégré à un circuit en ligne vérifiable, ce qui séparait la validation administrative, la facturation et l’activation de la plateforme. La mise en production n’était donc pas reliée de manière contrôlée à la confirmation d’un règlement.",
        "Le cas des concessionnaires met particulièrement en évidence cette fragmentation. Un concessionnaire partenaire de plusieurs banques devait gérer plusieurs comptes et répliquer les informations de son catalogue dans plusieurs espaces. Le stock devait être mis à jour manuellement dans chaque marketplace, ce qui favorisait les incohérences et ne permettait pas de disposer d’une vision centrale des produits, des partenariats et des demandes associées."
    ])
    add_block(doc, "1.2", [
        "La solution proposée par Matchia ne consiste donc pas à produire une nouvelle application autonome à chaque banque. Elle fournit un socle partagé, administré par un opérateur SaaS, dans lequel chaque banque est représentée par un tenant logique. La configuration de la marketplace — identité visuelle, stores, modules, contenus, utilisateurs, produits et souscription — devient une donnée pilotée par le back-office plutôt qu’une modification du code source.",
        "Cette réponse concilie mutualisation et autonomie. La plateforme partage les mécanismes techniques de sécurité, de paiement, de notification, de traçabilité et de déploiement. En revanche, l’administrateur Banque agit dans un périmètre restreint à son organisation. Il conserve le contrôle de sa marque, de son offre et de ses décisions opérationnelles, en particulier pour les partenariats, les publications et l’instruction des demandes de financement."
    ])
    add_block(doc, "1.4", [
        "Scrum est adapté à ce sujet parce que les fonctionnalités présentent des dépendances fortes mais peuvent être livrées progressivement. Le socle d’identité et de persistance doit précéder la multi-tenance ; l’onboarding et la marketplace précèdent les parcours publics ; les partenariats et la publication précèdent le financement d’un produit concessionnaire. La planification ne se réduit pas à une succession de tâches techniques : chaque sprint doit viser un incrément utilisable et démontrable.",
        "La qualité fait partie de la définition de terminé. Pour chaque user story, les critères d’acceptation décrivent les validations métier, les autorisations, les transitions de statut et les notifications attendues. Les tests unitaires et d’intégration, la revue de code, l’analyse SonarQube et la démonstration de l’écran concerné constituent des éléments de validation. Cette pratique limite le risque de réaliser une interface visuellement correcte mais incohérente avec les règles de gestion du backend."
    ])
    add_block(doc, "2.1", [
        "L’identification des acteurs repose sur les rôles applicatifs effectivement définis dans la sécurité et le routage : ADMIN_SAAS, ADMIN_BANK, DEALER_ADMIN et CLIENT. Les rôles ne constituent pas de simples libellés d’interface. Ils conditionnent les routes accessibles, les contrôles du backend et les relations de données consultables. Les acteurs externes sont distingués des utilisateurs humains car ils sont appelés par l’API et ne disposent pas d’un compte leur permettant de naviguer dans les espaces métiers.",
        "La séparation des responsabilités est un point essentiel du système. L’administrateur SaaS gouverne l’offre et les tenants sans traiter les opérations métier quotidiennes d’une banque. L’administrateur Banque possède l’autorité de configuration et de décision dans son tenant. Le concessionnaire fournit des produits et sollicite une visibilité sans contourner l’approbation bancaire. Le client dépose son dossier sans pouvoir influer sur l’instruction. Cette distribution formalise une séparation des droits et des responsabilités."
    ])
    add_block(doc, "2.2", [
        "Les besoins fonctionnels sont formulés comme des capacités observables. Une capacité ne peut être considérée comme réalisée que si l’interface, l’API, les validations, la persistance et les contrôles d’accès permettent d’achever le parcours. Par exemple, la création d’un dossier de financement ne se limite pas à un formulaire : elle comprend l’état brouillon, le contrôle des pièces, la soumission, l’instruction réservée à la banque, la décision, la notification du client et l’impact éventuel sur le stock du concessionnaire.",
        "Les sous-sections suivantes organisent les exigences par sous-systèmes afin d’éviter les recouvrements. Les fonctions de configuration du SaaS ne sont pas confondues avec les fonctions d’exploitation d’un tenant, et les fonctions d’un partenaire ne sont pas confondues avec celles d’un client. Cette décomposition sera reprise dans l’architecture logique et dans les diagrammes de cas d’utilisation du chapitre 3."
    ])
    add_block(doc, "2.3", [
        "Les exigences non fonctionnelles sont déterminantes dans le contexte bancaire. La confidentialité et l’isolation ne sont pas des options d’interface : elles doivent être appliquées par les services, les repositories et les contrôles de rôle. La traçabilité doit conserver les événements pertinents sans exposer de secrets. Enfin, le déploiement doit être reproductible, car une correction commune au socle SaaS doit pouvoir être testée et livrée de manière cohérente pour tous les tenants.",
        "Les choix d’implémentation reflètent cette orientation : DTO et validation côté API, authentification stateless, services métier transactionnels, tests automatisés, couverture JaCoCo, scans SonarQube et images Docker. Ces mécanismes ne garantissent pas à eux seuls une conformité réglementaire complète ; ils constituent le socle technique nécessaire à une évolution vers une exploitation plus exigeante."
    ])
    add_block(doc, "2.4", [
        "La multi-tenance retenue est de type base partagée avec isolation logique par banque. Les données ne sont donc pas dupliquées dans une instance applicative par organisation. Le contexte bancaire est porté par les relations métier et vérifié au moment des opérations. Les ressources configurables — stores, modules, contenus et branding — sont associées à la marketplace de la banque, tandis que les ressources sensibles — utilisateurs, demandes et décisions — sont filtrées par les services d’autorisation.",
        "Ce modèle convient particulièrement aux besoins d’ouverture progressive de nouvelles banques. Une nouvelle marketplace est obtenue par un workflow d’adhésion, de sélection et d’activation, plutôt que par un nouveau cycle complet de développement. La personnalisation reste déclarative et le pipeline CI/CD demeure commun, ce qui facilite la maintenance des correctifs et des améliorations de sécurité."
    ])
    add_block(doc, "2.5", [
        "Les règles de gestion constituent la traduction opérationnelle des besoins. Elles sont portées par les statuts d’entités, par les validations des services et par les relations de possession. Elles évitent notamment qu’un paiement non confirmé active une marketplace, qu’un produit soit publié sans relation contractuelle valide ou qu’un utilisateur traite un dossier appartenant à une autre banque. Les diagrammes d’activité des chapitres 4 et 5 illustrent ces règles sous forme de décisions et de branches."
    ])
    add_block(doc, "2.6", [
        "Le backlog est priorisé selon la valeur métier et les dépendances. Les fonctionnalités d’identité, de tenant et de sécurité sont classées en priorité haute, car elles conditionnent les autres parcours. Les fonctions de catalogue, de partenaires et de financement sont ensuite réalisées par incréments. Les capacités analytiques IA sont volontairement placées après la stabilisation des données, car elles doivent s’appuyer sur un modèle métier fiable et sur un contrôle strict de la lecture de la base."
    ])
    add_block(doc, "3.1", [
        "L’environnement retenu privilégie une séparation claire entre développement, qualité et exploitation. Le frontend est un projet Node/Vite ; le backend est un projet Maven Spring Boot. La base PostgreSQL porte les données métiers, tandis que les fichiers téléversés sont gérés séparément par les répertoires configurés. Les outils d’intégration et d’analyse sont conteneurisés, ce qui rend leur exécution reproductible sur une machine de développement ou un serveur d’intégration.",
        "L’usage de GitHub, Jenkins, Docker et Azure s’inscrit dans une chaîne de livraison continue. Le code versionné déclenche un pipeline ; le pipeline construit, teste et analyse ; les images résultantes sont publiées ; le déploiement met à jour les applications conteneurisées. Cette continuité réduit les différences entre le poste de développement et l’environnement de livraison."
    ])
    add_block(doc, "3.2", [
        "Le choix d’une SPA React répond à la richesse des interfaces de Matchia : tableaux de bord, formulaires par étapes, paramétrage de modules, fenêtres de détail, filtres et parcours protégés. TypeScript améliore la cohérence des contrats manipulés par les composants. Le client HTTP centralise la communication avec le backend et facilite l’injection de l’authentification ainsi que la gestion homogène des erreurs.",
        "Spring Boot a été retenu pour structurer une API REST mature autour de la sécurité, de la validation, de la persistance relationnelle et de l’envoi d’e-mails. La séparation controller–service–repository permet d’isoler la réception HTTP, les règles métier et l’accès aux données. PostgreSQL est cohérent avec les relations nombreuses entre banques, produits, partenariats, contrats, abonnements et demandes de financement."
    ])
    add_block(doc, "3.3", [
        "L’architecture physique distingue les nœuds qui exécutent la solution. Le navigateur héberge la SPA ; le serveur frontend sert les ressources statiques ; le backend expose les services REST et appelle les dépendances externes ; PostgreSQL persiste les entités. Cette représentation évite de confondre les éléments métiers avec les emplacements d’exécution. Elle permet aussi de visualiser les frontières de confiance : le navigateur ne contacte pas directement Stripe, SMTP, Gemini ou la base de données pour les opérations métier.",
        "Les flux principaux reposent sur HTTPS en production : navigateur vers frontend, frontend vers API, puis backend vers base et services externes. Les fichiers téléversés sont gérés par le backend sous contrôle d’autorisation. Les données de paiement et les secrets des services externes doivent être protégés dans l’environnement d’exécution et ne sont pas exposés au frontend."
    ])
    add_block(doc, "3.4", [
        "L’architecture logique applique le principe de responsabilité unique. Le contrôleur convertit l’échange HTTP en appel applicatif ; le service vérifie le rôle, le tenant, les règles et les transitions d’état ; le repository réalise les requêtes de persistance ; les DTO limitent les données renvoyées au client. Cette organisation rend les règles testables sans nécessiter l’interface React et évite que des validations importantes soient placées uniquement dans le navigateur.",
        "Les fonctions transverses traversent ces couches sans les remplacer. JwtAuthenticationFilter installe le contexte d’authentification ; les services de notification et d’e-mail diffusent les événements ; l’audit enregistre les actions ; le paiement confirme la transaction avant activation ; l’assistant IA applique une validation déterministe avant une lecture SQL. Les composants restent donc spécialisés mais coordonnés par les workflows métier."
    ])
    add_block(doc, "3.5", [
        "Le contexte tenant ne doit pas être déduit uniquement du paramètre présent dans une URL. L’application valide la cohérence entre le rôle connecté, la banque associée et la ressource demandée. Cette approche réduit le risque qu’un changement de paramètre côté client permette d’accéder à une ressource étrangère. Dans le cas d’un administrateur SaaS, l’accès transverse est explicite et justifié par son rôle de gouvernance ; dans le cas d’un administrateur Banque, les services restreignent la sélection au tenant courant.",
        "L’isolation est également visible dans la personnalisation publique. Le slug de la banque permet de charger la marketplace et son branding ; les stores, modules et contenus visibles sont ceux effectivement activés. La même SPA conserve ainsi une expérience cohérente tout en présentant des données et une identité différentes selon la marketplace consultée."
    ])
    add_block(doc, "3.6", [
        "Les diagrammes UML sont utilisés ici comme un langage de conception et non comme une décoration. Un diagramme de cas d’utilisation délimite le système, ses acteurs et leurs objectifs. Un scénario précise les conditions, les étapes, les alternatives et le résultat attendu. Le diagramme de classes montre les concepts persistants et les associations structurantes. Les attributs techniques secondaires et les classes de framework sont volontairement exclus afin de conserver une lecture métier.",
        "Les notations respectent les principes de base : acteurs externes aux frontières du système, cas d’utilisation sous forme d’ellipses, associations entre acteurs et objectifs, classes avec attributs synthétiques et associations munies de cardinalités lorsque cela est pertinent. Les activités et séquences, présentées ensuite dans les chapitres de réalisation, complètent ces vues en décrivant respectivement les décisions et les échanges temporels."
    ])
    add_block(doc, "3.6.1", [
        "La vue globale évite de dessiner tous les écrans du système. Elle regroupe les objectifs par domaines : gouvernance SaaS, configuration de tenant, partenaires et publications, financement, paiement et traçabilité. Les systèmes externes n’apparaissent que lorsqu’ils participent à un objectif réel : Stripe pour le paiement, SMTP pour les communications et Gemini pour la génération contrôlée d’une réponse analytique."
    ])
    add_block(doc, "3.6.2", [
        "Les scénarios détaillés complètent le diagramme par une description textuelle normalisée. Ils précisent les préconditions, car un diagramme seul ne suffit pas à exposer les contrôles d’état ou les droits. Ils isolent aussi les alternatives importantes : rejet de demande, paiement annulé, pièces manquantes, contrat expiré ou SQL non conforme. Cette formalisation garantit que l’implémentation et la recette s’appuient sur les mêmes règles métier."
    ])
    add_block(doc, "3.6.3", [
        "Bank est l’agrégat central de la multi-tenance. Marketplace décrit son exposition publique et son identité ; MarketplaceStore et MarketplaceStoreModule permettent d’activer et de configurer les éléments du catalogue global dans un tenant. Subscription et Payment représentent la dimension commerciale : une souscription porte l’état du service et le paiement participe à sa confirmation.",
        "Dealer, DealerBankPartnership et PartnershipContract représentent le cycle B2B. Ils sont distincts de DealerProduct afin qu’un même produit puisse être administré une seule fois tout en étant proposé dans plusieurs contextes bancaires. FinancingRequest relie le client, la banque, le store et le produit ; FinancingRequestDocument porte les pièces du dossier. Notification et AuditLog demeurent transverses, car ils documentent des événements produits par plusieurs domaines."
    ])
    add_block(doc, "4.7", [
        "Les diagrammes insérés dans ce chapitre sont placés là où les statuts et les interactions traversent plusieurs composants. Les activités explicitent notamment les branches d’approbation, de paiement et d’activation ; les séquences rendent visibles les responsabilités du formulaire React, de l’API, du service de paiement, de la base et de la messagerie. Les captures viennent ensuite ancrer ces représentations dans les interfaces réellement disponibles, sans les remplacer."
    ])
    add_block(doc, "5.7", [
        "Les diagrammes d’activité et de séquence du chapitre 5 soutiennent les processus les plus sensibles : cycle de publication, financement et assistant IA. Ils montrent que la décision bancaire reste séparée des actions du concessionnaire et du client, que la réservation de stock est conditionnée à une acceptation et que Gemini ne dispose d’aucun accès direct à PostgreSQL. Les captures associées documentent les écrans de suivi et les points de contrôle sans prétendre reproduire une fonctionnalité absente du prototype."
    ])
    add_block(doc, "6.1", [
        "La qualité est évaluée à plusieurs niveaux. Les tests unitaires vérifient des fonctions et composants isolés ; les tests de service vérifient les règles métier et les cas d’erreur ; les tests d’API permettent de contrôler l’intégration avec la sécurité et la persistance. JaCoCo produit un rapport de couverture pour le backend, tandis que Vitest s’appuie sur un environnement adapté aux composants frontend. SonarQube transforme ces résultats et l’analyse statique en indicateurs consultables par l’équipe.",
        "Le pipeline présent réalise les tests backend et les analyses Sonar. Une évolution recommandée consiste à enrichir cette chaîne par des tests frontend systématiques, des tests de bout en bout des parcours critiques et un Quality Gate bloquant. Cette distinction entre capacités actuellement codées et améliorations attendues est nécessaire pour maintenir la rigueur du mémoire."
    ])
    add_block(doc, "6.2", [
        "La conteneurisation sépare les responsabilités d’exécution. Le conteneur frontend sert le build statique via Nginx ; le conteneur backend exécute l’API Spring Boot ; la persistance des uploads est montée comme volume. Docker Compose fournit une topologie locale cohérente pour tester la communication des composants. Les configurations Jenkins et SonarQube disposent de fichiers Compose distincts afin de ne pas mélanger les services de développement avec l’infrastructure de qualité.",
        "Les images constituent un artefact immuable lorsqu’elles sont taguées par le numéro de build. Le tag latest facilite l’usage courant, mais le tag versionné est essentiel pour assurer la traçabilité, la reproductibilité et un éventuel retour arrière. En production, une politique de conservation et de signature des images renforcerait cette chaîne."
    ])
    add_block(doc, "6.3", [
        "Le déclencheur githubPush relie une modification de code à l’exécution du pipeline. Jenkins commence par récupérer exactement la révision concernée, puis réalise les étapes dans un ordre déterministe. Le build frontend reçoit l’URL de l’API cible au moment de la compilation ; les images backend et frontend sont construites séparément ; les credentials Docker Hub et Sonar évitent de placer les secrets dans le Jenkinsfile.",
        "La phase de déploiement utilise Azure CLI pour mettre à jour les applications conteneurisées. Le mémoire doit cependant rester prudent : un script de déploiement ne remplace pas les pratiques de gouvernance opérationnelle. Les validations de conformité, les sauvegardes, la supervision, le plan de reprise et la revue des droits d’accès doivent être définis avant une exploitation bancaire réelle."
    ])
    add_block(doc, "6.4", [
        "L’architecture de déploiement distingue le pipeline de production et l’application livrée. Jenkins, le moteur Docker et SonarQube peuvent être exécutés comme services d’infrastructure. Docker Hub sert de registre intermédiaire. Azure Container Apps exécute les images applicatives. PostgreSQL et les fichiers persistants doivent être protégés indépendamment de la durée de vie des conteneurs, car une reconstruction d’image ne doit pas entraîner de perte de données.",
        "La séparation frontend/backend facilite les mises à jour : une évolution d’interface peut être livrée sans reconstruire la logique métier si le contrat d’API demeure compatible. Réciproquement, une évolution de backend doit préserver ou versionner l’API attendue par la SPA. Cette discipline est particulièrement importante dans un modèle SaaS où une mise à jour est partagée par l’ensemble des tenants."
    ])
    doc.core_properties.title = "Rapport PFE complet et approfondi — Matchia"
    doc.save(OUT)
    print(OUT)

if __name__ == "__main__":
    main()
