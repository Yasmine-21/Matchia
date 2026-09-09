from pathlib import Path
import re
from docx import Document

SOURCE = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - diagrammes activités modèle harmonisé.docx")

doc = Document(SOURCE)

patterns = re.compile(
    r"(?:^Chapitre\s+[45]|^\d+(?:\.\d+)*\.?\s|Authent|marketplace|demande|concessionnaire|produit|client|financement|assistant|chatbot|Figure\s+[45]\.)",
    re.I,
)

for i, p in enumerate(doc.paragraphs):
    text = " ".join(p.text.split())
    if 400 <= i <= 615 and text:
        escaped = text.encode("unicode_escape").decode("ascii")
        print(f"{i:04d}\t{p.style.name}\t{escaped}")

print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} SHAPES={len(doc.inline_shapes)}")
