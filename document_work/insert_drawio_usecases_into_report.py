from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


SOURCE = Path(r"D:\PFE - BRI Technology\Rapport\Master Report.docx")
PNG_DIR = Path(r"D:\PFE M2\Platforme SaaS\deliverables\use_cases_drawio\aperçus_png")
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\deliverables\Master Report - diagrammes use case Draw.io intégrés.docx")


DIAGRAMS = [
    {
        "heading": "6.1. Cas d’utilisation détaillée : S’authentifier",
        "caption_prefix": "Figure 3.5",
        "image": "UC-01_Authentification.png",
        "alt": "Diagramme de cas d’utilisation détaillé UC-01 — Authentification",
    },
    {
        "heading": "6.2. Cas d’utilisation détaillée : Déposer une demande de marketplace",
        "caption_prefix": "Figure 3.6",
        "image": "UC-02_Demande_marketplace.png",
        "alt": "Diagramme de cas d’utilisation détaillé UC-02 — Demande de marketplace",
    },
    {
        "heading": "6.3. Cas d’utilisation détaillée : Traiter une demande et activer les services",
        "caption_prefix": "Figure 3.8",
        "image": "UC-03_Traitement_paiement_activation.png",
        "alt": "Diagramme de cas d’utilisation détaillé UC-03 — Traitement, paiement Stripe et activation",
    },
    {
        "heading": "6.4. Cas d’utilisation détaillée : Inscrire et valider un concessionnaire",
        "caption_prefix": "Figure 3.10",
        "image": "UC-04_Inscription_concessionnaire.png",
        "alt": "Diagramme de cas d’utilisation détaillé UC-04 — Inscription et validation d’un concessionnaire",
    },
    {
        "heading": "6.5. Cas d’utilisation détaillée : Gérer et publier un produit concessionnaire",
        "caption_prefix": "Figure 3.12",
        "image": "UC-05_Gestion_publication_produit.png",
        "alt": "Diagramme de cas d’utilisation détaillé UC-05 — Gestion et publication d’un produit",
    },
    {
        "heading": "6.6. Cas d’utilisation détaillée : Déposer et traiter une demande de financement",
        "caption_prefix": "Figure 3.13",
        "image": "UC-06_Demande_financement.png",
        "alt": "Diagramme de cas d’utilisation détaillé UC-06 — Demande de financement",
    },
]


def normalized(text: str) -> str:
    return " ".join(text.replace("\xa0", " ").split())


def remove_drawings(paragraph) -> int:
    removed = 0
    for xpath in (".//w:drawing", ".//w:pict"):
        for node in list(paragraph._p.xpath(xpath)):
            parent = node.getparent()
            parent.remove(node)
            removed += 1
    # Remove now-empty runs left behind by the former image.
    for run in list(paragraph._p.xpath("./w:r")):
        if not run.xpath(".//w:t") and not run.xpath(".//w:tab") and not run.xpath(".//w:br"):
            paragraph._p.remove(run)
    return removed


def insert_paragraph_before(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    return paragraph._parent.add_paragraph()._parent.paragraphs[-1] if False else paragraph.__class__(new_p, paragraph._parent)


def add_diagram(image_paragraph, image_path: Path, alt_text: str) -> None:
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.paragraph_format.space_before = Inches(0.05)
    image_paragraph.paragraph_format.space_after = Inches(0.04)
    run = image_paragraph.add_run()
    shape = run.add_picture(str(image_path), width=Inches(6.25))
    doc_pr = shape._inline.docPr
    doc_pr.set("title", alt_text)
    doc_pr.set("descr", alt_text)


def find_caption(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if normalized(p.text).startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one caption for {prefix!r}, found {len(matches)}")
    return matches[0]


def preceding_paragraph(paragraph):
    node = paragraph._p.getprevious()
    while node is not None and node.tag != qn("w:p"):
        node = node.getprevious()
    return paragraph.__class__(node, paragraph._parent) if node is not None else None


def style_caption(caption) -> None:
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = False
    caption.paragraph_format.space_before = Inches(0.02)
    caption.paragraph_format.space_after = Inches(0.08)
    for run in caption.runs:
        run.italic = True


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    for item in DIAGRAMS:
        image_path = PNG_DIR / item["image"]
        if not image_path.exists():
            raise FileNotFoundError(image_path)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document(SOURCE)
    initial_tables = len(document.tables)
    initial_shapes = len(document.inline_shapes)

    for item in DIAGRAMS:
        caption = find_caption(document, item["caption_prefix"])
        previous = preceding_paragraph(caption)

        # In the source, five figures already exist directly before their caption.
        # UC-02 is unusual: its old picture and caption share the same paragraph.
        removed_from_caption = remove_drawings(caption)
        if removed_from_caption:
            image_paragraph = insert_paragraph_before(caption)
        elif previous is not None and previous._p.xpath(".//w:drawing | .//w:pict"):
            remove_drawings(previous)
            image_paragraph = previous
        else:
            # UC-03 has a caption but no figure in the original report.
            image_paragraph = insert_paragraph_before(caption)

        add_diagram(image_paragraph, PNG_DIR / item["image"], item["alt"])
        style_caption(caption)

    # Keep each detailed use case visually coherent: from UC-02 onward, the
    # heading begins on a new page instead of being orphaned below the previous
    # scenario table while its diagram moves to the following page.
    for item in DIAGRAMS[1:]:
        heading_target = normalized(item["heading"])
        heading_matches = [p for p in document.paragraphs if normalized(p.text) == heading_target]
        if len(heading_matches) != 1:
            raise RuntimeError(f"Expected one heading {item['heading']!r}, found {len(heading_matches)}")
        heading_matches[0].paragraph_format.page_break_before = True

    # Correct the one scenario-reference typo while keeping every scenario table.
    for paragraph in document.paragraphs:
        if "UC-06" in paragraph.text and "Inscrire et valider un concessionnaire" in paragraph.text:
            for run in paragraph.runs:
                if "UC-06" in run.text:
                    run.text = run.text.replace("UC-06", "UC-04")

    document.save(OUTPUT)

    check = Document(OUTPUT)
    if len(check.tables) != initial_tables:
        raise RuntimeError(f"Table count changed: {initial_tables} -> {len(check.tables)}")
    if len(check.inline_shapes) != initial_shapes + 1:
        raise RuntimeError(f"Unexpected shape count: {initial_shapes} -> {len(check.inline_shapes)}")
    text = "\n".join(p.text for p in check.paragraphs)
    for item in DIAGRAMS:
        if item["caption_prefix"] not in text:
            raise RuntimeError(f"Missing caption {item['caption_prefix']}")

    print(f"Created: {OUTPUT}")
    print(f"Tables preserved: {len(check.tables)}")
    print(f"Inline shapes: {initial_shapes} -> {len(check.inline_shapes)}")


if __name__ == "__main__":
    main()
