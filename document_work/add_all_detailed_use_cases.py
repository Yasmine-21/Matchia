from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw
import sys

ROOT = Path(r"D:\PFE M2\Platforme SaaS")
sys.path.insert(0, str(ROOT / "document_work"))
import build_matchia_chapters_4_5 as c45

SOURCE = ROOT / "document_work" / "Rapport_PFE_Complet_Matchia_Approfondi.docx"
OUT = ROOT / "document_work" / "Rapport_PFE_Complet_Matchia_UseCases_Complets.docx"
DIAG = ROOT / "document_work" / "detailed_use_cases_uml"
DIAG.mkdir(parents=True, exist_ok=True)

def text_of(element):
    return ''.join(element.itertext())

def find_heading(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): return p
    raise ValueError(prefix)

def insert_after(anchor, text='', style='Normal'):
    xml = OxmlElement('w:p'); anchor._p.addnext(xml)
    p = Paragraph(xml, anchor._parent); p.style = style
    if text: p.add_run(text)
    return p

def remove_between(start, end):
    current = start._p.getnext()
    while current is not None and current is not end._p:
        nxt = current.getnext(); current.getparent().remove(current); current = nxt

def font(size, bold=False): return c45.font(size, bold)

def wrap(draw, value, f, width):
    rows=[]; row=''
    for word in value.split():
        trial=(row+' '+word).strip()
        if draw.textbbox((0,0),trial,font=f)[2] <= width: row=trial
        else: rows.append(row); row=word
    if row: rows.append(row)
    return rows

def actor(draw, x, y, name):
    col='#1F4D78'; draw.ellipse((x-18,y,x+18,y+36),outline=col,width=4)
    draw.line((x,y+36,x,y+97),fill=col,width=4); draw.line((x-38,y+58,x+38,y+58),fill=col,width=4)
    draw.line((x,y+97,x-34,y+143),fill=col,width=4); draw.line((x,y+97,x+34,y+143),fill=col,width=4)
    f=font(19,True); rows=wrap(draw,name,f,160); yy=y+154
    for r in rows:
        w=draw.textbbox((0,0),r,font=f)[2]; draw.text((x-w//2,yy),r,font=f,fill=col); yy+=25

def ellipse(draw, x, y, w, h, label):
    draw.ellipse((x,y,x+w,y+h),fill='#F7FBFF',outline='#2E74B5',width=4)
    f=font(20,True); rows=wrap(draw,label,f,w-45); heights=[draw.textbbox((0,0),r,font=f)[3]-draw.textbbox((0,0),r,font=f)[1] for r in rows]
    yy=y+(h-sum(heights)-5*(len(rows)-1))//2
    for r,hh in zip(rows,heights):
        tw=draw.textbbox((0,0),r,font=f)[2]; draw.text((x+(w-tw)//2,yy),r,font=f,fill='#0B2545'); yy+=hh+5

def usecase_diagram(filename, title, left, right, usecases):
    img=Image.new('RGB',(1800,1120),'white'); d=ImageDraw.Draw(img)
    d.rectangle((0,0,1800,95),fill='#0B2545'); d.text((48,26),title,font=font(35,True),fill='white')
    d.rectangle((360,145,1440,1045),outline='#0B2545',width=4); d.text((390,165),'Système Matchia',font=font(24,True),fill='#0B2545')
    ly=300
    for name in left:
        actor(d,145,ly,name); ly+=300
    ry=300
    for name in right:
        actor(d,1655,ry,name); ry+=300
    positions=[]
    for i,label in enumerate(usecases):
        col=i%2; row=i//2; x=510+col*490; y=270+row*230; positions.append((x,y,370,135,label)); ellipse(d,x,y,370,135,label)
    for idx,(x,y,w,h,label) in enumerate(positions):
        actor_y=370 if idx%2==0 else 670
        d.line((185,actor_y,x,y+h//2),fill='#5B6573',width=3)
        if right and idx in (1,3,5): d.line((x+w,y+h//2,1615,370 if idx%2 else 670),fill='#5B6573',width=3)
    img.save(DIAG/filename)

def figure_after(anchor, path, caption):
    p=insert_after(anchor); p.alignment=1; run=p.add_run(); run.add_picture(str(path),width=Inches(6.25))
    for node in run._r.xpath('.//wp:docPr'):
        node.set('descr',caption); node.set('title',caption)
    cap=insert_after(p,caption,'Caption'); cap.alignment=1
    return cap

def paragraph_after(anchor, text):
    return insert_after(anchor,text,'Normal')

def table_after(doc, anchor, rows):
    table=doc.add_table(rows=1,cols=2); table.style='Table Grid'; c45.set_table_geometry(table,[2600,6760]); c45.set_repeat_header(table.rows[0])
    for i,value in enumerate(['Élément','Description']):
        cell=table.rows[0].cells[i]; c45.set_cell_shading(cell,c45.LIGHT); p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); c45.apply_font(p.add_run(value),10,c45.NAVY,True)
    for label,value in rows:
        cells=table.add_row().cells
        for i,text in enumerate([label,value]):
            p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08; c45.apply_font(p.add_run(text),9.2)
    anchor._p.addnext(table._tbl)
    xml = OxmlElement('w:p'); table._tbl.addnext(xml)
    return Paragraph(xml, anchor._parent)

def scenario(name, principal, secondary, objective, pre, post, nominal, alternatives, exceptions):
    return [('Cas d’utilisation',name),('Acteur principal',principal),('Acteurs secondaires',secondary),('Objectif',objective),('Préconditions',pre),('Postconditions',post),('Scénario nominal',nominal),('Scénarios alternatifs',alternatives),('Exceptions',exceptions)]

def main():
    specs=[
        ('uc_admin_saas.png','Figure 3.5 — Diagramme de cas d’utilisation détaillé : administrateur SaaS.',['Administrateur SaaS'],['Stripe','SMTP','Gemini'],['Traiter les demandes de marketplace','Gérer banques, stores et modules','Gérer offres, abonnements et paiements','Administrer utilisateurs et concessionnaires','Consulter audit et indicateurs','Interroger l’assistant IA'],
         'UC-01 — Gouvernance de la plateforme SaaS','Administrateur SaaS','Stripe, SMTP et Gemini','Administrer l’offre globale et les tenants Matchia.','Utilisateur authentifié avec le rôle ADMIN_SAAS.','Les données et décisions sont persistées, notifiées ou auditées selon le workflow.','L’administrateur consulte le tableau de bord, gère les ressources globales, instruit les demandes et supervise abonnements, paiements et audit.','Une demande peut être rejetée avec motif ; une action peut être refusée si le statut ne permet pas la transition.','Erreur de validation, ressource introuvable, indisponibilité d’un service externe ou droit insuffisant.'),
        ('uc_admin_banque.png','Figure 3.6 — Diagramme de cas d’utilisation détaillé : administrateur Banque.',['Administrateur Banque'],['SMTP'],['Configurer marketplace et branding','Gérer utilisateurs et produits','Gérer concessionnaires et partenariats','Gérer contrats et publications','Traiter demandes de financement','Consulter abonnement et notifications'],
         'UC-02 — Exploitation d’une marketplace bancaire','Administrateur Banque','SMTP','Configurer et exploiter exclusivement le tenant bancaire concerné.','Compte ADMIN_BANK actif, marketplace et abonnement valides.','La configuration ou la décision est enregistrée dans le périmètre de la banque.','L’administrateur se connecte, configure ses ressources, traite les partenaires et instruit les dossiers de financement de son tenant.','Une publication peut être rejetée ; un abonnement peut faire l’objet d’une demande de renouvellement ou d’upgrade.','Tentative d’accès à une ressource d’une autre banque, données invalides ou tenant inactif.'),
        ('uc_dealer.png','Figure 3.7 — Diagramme de cas d’utilisation détaillé : concessionnaire.',['Concessionnaire'],['Administrateur Banque'],['Gérer profil et documents','Gérer catalogue, stock et médias','Demander un partenariat','Gérer contrat partenaire','Soumettre une publication produit','Suivre financements de ses produits'],
         'UC-03 — Gestion de l’activité concessionnaire','Concessionnaire','Administrateur Banque','Centraliser le catalogue et collaborer avec plusieurs banques.','Compte DEALER_ADMIN actif et dossier concessionnaire approuvé.','Le produit, partenariat, contrat ou suivi est mis à jour dans l’espace du concessionnaire.','Le concessionnaire met à jour ses produits et son stock, demande des partenariats, gère les contrats puis soumet les produits à publication.','La banque peut accepter, rejeter, suspendre ou terminer la relation ; une publication peut rester non visible.','Dossier non conforme, partenariat dupliqué, contrat non actif ou stock indisponible.'),
        ('uc_client.png','Figure 3.8 — Diagramme de cas d’utilisation détaillé : client.',['Client'],['SMTP'],['Créer et vérifier son compte','Gérer son profil','Consulter produits et simulateur','Constituer un dossier de financement','Téléverser les pièces requises','Suivre la décision et notifications'],
         'UC-04 — Parcours client de financement','Client','SMTP et administrateur Banque','Préparer, soumettre et suivre une demande de financement.','Client inscrit et rattaché à une marketplace bancaire active.','Le dossier est enregistré en brouillon, soumis ou traité selon les actions réalisées.','Le client consulte un produit, simule, crée un brouillon, ajoute les pièces exigées puis soumet sa demande.','Le client peut corriger son brouillon ; il reçoit une décision acceptée ou rejetée.','Compte non vérifié, pièces manquantes, produit hors tenant ou délai de code e-mail expiré.'),
        ('uc_onboarding.png','Figure 3.9 — Diagramme de cas d’utilisation détaillé : onboarding bancaire.',['Représentant Banque'],['Administrateur SaaS','Stripe','SMTP'],['Saisir informations banque','Vérifier adresse e-mail','Sélectionner stores et modules','Soumettre demande','Approuver ou rejeter','Payer et activer marketplace'],
         'UC-05 — Demande de création et activation d’une marketplace','Représentant de la banque','Administrateur SaaS, Stripe et SMTP','Obtenir une marketplace active à partir d’une demande validée et payée.','Formulaire accessible ; e-mail non déjà utilisé ; offres disponibles.','Demande persistée ; après paiement confirmé, banque, marketplace, abonnement et compte sont activés.','Le représentant renseigne les données, vérifie l’e-mail, choisit stores/modules et soumet. Le SaaS analyse et approuve. Le contact règle le paiement et le backend active les ressources.','Le SaaS rejette la demande ou le paiement est annulé/non confirmé ; l’activation ne se produit pas.','Code incorrect/expiré, données invalides, paiement refusé ou service e-mail indisponible.'),
        ('uc_partnership.png','Figure 3.10 — Diagramme de cas d’utilisation détaillé : partenariat et contrat.',['Concessionnaire'],['Administrateur Banque','SMTP'],['Rechercher une banque compatible','Créer demande partenariat','Accepter ou rejeter partenariat','Créer et envoyer contrat','Accepter contrat','Soumettre publication produit'],
         'UC-06 — Partenariat, contrat et publication','Concessionnaire','Administrateur Banque et SMTP','Établir une relation commerciale permettant la publication d’un produit.','Concessionnaire approuvé ; store compatible ; banque accessible.','Partenariat et contrat actifs, ou décision motivée en cas de rejet ; publication enregistrée.','Le concessionnaire demande un partenariat. La banque le traite. Un contrat est préparé, envoyé et accepté. Le concessionnaire soumet ensuite son produit à la publication.','Le partenariat est rejeté, suspendu ou terminé ; le contrat expire ou est annulé ; la publication est refusée.','Relation identique déjà existante, contrat incomplet, produit inactif ou utilisateur non autorisé.'),
        ('uc_financing.png','Figure 3.11 — Diagramme de cas d’utilisation détaillé : demande de financement.',['Client'],['Administrateur Banque','Concessionnaire','SMTP'],['Créer brouillon depuis simulation','Ajouter ou retirer documents','Soumettre dossier PENDING','Consulter et instruire dossier','Accepter ou rejeter avec motif','Réserver stock et notifier'],
         'UC-07 — Instruction d’une demande de financement','Client','Administrateur Banque, concessionnaire et SMTP','Faire traiter un dossier de financement dans le tenant concerné.','Client connecté ; store actif ; produit éligible ; exigences documentaires disponibles.','Dossier ACCEPTED ou REJECTED ; client informé ; stock réservé si le produit dealer est accepté.','Le client crée un brouillon et charge les pièces. Après la soumission, la banque consulte le dossier, accepte ou rejette ; le système informe le client et réserve le stock si nécessaire.','Le client complète les pièces manquantes ; la banque rejette avec motif ; le concessionnaire consulte le nouveau statut de son produit.','Soumission d’un dossier incomplet, statut non PENDING, motif de rejet absent, accès hors tenant ou stock insuffisant.')]
    for filename,caption,left,right,uses,*_ in specs: usecase_diagram(filename,caption.split('— ')[1].replace('.',''),left,right,uses)
    doc=Document(SOURCE)
    start=find_heading(doc,'3.6.2'); end=find_heading(doc,'3.6.3'); remove_between(start,end)
    # Keep figure numbering coherent after replacing the prior three illustrative detailed diagrams.
    for p in doc.paragraphs:
        if p.text.startswith('Figure 3.8 — Diagramme de classes'): p.text=p.text.replace('Figure 3.8','Figure 3.12'); p.style='Caption'
    current=paragraph_after(start,'Les figures suivantes complètent la vue globale par des cas d’utilisation détaillés. Chaque diagramme conserve la frontière du système Matchia, représente les acteurs à l’extérieur et associe les objectifs métier sous forme d’ellipses. Le tableau placé après chaque figure formalise le scénario académique correspondant.')
    for filename,caption,left,right,uses,name,principal,secondary,objective,pre,post,nominal,alternatives,exceptions in specs:
        current=figure_after(current,DIAG/filename,caption)
        current=paragraph_after(current,'Le diagramme est suivi d’un scénario afin de préciser les conditions de déclenchement, les transitions possibles et les résultats attendus.')
        current=table_after(doc,current,scenario(name,principal,secondary,objective,pre,post,nominal,alternatives,exceptions))
        current=paragraph_after(current,'')
    doc.core_properties.title='Rapport PFE — Matchia : cas d’utilisation détaillés complets'
    doc.save(OUT); print(OUT)

if __name__=='__main__': main()
