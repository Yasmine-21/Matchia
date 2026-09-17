from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(r"D:\PFE M2\Platforme SaaS")
PNG_DIR = ROOT / "deliverables" / "use_cases_drawio" / "aperçus_png"
OUTPUT = ROOT / "document_work" / "rendered_usecases_report" / "usecases_visual_qa.docx"

ITEMS = [
    ("UC-01_Authentification.png", "Figure 3.5 — Authentification"),
    ("UC-02_Demande_marketplace.png", "Figure 3.6 — Demande de création d’une marketplace"),
    ("UC-03_Traitement_paiement_activation.png", "Figure 3.8 — Traitement, paiement Stripe et activation"),
    ("UC-04_Inscription_concessionnaire.png", "Figure 3.10 — Inscription et validation d’un concessionnaire"),
    ("UC-05_Gestion_publication_produit.png", "Figure 3.12 — Gestion et publication des produits"),
    ("UC-06_Demande_financement.png", "Figure 3.13 — Demande de financement"),
]


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    for index, (filename, caption) in enumerate(ITEMS):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(PNG_DIR / filename), width=Inches(6.25))

        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cp.add_run(caption)
        run.italic = True
        run.font.size = Pt(10)
        if index != len(ITEMS) - 1:
            doc.add_page_break()

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
