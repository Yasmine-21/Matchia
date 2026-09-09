from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "document_work" / "Architecture_Matchia_Redaction_Amelioree.docx"
LOGICAL = Path(r"C:\Users\ASUS\AppData\Local\Temp\codex-clipboard-51bf0c32-504e-4d1d-a3f6-e8eb1e562085.png")
PHYSICAL = Path(r"C:\Users\ASUS\AppData\Local\Temp\codex-clipboard-61badc67-f4d7-48eb-8a0a-456fab894b08.png")

NAVY = "0B2A52"
BLUE = "2E78C7"
LIGHT_BLUE = "EAF3FC"
LIGHT_GRAY = "D9E1EA"
TEXT = "1F2937"


def set_run_font(run, name="Times New Roman", size=11.5, bold=None, italic=None, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(16)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in [
        ("Heading 1", 16, 14, 8),
        ("Heading 2", 13.5, 12, 6),
        ("Heading 3", 12, 9, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.font.size = Pt(9.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(TEXT)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(10)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.append(fld)


def add_body(doc, text, lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.6)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if lead:
        set_run_font(p.add_run(lead), bold=True)
    set_run_font(p.add_run(text))
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_GRAY, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        elem = borders.find(qn(tag))
        if elem is None:
            elem = OxmlElement(tag)
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:color"), color)


def add_comparison_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Cm(3.2), Cm(5.7), Cm(7.0)]
    headers = ["Vue", "Question traitée", "Éléments représentés"]
    for i, cell in enumerate(table.rows[0].cells):
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(headers[i]), name="Arial", size=10, bold=True, color="FFFFFF")

    rows = [
        ("Logique", "Comment le logiciel est-il organisé ?",
         "Vue React, contrôleurs REST, services métier, modèle, repositories et persistance."),
        ("Physique", "Quels composants communiquent à l’exécution ?",
         "Navigateur, frontend, backend, PostgreSQL et services externes."),
    ]
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[i])
            set_cell_shading(cells[i], "FFFFFF" if row_idx % 2 == 0 else LIGHT_BLUE)
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=10, bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc, path, caption, alt_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline = run.add_picture(str(path), width=Cm(16.4))
    doc_pr = inline._inline.docPr
    doc_pr.set("title", alt_text)
    doc_pr.set("descr", alt_text)
    cap = doc.add_paragraph(caption, style="Caption")
    cap.paragraph_format.keep_with_next = False


def build():
    if not LOGICAL.exists() or not PHYSICAL.exists():
        missing = [str(p) for p in (LOGICAL, PHYSICAL) if not p.exists()]
        raise FileNotFoundError("Images introuvables : " + ", ".join(missing))

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    configure_styles(doc)

    footer = section.footer
    add_page_number(footer.paragraphs[0])

    title = doc.add_paragraph("Architecture de la solution Matchia", style="Title")
    title.paragraph_format.keep_with_next = True
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    add_body(
        doc,
        "La solution Matchia repose sur une architecture client-serveur en couches. Le frontend, développé avec React et TypeScript, constitue une application web autonome qui communique avec une API REST réalisée avec Spring Boot. Le backend centralise la sécurité, les règles métier et l’accès aux données. Cette séparation permet de faire évoluer l’interface et la logique applicative de manière indépendante, tout en conservant un contrat d’échange explicite fondé sur HTTP et JSON."
    )
    add_body(
        doc,
        " Matchia ne met pas en œuvre une architecture microservices : les capacités fonctionnelles sont regroupées dans une même application backend. Ce backend reste cependant organisé en modules métier cohérents, notamment l’administration SaaS, les marketplaces, les stores et modules, les produits, les abonnements, les partenariats, les contrats et les demandes de financement. La qualification la plus précise est donc celle d’une architecture monolithique modulaire en couches, avec un frontend séparé et une gestion multi-tenant logique.",
        lead="Choix architectural retenu."
    )
    add_body(
        doc,
        "Deux vues complémentaires sont nécessaires pour décrire cette organisation. La vue logique explique la répartition des responsabilités dans le code ; la vue physique montre les composants qui communiquent pendant l’exécution."
    )
    add_comparison_table(doc)

    doc.add_heading("Architecture logique", level=1)
    add_body(
        doc,
        "L’architecture logique décrit l’enchaînement des responsabilités depuis l’interaction utilisateur jusqu’à la persistance. La figure 1 propose une lecture inspirée du modèle MVC. Cette correspondance est pédagogique plutôt que stricte : la vue n’est pas générée par le serveur, mais portée par une application React séparée, tandis que le backend expose des ressources REST et suit principalement l’organisation Controller-Service-Repository."
    )
    add_figure(
        doc,
        LOGICAL,
        "Figure 1 — Architecture logique de la plateforme Matchia",
        "Architecture logique de Matchia montrant React, les contrôleurs REST, les services métier, les repositories JPA et PostgreSQL."
    )

    doc.add_heading("Couche de présentation", level=2)
    add_body(
        doc,
        "La couche de présentation est développée avec React et TypeScript. Elle regroupe les pages, les composants d’interface, les layouts, la navigation et les services HTTP utilisés par les espaces Public, Client, Administrateur Banque, Concessionnaire et Administrateur SaaS. Elle collecte les actions de l’utilisateur, applique les validations d’ergonomie immédiates et transmet les demandes au backend. Les décisions de sécurité et les règles métier sensibles ne reposent toutefois jamais uniquement sur le navigateur."
    )

    doc.add_heading("Couche API et contrôleurs", level=2)
    add_body(
        doc,
        "Les contrôleurs Spring Boot constituent le point d’entrée du backend. Ils exposent les endpoints REST, reçoivent les paramètres et les corps de requête, déclenchent la validation des données et délèguent le traitement aux services appropriés. Les échanges avec le frontend utilisent principalement le format JSON. Les objets de transfert de données, ou DTO, définissent le contrat de l’API et limitent les informations exposées ; ils appartiennent donc à la frontière applicative plutôt qu’au modèle de persistance lui-même."
    )

    doc.add_heading("Services métier et modèle", level=2)
    add_body(
        doc,
        "La couche métier porte les règles qui donnent son sens fonctionnel à la plateforme. Elle orchestre notamment la création et la configuration des marketplaces, l’affectation des stores et modules, la gestion des produits, la validation des demandes, les abonnements et paiements, les partenariats avec les concessionnaires, les contrats et l’instruction des dossiers de financement. Les entités représentent l’état persistant du domaine, tandis que les services contrôlent les autorisations, les transitions d’état et la cohérence entre les agrégats. Cette séparation facilite les tests et évite de disperser les règles métier dans les contrôleurs ou dans l’interface."
    )

    doc.add_heading("Accès aux données et persistance", level=2)
    add_body(
        doc,
        "L’accès aux données est assuré par les repositories Spring Data JPA. Ceux-ci encapsulent les opérations de lecture et d’écriture et isolent les services métier des détails du pilote JDBC et des requêtes de persistance. Hibernate réalise la correspondance entre les entités Java et les tables PostgreSQL. Le flux logique principal est ainsi : React, contrôleurs REST, services métier, repositories JPA, puis PostgreSQL."
    )

    doc.add_heading("Sécurité et multi-tenancy", level=2)
    add_body(
        doc,
        "La sécurité et le contexte tenant sont des préoccupations transversales à toutes les couches du backend. Spring Security, les jetons JWT, les rôles et les contrôles applicatifs déterminent les opérations autorisées. La multi-tenancy repose sur une base et un schéma partagés : la banque constitue l’ancrage du tenant et les données sont reliées à son périmètre métier. L’isolation est donc logique, au niveau des relations et des requêtes filtrées, et non physique au moyen d’une base distincte pour chaque établissement."
    )

    doc.add_heading("Architecture physique", level=1)
    add_body(
        doc,
        "L’architecture physique identifie les composants mobilisés pendant l’exécution et les protocoles qui les relient. Elle ne décrit pas ici la chaîne CI/CD ni le mécanisme de conteneurisation, mais uniquement le trajet des requêtes et les dépendances nécessaires au fonctionnement de la plateforme."
    )
    add_figure(
        doc,
        PHYSICAL,
        "Figure 2 — Architecture physique de la plateforme Matchia",
        "Architecture physique de Matchia montrant les utilisateurs, le frontend, le backend, PostgreSQL, Stripe, Google Gemini et le serveur SMTP."
    )

    doc.add_heading("Accès utilisateur et frontend", level=2)
    add_body(
        doc,
        "Les utilisateurs accèdent à Matchia depuis un navigateur web. Le frontend React, compilé avec Vite et servi par Nginx dans l’environnement de déploiement présenté, fournit les ressources statiques et exécute l’interface dans le navigateur. Les échanges sont protégés par HTTPS. Le frontend ne communique ni directement avec PostgreSQL ni avec les services métier externes : il transmet les opérations applicatives au backend par l’intermédiaire de l’API REST."
    )

    doc.add_heading("Backend et base PostgreSQL", level=2)
    add_body(
        doc,
        "Le backend Spring Boot constitue le centre de confiance de l’application. Il authentifie les requêtes, applique les droits liés aux rôles et au tenant, exécute les règles métier et prépare les réponses retournées au frontend. L’accès à PostgreSQL s’effectue à travers JPA, Hibernate et le pilote JDBC. La base conserve notamment les utilisateurs, banques, marketplaces, stores, modules, produits, abonnements, paiements, demandes, partenariats, contrats et dossiers de financement."
    )
    add_body(
        doc,
        "Dans ce modèle multi-tenant, plusieurs banques utilisent la même infrastructure applicative et la même base relationnelle. Les services doivent donc systématiquement appliquer le contexte de banque et les restrictions de rôle avant toute consultation ou modification. Ce contrôle côté serveur constitue la garantie principale d’étanchéité entre les tenants."
    )

    doc.add_heading("Services externes", level=2)
    add_body(
        doc,
        "Le backend communique avec Stripe au moyen d’une API HTTPS pour préparer et confirmer les paiements associés aux abonnements. Google Gemini intervient dans les fonctions d’assistance intelligente ; l’appel est effectué par le backend afin de conserver le contrôle sur le contexte envoyé, de valider les traitements générés et d’éviter tout accès direct du modèle à PostgreSQL. Enfin, le serveur SMTP prend en charge les courriers de vérification, les notifications de traitement et les communications liées aux parcours métier. Ces dépendances restent extérieures au cœur applicatif et sont accessibles au moyen de paramètres de configuration dédiés."
    )

    doc.add_heading("Synthèse du choix architectural", level=1)
    add_body(
        doc,
        "L’architecture retenue répond aux besoins actuels de Matchia par un compromis cohérent entre simplicité d’exploitation et séparation des responsabilités. Le frontend autonome fournit une expérience adaptée aux différents profils, tandis que le monolithe modulaire centralise la sécurité et la cohérence transactionnelle. Les couches Controller, Service et Repository rendent le code plus lisible et testable ; PostgreSQL offre un modèle relationnel adapté aux liens entre banques, marketplaces, utilisateurs, offres et processus de financement."
    )
    add_body(
        doc,
        "Cette organisation permet également une évolution progressive. Si la charge, l’autonomie des équipes ou les contraintes réglementaires l’exigent ultérieurement, certains domaines pourront être isolés derrière des contrats plus indépendants. Une telle évolution devra toutefois être justifiée par des besoins mesurés. Dans l’état actuel du projet, l’architecture monolithique modulaire en couches demeure plus adaptée qu’une architecture microservices, car elle limite la complexité distribuée sans sacrifier la modularité fonctionnelle."
    )

    doc.core_properties.title = "Architecture de la solution Matchia"
    doc.core_properties.subject = "Architecture logique et physique de la plateforme Matchia"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "Matchia, architecture, React, Spring Boot, PostgreSQL, multi-tenant"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
