from copy import deepcopy
from pathlib import Path
import math

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\PFE M2\Platforme SaaS")
INPUT = Path(r"D:\PFE - BRI Technology\Rapport\Master Report.docx")
OUTPUT = ROOT / "document_work" / "Master Report - Chapitre 6 DevOps CI CD Azure détaillé.docx"
FIG = ROOT / "figure"
GEN = ROOT / "document_work" / "chapter6_complete_figures"


def get_font(size, bold=False):
    filename = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / filename), size)


def text_center(draw, rect, text, font, fill="#1f2937"):
    x1, y1, x2, y2 = rect
    lines, current = [], ""
    for word in text.split():
        trial = (current + " " + word).strip()
        if not current or draw.textbbox((0, 0), trial, font=font)[2] <= x2 - x1 - 32:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    line_h = font.getbbox("Ag")[3] + 5
    y = (y1 + y2 - line_h * len(lines)) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=font, fill=fill)
        y += line_h


def box(draw, rect, text, fill="#edf5ff", outline="#5d82a8"):
    draw.rounded_rectangle(rect, radius=22, fill=fill, outline=outline, width=3)
    text_center(draw, rect, text, get_font(23, True))


def arrow(draw, start, end, color="#3f5469"):
    draw.line((*start, *end), fill=color, width=4)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    head = [end]
    for delta in (math.pi - .42, math.pi + .42):
        head.append((end[0] + length * math.cos(angle + delta), end[1] + length * math.sin(angle + delta)))
    draw.polygon(head, fill=color)


def make_overview():
    GEN.mkdir(parents=True, exist_ok=True)
    out = GEN / "fig_6_1_vue_chaine_devops.png"
    im = Image.new("RGB", (2200, 700), "white")
    draw = ImageDraw.Draw(im)
    items = [
        (80, "Développement"), (410, "Git et GitHub"), (740, "Jenkins"),
        (1070, "Build Tests SonarQube"), (1400, "Docker Hub"), (1730, "Microsoft Azure"),
    ]
    for x, label in items:
        box(draw, (x, 260, x + 270, 410), label)
    for (x, _), (next_x, _) in zip(items, items[1:]):
        arrow(draw, (x + 280, 335), (next_x - 10, 335))
    text_center(draw, (100, 500, 2100, 610), "Code source → validation automatisée → images conteneurisées → déploiement Cloud", get_font(28, True), "#374151")
    im.save(out)
    return out


def make_repo_structure():
    out = GEN / "fig_6_2_organisation_depot.png"
    im = Image.new("RGB", (1600, 1050), "white")
    draw = ImageDraw.Draw(im)
    draw.rounded_rectangle((500, 60, 1100, 160), radius=20, fill="#dbeafe", outline="#5d82a8", width=3)
    text_center(draw, (500, 60, 1100, 160), "Dépôt GitHub Matchia", get_font(30, True))
    children = [
        (120, 300, "MatchiaFrontend", "React TypeScript Vite Dockerfile"),
        (610, 300, "MatchiaBackend", "Spring Boot Java 17 Dockerfile"),
        (1100, 300, "Fichiers racine", "Jenkinsfile docker-compose.yml"),
    ]
    for x, y, title, details in children:
        arrow(draw, (800, 160), (x + 190, y))
        draw.rounded_rectangle((x, y, x + 380, y + 220), radius=20, fill="#eefaf6", outline="#587b78", width=3)
        text_center(draw, (x + 20, y + 25, x + 360, y + 100), title, get_font(25, True))
        text_center(draw, (x + 20, y + 105, x + 360, y + 195), details, get_font(19))
    text_center(draw, (150, 790, 1450, 920), "Les composants applicatifs et les fichiers de déploiement sont versionnés dans le même dépôt afin que Jenkins récupère une version cohérente de la solution.", get_font(23), "#374151")
    im.save(out)
    return out


def make_global_architecture():
    out = GEN / "fig_6_20_architecture_globale.png"
    im = Image.new("RGB", (2200, 2050), "white")
    draw = ImageDraw.Draw(im)
    top = [(100, "Développeur\ngit push"), (450, "GitHub"), (800, "Webhook"), (1150, "Cloudflare\nTunnel"), (1500, "Jenkins")]
    for x, label in top:
        box(draw, (x, 100, x + 260, 230), label)
    for (x, _), (nx, _) in zip(top, top[1:]):
        arrow(draw, (x + 270, 165), (nx - 10, 165))
    box(draw, (1300, 400, 1580, 530), "Build Backend\nMaven")
    box(draw, (1650, 400, 1930, 530), "Build Frontend\nNode.js Vite")
    arrow(draw, (1630, 230), (1440, 400))
    arrow(draw, (1680, 230), (1790, 400))
    stages = [(1500, 650, "Tests Backend"), (1500, 830, "SonarQube"), (1500, 1010, "Docker Build\nDocker Hub"), (1500, 1190, "Azure CLI")]
    for x, y, label in stages:
        box(draw, (x, y, x + 340, y + 120), label)
    arrow(draw, (1440, 530), (1640, 650))
    arrow(draw, (1790, 530), (1710, 650))
    for (_, y, _), (_, ny, _) in zip(stages, stages[1:]):
        arrow(draw, (1670, y + 120), (1670, ny))
    # Environment target: separate deployment and persistent data paths so that flows remain readable.
    draw.rounded_rectangle((150, 1500, 1950, 1970), radius=28, outline="#476178", width=3, fill="#fbfdff")
    draw.text((195, 1540), "Microsoft Azure", font=get_font(30, True), fill="#1f2937")
    box(draw, (350, 1660, 750, 1790), "Container App\nFrontend")
    box(draw, (1050, 1660, 1450, 1790), "Container App\nBackend")
    box(draw, (1050, 1840, 1370, 1950), "PostgreSQL\nAzure")
    box(draw, (1550, 1840, 1870, 1950), "Azure Files\napp uploads")
    # Azure CLI deploys both application revisions through an explicit deployment junction.
    box(draw, (1450, 1350, 1790, 1460), "Déploiement des\nrévisions")
    arrow(draw, (1670, 1310), (1620, 1350))
    draw.line((1620, 1460, 1620, 1590), fill="#3f5469", width=4)
    draw.line((550, 1590, 1250, 1590), fill="#3f5469", width=4)
    arrow(draw, (550, 1590), (550, 1660))
    arrow(draw, (1250, 1590), (1250, 1660))
    arrow(draw, (750, 1725), (1050, 1725))
    arrow(draw, (1250, 1790), (1210, 1840))
    arrow(draw, (1450, 1725), (1710, 1840))
    im.save(out)
    return out


def first_run(paragraph):
    return next((r for r in paragraph.runs if r.text), None)


def copy_ppr(target, source):
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if source._p.pPr is not None:
        target._p.insert(0, deepcopy(source._p.pPr))


def add_before(anchor, text, template, alignment=None):
    p = anchor.insert_paragraph_before()
    copy_ppr(p, template)
    run = p.add_run(text)
    src = first_run(template)
    if src is not None and src._r.rPr is not None:
        run._r.insert(0, deepcopy(src._r.rPr))
    if alignment is not None:
        p.alignment = alignment
    return p


def add_code(anchor, text, template):
    p = add_before(anchor, text, template)
    p.paragraph_format.left_indent = Inches(.35)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    for run in p.runs:
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(9.5)
    return p


def move_before(anchor, element):
    element.getparent().remove(element)
    anchor._p.addprevious(element)


def add_figure(document, anchor, image_path, caption, caption_template):
    image = Image.open(image_path)
    width = Inches(5.9 if image.width / image.height >= 1.35 else 4.8)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(image_path), width=width)
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    move_before(anchor, p._p)
    caption_p = add_before(anchor, caption, caption_template, WD_ALIGN_PARAGRAPH.CENTER)
    caption_p.paragraph_format.space_before = Pt(3)
    caption_p.paragraph_format.space_after = Pt(7)
    return caption_p


def remove_paragraph(p):
    p._element.getparent().remove(p._element)
    p._p = p._element = None


def main():
    overview, repo, global_arch = make_overview(), make_repo_structure(), make_global_architecture()
    doc = Document(INPUT)
    paragraphs = doc.paragraphs
    chapter_start = next(i for i, p in enumerate(paragraphs) if p.text.strip().startswith("Chapitre 6"))
    old = paragraphs[chapter_start:]
    anchor = paragraphs[chapter_start]
    chapter_template = anchor
    heading_template = paragraphs[chapter_start + 3]
    body_template = paragraphs[chapter_start + 4]
    subheading_template = paragraphs[chapter_start + 8]
    caption_template = next(p for p in paragraphs if "Figure 3.7" in p.text)

    def chapter_title(text): add_before(anchor, text, chapter_template)
    def heading(text): add_before(anchor, text, heading_template)
    def subheading(text): add_before(anchor, text, subheading_template)
    def para(text): add_before(anchor, text, body_template)
    def code(text): add_code(anchor, text, body_template)
    def fig(path, caption): add_figure(doc, anchor, path, caption, caption_template)

    chapter_title("Chapitre 6 : Mise en place de la démarche DevOps, du pipeline CI/CD et déploiement de Matchia sur Microsoft Azure")

    heading("1. Introduction")
    para("Après la conception et le développement de la plateforme Matchia, une démarche de livraison et de déploiement a été mise en place afin que chaque nouvelle version puisse être construite, vérifiée et déployée de manière reproductible.")
    para("L’objectif ne consiste pas uniquement à rendre l’application accessible dans un environnement Cloud. Il vise également à automatiser les opérations précédant le déploiement : centralisation du code source, construction des composants, exécution des tests, analyse de la qualité, génération des images Docker et mise à jour de l’environnement cible.")
    para("La solution repose sur Git et GitHub pour le versionnement, Docker pour la conteneurisation, Jenkins pour l’orchestration du pipeline, SonarQube pour l’analyse de qualité, Docker Hub pour le registre d’images et Microsoft Azure pour l’hébergement. Azure CLI permet à Jenkins d’exécuter les opérations nécessaires à la mise à jour des services Cloud.")
    para("La figure suivante présente la chaîne globale mise en place entre le développement d’une version et son déploiement.")
    fig(overview, "Figure 6.1 — Vue générale de la chaîne DevOps et CI/CD de Matchia")

    heading("2. Démarche DevOps et principes de la CI/CD")
    subheading("2.1. Démarche DevOps")
    para("Le terme DevOps associe les activités de développement et d’exploitation. Il désigne une démarche qui améliore la collaboration entre ces deux domaines et qui automatise les opérations répétitives du cycle de livraison d’une application.")
    para("Dans Matchia, cette démarche relie la gestion du code source, la construction du frontend et du backend, les tests, le contrôle de la qualité, la conteneurisation et le déploiement. Elle favorise des mises en production plus fiables, plus rapides et reproductibles, car chaque version suit les mêmes contrôles avant d’être publiée.")
    subheading("2.2. Intégration et déploiement continus")
    para("L’intégration continue, ou CI, consiste à intégrer régulièrement les modifications du code et à vérifier automatiquement leur validité. Pour Matchia, elle comprend la récupération du code, la construction des composants, les tests backend et l’analyse de qualité du backend et du frontend.")
    code("Code source → Build → Tests → Analyse de qualité")
    para("Le déploiement continu, ou CD, intervient après cette validation technique. Il construit les images Docker, les publie dans Docker Hub, puis met à jour les services hébergés dans Microsoft Azure. Une erreur bloquante interrompt le pipeline afin d’empêcher le déploiement d’une version non validée.")
    code("Version validée → Images Docker → Docker Hub → Microsoft Azure")

    heading("3. Gestion des versions avec Git et GitHub")
    subheading("3.1. Principe de la gestion des versions")
    para("La gestion des versions conserve l’historique des modifications apportées au projet et permet d’identifier précisément chaque évolution du code source. Elle facilite la collaboration, la comparaison entre versions et le retour vers une version antérieure lorsqu’une anomalie est détectée.")
    subheading("3.2. Utilisation de Git")
    para("Git est utilisé localement pour enregistrer les modifications sous forme de commits. Après vérification, le développeur transmet la version concernée vers le dépôt distant ; Git représente ainsi le premier maillon de la chaîne de livraison.")
    code("Développeur → Git")
    subheading("3.3. Centralisation du code avec GitHub")
    para("GitHub héberge le dépôt distant de Matchia. Il centralise le frontend React, le backend Spring Boot, les Dockerfiles, le fichier docker-compose.yml et le Jenkinsfile. Lorsqu’une évolution est prête, la commande git push publie la version concernée et rend celle-ci disponible pour le stage Checkout de Jenkins.")
    code("Développeur → Git → GitHub → Jenkins Checkout")
    para("La figure suivante présente l’organisation fonctionnelle du dépôt utilisée par la chaîne CI/CD.")
    fig(repo, "Figure 6.2 — Organisation du dépôt GitHub de la plateforme Matchia")

    heading("4. Conteneurisation de Matchia")
    subheading("4.1. Principe de la conteneurisation")
    para("La conteneurisation regroupe une application et les dépendances nécessaires à son exécution dans une unité isolée appelée conteneur. Cette approche réduit les écarts de configuration entre la machine de développement, l’environnement de test et l’environnement Cloud, tout en conservant un format de déploiement léger.")
    subheading("4.2. Choix de Docker")
    para("Docker est la technologie retenue pour créer et exécuter les conteneurs de Matchia. Une image Docker constitue le modèle contenant l’application et ses dépendances ; un conteneur est l’instance en cours d’exécution de cette image. Docker facilite ainsi l’isolation du frontend et du backend ainsi que leur intégration avec Jenkins, Docker Hub et Azure.")
    subheading("4.3. Conteneurisation du backend Spring Boot")
    para("Le backend est développé avec Spring Boot et Java 17. Son Dockerfile adopte une construction en plusieurs étapes : Maven produit le fichier JAR, puis une image d’exécution ne conserve que l’environnement Java et l’artefact nécessaire au démarrage. Cette organisation limite la taille de l’image finale et expose les API REST sur le port 8081.")
    para("Avant son intégration dans le pipeline, l’image a été construite et exécutée manuellement afin de vérifier le démarrage correct du backend et de corriger les éventuels problèmes de configuration.")
    fig(FIG / "imageBackend.png", "Figure 6.3 — Dockerfile utilisé pour la conteneurisation du backend Matchia")
    subheading("4.4. Conteneurisation du frontend React")
    para("Le frontend est développé avec React, TypeScript et Vite. Node.js installe les dépendances et produit les fichiers statiques de production ; ceux-ci sont ensuite servis par Nginx sur le port 80. La variable VITE_API_URL est transmise pendant le build afin que la version frontend produite puisse contacter l’URL du backend correspondant à son environnement.")
    para("Cette image a également été testée manuellement avant l’automatisation afin de valider le build Vite, la configuration Nginx et la communication avec le backend.")
    fig(FIG / "imageFrontend.png", "Figure 6.4 — Dockerfile utilisé pour la conteneurisation du frontend Matchia")
    subheading("4.5. Orchestration locale avec Docker Compose")
    para("Docker Compose permet de définir plusieurs services dans un même fichier de configuration. Dans Matchia, il décrit le frontend, le backend, les ports exposés, les variables d’environnement, le réseau de communication et le volume associé aux fichiers téléversés.")
    code("docker compose up -d")
    para("Cette étape locale a permis de valider le fonctionnement conjoint des deux composants avant que leur construction soit intégrée dans Jenkins.")
    fig(FIG / "dockercompose.png", "Figure 6.5 — Configuration locale des services Matchia avec Docker Compose")

    heading("5. Mise en place du pipeline CI/CD avec Jenkins")
    subheading("5.1. Rôle de Jenkins")
    para("Jenkins est un serveur d’automatisation chargé d’exécuter et d’orchestrer les étapes d’un pipeline CI/CD. Dans Matchia, il coordonne GitHub, Maven, Node.js, les tests, SonarQube, Docker, Docker Hub et Azure CLI. Jenkins est exécuté dans un conteneur Docker afin de disposer d’un environnement de travail reproductible.")
    para("Les plugins Jenkins permettent notamment l’exécution des pipelines, l’intégration GitHub, l’utilisation de Node.js, la communication avec SonarQube et la gestion des credentials. Les informations sensibles, telles que les accès à Docker Hub, aux services d’analyse ou à Azure, sont enregistrées dans les credentials Jenkins plutôt que dans le code source.")
    fig(FIG / "jenkins.png", "Figure 6.6 — Construction de l’environnement Docker utilisé pour Jenkins")
    subheading("5.2. Définition du pipeline avec le Jenkinsfile")
    para("Le pipeline de Matchia est défini dans un Jenkinsfile placé à la racine du dépôt. Cette approche, appelée Pipeline as Code, permet de versionner la définition de la chaîne CI/CD avec le reste du projet et d’assurer que Jenkins exécute les étapes correspondant exactement à la version récupérée.")
    code("Checkout → Build Backend → Build Frontend → Tests Backend → SonarQube Backend → SonarQube Frontend → Docker Build → Docker Push → Deploy Backend → Deploy Frontend")
    para("La séparation des stages facilite l’identification de l’étape responsable d’un échec et rend l’exécution plus lisible depuis l’interface Jenkins.")
    fig(FIG / "CréationPipeline.png", "Figure 6.7 — Création du pipeline Jenkins associé au projet Matchia")
    subheading("5.3. Déclenchement avec GitHub Webhook")
    para("Un Webhook GitHub est configuré afin de déclencher automatiquement le pipeline lorsqu’un push est effectué sur le dépôt. GitHub envoie alors un événement vers Jenkins, qui lance le Jenkinsfile sans intervention manuelle.")
    code("git push → GitHub → Webhook → Jenkins")
    fig(FIG / "ConfigWebhookSurGithub.png", "Figure 6.8 — Configuration du Webhook GitHub déclenchant Jenkins")
    subheading("5.4. Utilisation de Cloudflare Tunnel")
    para("Pendant la phase de développement, Jenkins était exécuté localement et son adresse localhost ne pouvait pas être appelée directement par GitHub. Cloudflare Tunnel a donc été utilisé temporairement pour fournir une adresse HTTPS publique redirigeant les requêtes vers l’instance Jenkins locale.")
    code("GitHub → Cloudflare Tunnel → Jenkins local")
    fig(FIG / "CréationTunnelCloudflareTempVersJenkins.png", "Figure 6.9 — Exposition temporaire de Jenkins au moyen de Cloudflare Tunnel")

    heading("6. Exécution automatisée du pipeline CI/CD")
    subheading("6.1. Récupération du code et construction")
    para("Après son déclenchement, Jenkins exécute le stage Checkout afin de récupérer la version du projet disponible sur GitHub. Le backend est construit avec Maven ; dans la configuration finale, le stage Build Backend utilise la commande ./mvnw clean package -DskipTests. Les tests sont volontairement exécutés dans un stage distinct afin de localiser clairement les erreurs de construction et de test.")
    para("Le frontend est construit avec Node.js. La commande npm ci installe les dépendances de manière reproductible, puis npm run build génère la version de production. VITE_API_URL est transmis durant cette étape pour intégrer l’adresse publique du backend dans la version frontend déployée.")
    code("npm ci")
    code("VITE_API_URL=$AZURE_BACKEND_URL npm run build")
    fig(FIG / "CommandeBackend.png", "Figure 6.10 — Exécution de la phase de construction du backend")
    subheading("6.2. Exécution automatique des tests")
    para("Le stage Tests Backend exécute ./mvnw test. Il constitue une barrière de validation : un test en échec interrompt le pipeline avant la création des images Docker ; en cas de succès, l’exécution se poursuit vers l’analyse de qualité.")
    code("Tests réussis → poursuite du pipeline ; Tests échoués → arrêt du pipeline")
    para("La couverture des tests est générée avec JaCoCo et peut ensuite être exploitée par SonarQube. La suite validée comporte notamment 338 tests exécutés, sans échec ni erreur.")
    fig(FIG / "testBackend.png", "Figure 6.11 — Résultat de l’exécution automatique des tests backend")
    subheading("6.3. Analyse de la qualité avec SonarQube")
    para("SonarQube réalise une analyse statique du code afin d’identifier les bugs potentiels, les vulnérabilités, les duplications et les odeurs de code. Le backend utilise les informations de couverture JaCoCo ; une analyse distincte est effectuée pour le frontend React à l’aide du scanner approprié.")
    para("Cette étape complète les tests fonctionnels en fournissant des indicateurs de qualité et de maintenabilité avant la publication des images destinées au déploiement.")
    fig(FIG / "SonarQubeProjects.png", "Figure 6.12 — Projets Matchia analysés dans SonarQube")
    subheading("6.4. Construction et publication des images Docker")
    para("Lorsque la version est validée, Jenkins réutilise les Dockerfiles du frontend et du backend afin de construire automatiquement les deux images Docker. Chaque image reçoit le tag latest ainsi qu’un tag contenant le numéro du build Jenkins ; ce dernier assure la traçabilité entre une exécution du pipeline et l’image produite.")
    para("Les images matchia-backend et matchia-frontend sont ensuite publiées dans Docker Hub, utilisé comme registre distant. Elles deviennent ainsi disponibles pour les services Azure Container Apps qui exécuteront la nouvelle version.")
    code("Jenkins → Docker Build → Docker Hub")
    fig(FIG / "créationRepoDockerHub.png", "Figure 6.13 — Registre Docker Hub utilisé pour les images Matchia")

    heading("7. Déploiement de Matchia sur Microsoft Azure")
    subheading("7.1. Principe du Cloud Computing et choix d’Azure")
    para("Le Cloud Computing fournit à distance les ressources nécessaires à l’hébergement d’une application : puissance de calcul, stockage, base de données et services réseau. Microsoft Azure a été choisi comme environnement cible car il permet d’héberger les conteneurs, la base PostgreSQL et les fichiers persistants dans des services distincts.")
    para("L’environnement Matchia comprend deux Azure Container Apps pour le frontend et le backend, Azure Database for PostgreSQL Flexible Server pour les données métier et Azure Files pour les contenus téléversés. Cette séparation évite qu’un remplacement de conteneur supprime la base de données ou les fichiers de l’application.")
    fig(FIG / "azure.png", "Figure 6.14 — Ressources Azure préparées pour l’hébergement de Matchia")
    subheading("7.2. Préparation de la base PostgreSQL Azure")
    para("La base de données est hébergée dans Azure Database for PostgreSQL Flexible Server. Le transfert initial depuis l’environnement local est effectué avec pg_dump, qui crée une sauvegarde, puis pg_restore, qui restaure les données sur le serveur PostgreSQL Azure.")
    code("PostgreSQL local → pg_dump → pg_restore → PostgreSQL Azure")
    para("Le backend utilise ensuite une chaîne de connexion JDBC fournie sous forme de variables d’environnement. Les paramètres de connexion restent ainsi séparés du code source et peuvent être adaptés à l’environnement de déploiement.")
    fig(FIG / "ConfigurationDataBase.png", "Figure 6.15 — Configuration de la base PostgreSQL Azure de Matchia")
    subheading("7.3. Préparation du stockage persistant avec Azure Files")
    para("Les conteneurs peuvent être remplacés lors de la création d’une nouvelle révision. Azure Files est donc utilisé pour préserver les documents, les images de produits, les logos et les autres fichiers téléversés indépendamment du cycle de vie du backend.")
    para("Le partage est monté dans le backend au niveau du répertoire /app/uploads. Les fichiers restent ainsi disponibles après la mise à jour d’une image ou le redémarrage d’un conteneur.")
    fig(FIG / "StockageAzureDataBase.png", "Figure 6.16 — Ressource Azure Files associée au stockage persistant de Matchia")
    subheading("7.4. Préparation des Azure Container Apps")
    para("Deux Azure Container Apps distinctes exécutent les images publiées par Jenkins : matchia-frontend contient l’application React servie par Nginx, tandis que matchia-backend contient l’API Spring Boot. Le frontend communique avec le backend, qui accède lui-même à PostgreSQL Azure et à Azure Files.")
    para("La séparation des services permet de mettre à jour les deux composants de manière indépendante. L’Ingress d’Azure Container Apps fournit l’accès HTTPS nécessaire à la consultation de la plateforme.")
    fig(FIG / "ConfigContainerFrontend.png", "Figure 6.17 — Configuration de l’Azure Container App frontend")
    fig(FIG / "ConfigurationContainerBackend.png", "Figure 6.18 — Configuration de l’Azure Container App backend")
    subheading("7.5. Déploiement automatisé avec Azure CLI")
    para("Azure CLI est l’interface en ligne de commande de Microsoft Azure. Jenkins orchestre le pipeline, tandis qu’Azure CLI exécute les commandes de mise à jour des Container Apps. Une authentification initiale peut être effectuée par code appareil afin de respecter les contraintes d’authentification multifacteur.")
    code("az login --use-device-code")
    para("Après la publication dans Docker Hub, Jenkins sélectionne l’abonnement configuré et exécute une commande az containerapp update pour mettre à jour le backend puis le frontend. Azure Container Apps crée alors une nouvelle révision correspondant à la version déployée, ce qui facilite le suivi des mises à jour.")
    code("az containerapp update --name matchia-backend --resource-group rg-matchia --image <registre>/matchia-backend:latest")
    fig(FIG / "déploy.png", "Figure 6.19 — Exécution réussie du pipeline Jenkins jusqu’au déploiement sur Azure")

    heading("8. Architecture globale et validation du processus")
    subheading("8.1. Architecture globale")
    para("L’architecture suivante récapitule le parcours d’une version depuis le push réalisé par le développeur. Elle présente les étapes de déclenchement, de construction, de contrôle de qualité, de publication des images et de mise à jour des services Azure.")
    fig(global_arch, "Figure 6.20 — Architecture globale de la chaîne CI/CD et du déploiement Cloud de Matchia")
    subheading("8.2. Validation du processus complet")
    para("La validation de bout en bout consiste à envoyer une modification contrôlée vers la branche suivie par Jenkins. Le webhook GitHub déclenche alors automatiquement le pipeline ; chaque stage doit aboutir avec succès avant le passage au suivant.")
    code("git push origin main")
    code("GitHub ✓ → Webhook ✓ → Jenkins ✓ → Build ✓ → Tests ✓ → SonarQube ✓ → Docker ✓ → Docker Hub ✓ → Azure CLI ✓ → Azure Container Apps ✓")
    para("La capture de l’exécution Jenkins illustre cette validation : les étapes Checkout, Build Backend, Build Frontend, Tests Backend, SonarQube, Docker Build, Docker Push et déploiement des deux services Azure sont réalisées avec succès. La disponibilité de l’application après le déploiement constitue une vérification fonctionnelle complémentaire du processus.")

    heading("9. Conclusion")
    para("Ce chapitre a présenté la démarche DevOps mise en œuvre pour Matchia, depuis la gestion des versions avec Git et GitHub jusqu’au déploiement des composants sur Microsoft Azure. La validation locale de Docker et Docker Compose a précédé l’intégration progressive des étapes de construction, de test, d’analyse SonarQube, de publication Docker Hub et de déploiement dans Jenkins.")
    para("Le pipeline CI/CD automatise désormais la récupération du code, la construction du frontend et du backend, les contrôles de qualité, la création des images Docker et la mise à jour des Azure Container Apps à l’aide d’Azure CLI. Cette organisation rend la livraison des nouvelles versions plus cohérente, contrôlée et reproductible.")

    for p in old:
        remove_paragraph(p)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
