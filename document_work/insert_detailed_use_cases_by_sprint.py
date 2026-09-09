from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\PFE M2\Platforme SaaS")
INPUT = ROOT / "document_work" / "Master Report - backlog chapitre 4 détaillé.docx"
OUTPUT = ROOT / "document_work" / "Master Report - cas utilisation détaillés sprints.docx"
DIAGRAM_DIR = ROOT / "document_work" / "use_cases_sprints_uml"


def get_font(size, bold=False):
    filename = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / filename), size)


def wrap(draw, value, font, width):
    lines, current = [], ""
    for word in value.split():
        candidate = (current + " " + word).strip()
        if not current or draw.textbbox((0, 0), candidate, font=font)[2] <= width:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def dashed_line(draw, start, end, dash=12, gap=8, fill="#444444", width=3):
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    position = 0
    while position < length:
        nxt = min(position + dash, length)
        draw.line((x1 + ux * position, y1 + uy * position, x1 + ux * nxt, y1 + uy * nxt), fill=fill, width=width)
        position += dash + gap


def arrow_head(draw, x, y, angle, fill="#444444"):
    import math
    spread = math.pi / 7
    size = 16
    points = [(x, y)]
    for delta in (math.pi - spread, math.pi + spread):
        points.append((x + size * math.cos(angle + delta), y + size * math.sin(angle + delta)))
    draw.polygon(points, fill=fill)


def draw_actor(draw, x, y, label):
    stroke = "#111111"
    draw.ellipse((x - 18, y, x + 18, y + 36), outline=stroke, width=3)
    draw.line((x, y + 36, x, y + 95), fill=stroke, width=3)
    draw.line((x - 34, y + 55, x + 34, y + 55), fill=stroke, width=3)
    draw.line((x, y + 95, x - 28, y + 135), fill=stroke, width=3)
    draw.line((x, y + 95, x + 28, y + 135), fill=stroke, width=3)
    font = get_font(18, True)
    lines = wrap(draw, label, font, 210)
    yy = y + 148
    for line in lines:
        width = draw.textbbox((0, 0), line, font=font)[2]
        draw.text((x - width / 2, yy), line, font=font, fill=stroke)
        yy += 22


def draw_oval(draw, cx, cy, w, h, label, main=False):
    outline = "#000000"
    fill = "#F2F2F2" if main else "#FFFFFF"
    draw.ellipse((cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2), fill=fill, outline=outline, width=3)
    font = get_font(19 if main else 16, main)
    lines = wrap(draw, label, font, w - 36)
    heights = [draw.textbbox((0, 0), line, font=font)[3] - draw.textbbox((0, 0), line, font=font)[1] for line in lines]
    yy = cy - (sum(heights) + (len(lines) - 1) * 4) / 2
    for line, height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=font)[2]
        draw.text((cx - width / 2, yy), line, font=font, fill="#111111")
        yy += height + 4


def draw_relationship(draw, source, destination, label, relation):
    # Association is solid and has no arrow; include/extend is a dashed directed dependency.
    if relation == "association":
        draw.line((source, destination), fill="#333333", width=3)
        return
    dashed_line(draw, source, destination)
    import math
    angle = math.atan2(destination[1] - source[1], destination[0] - source[0])
    arrow_head(draw, destination[0], destination[1], angle)
    font = get_font(14, False)
    mx, my = (source[0] + destination[0]) / 2, (source[1] + destination[1]) / 2
    draw.rectangle((mx - 52, my - 13, mx + 52, my + 13), fill="white")
    width = draw.textbbox((0, 0), label, font=font)[2]
    draw.text((mx - width / 2, my - 10), label, font=font, fill="#222222")


def build_diagram(path, title, actors_left, actors_right, main_label, children, associations, includes, extends):
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (2200, 1450), "white")
    draw = ImageDraw.Draw(image)
    title_font = get_font(31, True)
    draw.text((80, 36), title, font=title_font, fill="#000000")
    draw.line((80, 90, 2120, 90), fill="#000000", width=2)
    bx1, by1, bx2, by2 = 370, 130, 1830, 1360
    draw.rectangle((bx1, by1, bx2, by2), outline="#000000", width=3)
    draw.text((bx1 + 28, by1 + 20), "Système Matchia", font=get_font(23, True), fill="#000000")

    # Actor positions.
    left_positions = {}
    right_positions = {}
    left_start = 290 if len(actors_left) <= 2 else 220
    left_step = 410 if len(actors_left) <= 2 else 280
    for idx, label in enumerate(actors_left):
        point = (170, left_start + idx * left_step)
        draw_actor(draw, *point, label)
        left_positions[label] = (point[0] + 38, point[1] + 65)
    right_start = 310 if len(actors_right) <= 2 else 220
    right_step = 410 if len(actors_right) <= 2 else 280
    for idx, label in enumerate(actors_right):
        point = (2025, right_start + idx * right_step)
        draw_actor(draw, *point, label)
        right_positions[label] = (point[0] - 38, point[1] + 65)

    positions = {"main": (1100, 300)}
    child_slots = [(700, 520), (1500, 520), (700, 690), (1500, 690), (700, 860), (1500, 860), (700, 1030), (1500, 1030), (700, 1200), (1500, 1200)]
    for index, child in enumerate(children):
        positions[child] = child_slots[index]
    draw_oval(draw, *positions["main"], 520, 125, main_label, main=True)
    for child in children:
        draw_oval(draw, *positions[child], 440, 105, child)

    for source_actor, target in associations:
        if source_actor in left_positions:
            source = left_positions[source_actor]
        else:
            source = right_positions[source_actor]
        target_point = positions["main"] if target == "main" else positions[target]
        draw_relationship(draw, source, (target_point[0] - 220 if source[0] < target_point[0] else target_point[0] + 220, target_point[1]), "", "association")
    def dependency_edges(source, target):
        source_point = positions["main"] if source == "main" else positions[source]
        destination_point = positions["main"] if target == "main" else positions[target]
        source_half_h = 62 if source == "main" else 52
        destination_half_h = 62 if target == "main" else 52
        if destination_point[1] > source_point[1]:
            return (source_point[0], source_point[1] + source_half_h), (destination_point[0], destination_point[1] - destination_half_h)
        return (source_point[0], source_point[1] - source_half_h), (destination_point[0], destination_point[1] + destination_half_h)

    for source, target in includes:
        start, end = dependency_edges(source, target)
        draw_relationship(draw, start, end, "<<include>>", "include")
    for source, target in extends:
        start, end = dependency_edges(source, target)
        draw_relationship(draw, start, end, "<<extend>>", "extend")

    legend_font = get_font(14)
    draw.text((420, 1322), "Association : ligne pleine     <<include>> : sous-fonction obligatoire     <<extend>> : comportement conditionnel", font=legend_font, fill="#222222")
    image.save(path)


def copy_paragraph_properties(target, source):
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if source._p.pPr is not None:
        target._p.insert(0, deepcopy(source._p.pPr))


def styled_insert_before(anchor, text, template, alignment=None):
    paragraph = anchor.insert_paragraph_before()
    copy_paragraph_properties(paragraph, template)
    source_run = next((run for run in template.runs if run.text), None)
    run = paragraph.add_run(text)
    if source_run is not None and source_run._r.rPr is not None:
        run._r.insert(0, deepcopy(source_run._r.rPr))
    if alignment is not None:
        paragraph.alignment = alignment
    return paragraph


def move_before(anchor, element):
    element.getparent().remove(element)
    anchor._p.addprevious(element)


def add_figure(document, anchor, image_path, caption, caption_template):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(image_path), width=Inches(6.15))
    picture._inline.docPr.set('descr', caption)
    picture._inline.docPr.set('title', caption)
    move_before(anchor, paragraph._p)
    return styled_insert_before(anchor, caption, caption_template, WD_ALIGN_PARAGRAPH.CENTER)


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        tc_pr.append(shading)
    shading.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_text(cell, text, bold=False, size=8.8):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_scenario_table(document, anchor, rows):
    table = document.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.autofit = False
    move_before(anchor, table._tbl)
    headers = ("Élément", "Description")
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(1.5 if idx == 0 else 5.0)
        set_shading(cell, "D9D9D9")
        set_cell_text(cell, value, bold=True, size=9)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    table_header = OxmlElement('w:tblHeader')
    table_header.set(qn('w:val'), 'true')
    tr_pr.append(table_header)
    for label, description in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(1.5)
        cells[1].width = Inches(5.0)
        set_cell_text(cells[0], label, bold=True)
        set_cell_text(cells[1], description)
    return table


def scenario_rows(spec):
    return [
        ("Identifiant", spec["id"]),
        ("Sprint associé", spec["sprint"]),
        ("Acteur principal", spec["primary"]),
        ("Acteurs secondaires", spec["secondary"]),
        ("Objectif", spec["objective"]),
        ("Préconditions", spec["preconditions"]),
        ("Déclencheur", spec["trigger"]),
        ("Scénario nominal", spec["nominal"]),
        ("Alternatives et exceptions", spec["alternatives"]),
        ("Postconditions", spec["postconditions"]),
    ]


def make_specs():
    return [
        {
            "id": "UC-01 — Accéder à un espace sécurisé", "sprint": "S1 — Socle SaaS multi-tenant",
            "title": "Accès sécurisé et contexte multi-tenant", "caption": "Figure 3.5 — Diagramme de cas d’utilisation détaillé de l’accès sécurisé multi-tenant",
            "actors_left": ["Administrateur SaaS", "Administrateur Banque", "Concessionnaire", "Client"], "actors_right": [],
            "main": "S’authentifier", "children": ["Vérifier les identifiants", "Charger le rôle", "Appliquer le contexte tenant"],
            "associations": [("Administrateur SaaS", "main"), ("Administrateur Banque", "main"), ("Concessionnaire", "main"), ("Client", "main")],
            "includes": [("main", "Vérifier les identifiants"), ("main", "Charger le rôle"), ("main", "Appliquer le contexte tenant")], "extends": [],
            "primary": "Utilisateur autorisé", "secondary": "—", "objective": "Accéder à l’espace correspondant à son rôle sans consulter les données d’un autre tenant.",
            "preconditions": "Compte actif et identifiants valides.", "trigger": "L’utilisateur soumet le formulaire de connexion.",
            "nominal": "1. L’utilisateur saisit son e-mail et son mot de passe. 2. Le système vérifie les identifiants. 3. Il récupère le rôle associé. 4. Pour les espaces bancaires, il applique le contexte de la banque concernée. 5. Il ouvre l’espace autorisé.",
            "alternatives": "Identifiants erronés ou compte inactif : l’accès est refusé. Jeton expiré : l’utilisateur doit se reconnecter. Une ressource appartenant à un autre tenant est inaccessible.",
            "postconditions": "Une session authentifiée est ouverte avec les droits et le périmètre de données appropriés.",
        },
        {
            "id": "UC-02 — Déposer une demande de marketplace", "sprint": "S2 — Demande de création de marketplace",
            "title": "Demande de création d’une marketplace bancaire", "caption": "Figure 3.6 — Diagramme de cas d’utilisation détaillé de la demande de création d’une marketplace",
            "actors_left": ["Représentant Banque"], "actors_right": ["Service e-mail"],
            "main": "Soumettre la demande", "children": ["Saisir informations banque", "Saisir futur administrateur", "Vérifier l’adresse e-mail", "Configurer la marketplace", "Sélectionner stores et modules", "Consulter le récapitulatif"],
            "associations": [("Représentant Banque", "main"), ("Service e-mail", "Vérifier l’adresse e-mail")],
            "includes": [("main", "Saisir informations banque"), ("main", "Saisir futur administrateur"), ("main", "Vérifier l’adresse e-mail"), ("main", "Configurer la marketplace"), ("main", "Sélectionner stores et modules"), ("main", "Consulter le récapitulatif")], "extends": [],
            "primary": "Représentant de la banque", "secondary": "Service e-mail", "objective": "Transmettre une demande complète en vue de créer une marketplace bancaire.",
            "preconditions": "Le formulaire d’adhésion est accessible et l’adresse e-mail du futur administrateur n’est pas déjà utilisée.", "trigger": "Le représentant accède au formulaire public de demande.",
            "nominal": "1. Le représentant renseigne les données de la banque. 2. Il saisit les informations du futur Administrateur Banque. 3. Il valide le code reçu par e-mail. 4. Il configure le slug et les éléments d’identité de la marketplace. 5. Il choisit les stores et modules. 6. Il vérifie le récapitulatif puis soumet.",
            "alternatives": "Une donnée invalide impose une correction. Un code e-mail incorrect ou expiré peut être renvoyé. Le représentant peut revenir sur les choix de stores ou de modules avant la soumission.",
            "postconditions": "La demande est enregistrée à l’état en attente et une notification est créée pour l’Administrateur SaaS.",
        },
        {
            "id": "UC-03 — Administrer la plateforme SaaS", "sprint": "S3 — Back-office SaaS",
            "title": "Administration et supervision de la plateforme SaaS", "caption": "Figure 3.7 — Diagramme de cas d’utilisation détaillé de l’administration SaaS",
            "actors_left": ["Administrateur SaaS"], "actors_right": ["Gemini"],
            "main": "Administrer la plateforme SaaS", "children": ["Consulter le tableau de bord", "Gérer stores et modules", "Gérer banques et marketplaces", "Gérer utilisateurs et rôles", "Gérer offres et abonnements", "Consulter audit et logs", "Interroger l’assistant IA"],
            "associations": [("Administrateur SaaS", "main"), ("Administrateur SaaS", "Consulter audit et logs"), ("Gemini", "Interroger l’assistant IA")],
            "includes": [("main", "Consulter le tableau de bord"), ("main", "Gérer stores et modules"), ("main", "Gérer banques et marketplaces"), ("main", "Gérer utilisateurs et rôles"), ("main", "Gérer offres et abonnements"), ("main", "Consulter audit et logs")], "extends": [],
            "primary": "Administrateur SaaS", "secondary": "Gemini pour l’assistant analytique", "objective": "Administrer les ressources globales et superviser les tenants de la plateforme.",
            "preconditions": "Utilisateur authentifié avec le rôle Administrateur SaaS.", "trigger": "L’administrateur accède au Back-office SaaS.",
            "nominal": "1. L’administrateur consulte les indicateurs. 2. Il gère les banques, marketplaces, utilisateurs, stores et modules. 3. Il configure les offres et consulte les abonnements. 4. Il consulte les journaux d’audit et les paramètres autorisés.",
            "alternatives": "Une ressource peut être filtrée, recherchée ou désactivée selon son statut. L’assistant IA peut être interrogé pour une consultation analytique.",
            "postconditions": "Les modifications autorisées sont persistées et les actions sensibles sont traçables dans les journaux d’audit.",
        },
        {
            "id": "UC-04 — Traiter une demande et activer les services", "sprint": "S4 — Traitement, paiement et activation",
            "title": "Traitement SaaS, paiement Stripe et activation", "caption": "Figure 3.8 — Diagramme de cas d’utilisation détaillé du traitement, du paiement et de l’activation",
            "actors_left": ["Administrateur SaaS"], "actors_right": ["Stripe", "Service e-mail"],
            "main": "Traiter une demande de marketplace", "children": ["Consulter le dossier", "Approuver la demande", "Rejeter avec motif", "Générer le lien de paiement", "Vérifier le paiement Stripe", "Activer les services", "Envoyer les identifiants"],
            "associations": [("Administrateur SaaS", "main"), ("Stripe", "Vérifier le paiement Stripe"), ("Service e-mail", "Générer le lien de paiement"), ("Service e-mail", "Envoyer les identifiants")],
            "includes": [("main", "Consulter le dossier"), ("Approuver la demande", "Générer le lien de paiement"), ("Générer le lien de paiement", "Vérifier le paiement Stripe"), ("Vérifier le paiement Stripe", "Activer les services"), ("Activer les services", "Envoyer les identifiants")], "extends": [("Approuver la demande", "main"), ("Rejeter avec motif", "main")],
            "primary": "Administrateur SaaS", "secondary": "Stripe et service e-mail", "objective": "Décider d’une demande puis activer les ressources uniquement après confirmation du paiement.",
            "preconditions": "Demande complète au statut en attente.", "trigger": "L’Administrateur SaaS ouvre le détail d’une demande.",
            "nominal": "1. L’administrateur vérifie le dossier. 2. Il approuve la demande. 3. Le système crée l’opération Stripe et envoie le lien de paiement. 4. Le contact bancaire réalise le paiement. 5. Le backend vérifie le statut confirmé auprès de Stripe. 6. Il active la banque, la marketplace, l’abonnement et le compte Administrateur Banque. 7. Il envoie les identifiants de connexion.",
            "alternatives": "L’administrateur rejette la demande avec un motif et un e-mail est envoyé. Si le paiement est annulé ou non confirmé, les ressources restent inactives.",
            "postconditions": "La demande est rejetée avec un motif, ou la banque et ses services sont activés après paiement confirmé.",
        },
        {
            "id": "UC-05 — Configurer et utiliser une marketplace bancaire", "sprint": "S5 — Back-office Banque et marketplace publique",
            "title": "Configuration bancaire et consultation de la marketplace publique", "caption": "Figure 3.9 — Diagramme de cas d’utilisation détaillé de la configuration et de la marketplace publique",
            "actors_left": ["Administrateur Banque", "Visiteur"], "actors_right": [],
            "main": "Gérer la marketplace bancaire", "children": ["Personnaliser identité et contenu", "Configurer stores et modules", "Demander renouvellement ou extension", "Consulter les offres publiques", "Comparer et simuler", "Utiliser le chatbot public"],
            "associations": [("Administrateur Banque", "main"), ("Visiteur", "Consulter les offres publiques"), ("Visiteur", "Comparer et simuler"), ("Visiteur", "Utiliser le chatbot public")],
            "includes": [("main", "Personnaliser identité et contenu"), ("main", "Configurer stores et modules"), ("main", "Demander renouvellement ou extension")], "extends": [("Comparer et simuler", "Consulter les offres publiques"), ("Utiliser le chatbot public", "Consulter les offres publiques")],
            "primary": "Administrateur Banque", "secondary": "Visiteur de la marketplace", "objective": "Configurer les services souscrits et offrir un parcours public adapté à la banque.",
            "preconditions": "Marketplace active, abonnement valide et Administrateur Banque authentifié pour les fonctions privées.", "trigger": "L’administrateur ouvre le Back-office Banque ou le visiteur accède à la marketplace publique.",
            "nominal": "1. L’administrateur personnalise les contenus et l’identité de la marketplace. 2. Il configure les stores et modules autorisés. 3. Le visiteur consulte les offres publiées. 4. Si les modules sont actifs, il compare les produits, simule un financement ou utilise le chatbot.",
            "alternatives": "Une demande de renouvellement ou d’ajout de services est transmise au SaaS. Un module non souscrit n’est pas disponible sur la marketplace.",
            "postconditions": "La configuration est enregistrée dans le tenant bancaire et les services publics autorisés sont visibles aux visiteurs.",
        },
        {
            "id": "UC-06 — Inscrire et valider un concessionnaire", "sprint": "S6 — Inscription et validation du concessionnaire",
            "title": "Inscription et validation d’un concessionnaire", "caption": "Figure 3.10 — Diagramme de cas d’utilisation détaillé de l’inscription et de la validation d’un concessionnaire",
            "actors_left": ["Représentant Concessionnaire"], "actors_right": ["Administrateur SaaS", "Service e-mail"],
            "main": "Soumettre une demande concessionnaire", "children": ["Saisir les données de l’entreprise", "Téléverser les justificatifs", "Consulter la demande", "Approuver le dossier", "Rejeter avec motif", "Créer le compte DEALER_ADMIN", "Envoyer les identifiants"],
            "associations": [("Représentant Concessionnaire", "main"), ("Administrateur SaaS", "Consulter la demande"), ("Administrateur SaaS", "Approuver le dossier"), ("Service e-mail", "Envoyer les identifiants")],
            "includes": [("main", "Saisir les données de l’entreprise"), ("main", "Téléverser les justificatifs"), ("Approuver le dossier", "Créer le compte DEALER_ADMIN"), ("Créer le compte DEALER_ADMIN", "Envoyer les identifiants")], "extends": [("Rejeter avec motif", "Consulter la demande")],
            "primary": "Représentant du concessionnaire", "secondary": "Administrateur SaaS et service e-mail", "objective": "Obtenir un compte concessionnaire après contrôle du dossier et des pièces justificatives.",
            "preconditions": "Formulaire public accessible et documents requis disponibles.", "trigger": "Le représentant démarre une demande d’inscription concessionnaire.",
            "nominal": "1. Le représentant renseigne les données de l’entreprise. 2. Il joint les justificatifs. 3. Il soumet la demande. 4. L’Administrateur SaaS consulte le dossier. 5. Il l’approuve. 6. Le système crée le compte DEALER_ADMIN et transmet les identifiants par e-mail.",
            "alternatives": "Le SaaS rejette le dossier avec un motif, qui est communiqué au demandeur. Le représentant peut soumettre un nouveau dossier corrigé.",
            "postconditions": "Le concessionnaire possède un compte actif, ou sa demande est rejetée avec traçabilité de la décision.",
        },
        {
            "id": "UC-07 — Gérer un partenariat et un contrat", "sprint": "S7 — Partenariats et contrats",
            "title": "Partenariat banque-concessionnaire et contrat", "caption": "Figure 3.11 — Diagramme de cas d’utilisation détaillé du partenariat et du contrat",
            "actors_left": ["Concessionnaire"], "actors_right": ["Administrateur Banque", "Service e-mail"],
            "main": "Gérer un partenariat", "children": ["Initier une demande", "Examiner la demande", "Accepter le partenariat", "Rejeter le partenariat", "Préparer le contrat", "Envoyer le contrat", "Accepter le contrat", "Résilier le contrat"],
            "associations": [("Concessionnaire", "main"), ("Administrateur Banque", "main"), ("Administrateur Banque", "Examiner la demande"), ("Service e-mail", "Envoyer le contrat")],
            "includes": [("main", "Initier une demande"), ("Examiner la demande", "Accepter le partenariat"), ("Accepter le partenariat", "Préparer le contrat"), ("Préparer le contrat", "Envoyer le contrat"), ("Envoyer le contrat", "Accepter le contrat")], "extends": [("Rejeter le partenariat", "Examiner la demande"), ("Résilier le contrat", "Accepter le contrat")],
            "primary": "Concessionnaire ou Administrateur Banque", "secondary": "Autre partie et service e-mail", "objective": "Établir une relation commerciale active permettant la publication des produits dans la marketplace partenaire.",
            "preconditions": "Concessionnaire approuvé, banque et store compatibles.", "trigger": "La banque ou le concessionnaire initie une demande de partenariat.",
            "nominal": "1. Une partie crée la demande. 2. La partie concernée l’examine. 3. Elle l’accepte. 4. La banque prépare le contrat. 5. Le contrat est envoyé au concessionnaire. 6. Le concessionnaire l’accepte. 7. La relation commerciale devient active.",
            "alternatives": "La demande peut être rejetée avec un motif. Le contrat peut rester en brouillon, être refusé, expirer ou être résilié selon son cycle de vie.",
            "postconditions": "Un partenariat et un contrat actifs autorisent les actions de publication ; sinon, la décision est enregistrée.",
        },
        {
            "id": "UC-08 — Gérer et publier un produit concessionnaire", "sprint": "S8 — Produits, publications et suivi des dossiers",
            "title": "Gestion du catalogue, du stock et des publications", "caption": "Figure 3.12 — Diagramme de cas d’utilisation détaillé de la gestion et de la publication des produits",
            "actors_left": ["Concessionnaire"], "actors_right": ["Administrateur Banque"],
            "main": "Gérer le catalogue produit", "children": ["Créer ou modifier un produit", "Gérer stock et documents", "Soumettre une publication", "Examiner la publication", "Approuver la publication", "Rejeter la publication", "Suivre les dossiers liés"],
            "associations": [("Concessionnaire", "main"), ("Concessionnaire", "Soumettre une publication"), ("Administrateur Banque", "Examiner la publication"), ("Administrateur Banque", "Approuver la publication")],
            "includes": [("main", "Créer ou modifier un produit"), ("main", "Gérer stock et documents"), ("main", "Soumettre une publication"), ("Soumettre une publication", "Examiner la publication"), ("Examiner la publication", "Approuver la publication"), ("main", "Suivre les dossiers liés")], "extends": [("Rejeter la publication", "Examiner la publication")],
            "primary": "Concessionnaire", "secondary": "Administrateur Banque", "objective": "Centraliser le catalogue et rendre un produit visible dans une marketplace partenaire après validation.",
            "preconditions": "Concessionnaire actif, produit renseigné, partenariat et contrat actifs pour la banque ciblée.", "trigger": "Le concessionnaire crée, met à jour ou souhaite publier un produit.",
            "nominal": "1. Le concessionnaire crée ou modifie la fiche produit. 2. Il gère les documents et le stock centralisé. 3. Il soumet une publication pour une banque partenaire. 4. La banque l’examine et l’approuve. 5. Le produit devient visible dans la marketplace correspondante. 6. Le concessionnaire suit les dossiers liés à ses produits.",
            "alternatives": "La banque rejette la publication avec un motif. Le concessionnaire corrige les informations puis soumet à nouveau. Une demande de financement rejetée ne réserve pas de stock.",
            "postconditions": "La publication est approuvée ou rejetée dans le périmètre de la banque ; le produit source et son stock restent centralisés.",
        },
        {
            "id": "UC-09 — Déposer et traiter une demande de financement", "sprint": "S9 — Parcours client et demande de financement",
            "title": "Parcours client et demande de financement", "caption": "Figure 3.13 — Diagramme de cas d’utilisation détaillé de la demande de financement",
            "actors_left": ["Client"], "actors_right": ["Administrateur Banque", "Concessionnaire", "Service e-mail"],
            "main": "Déposer une demande de financement", "children": ["Créer et vérifier le compte", "Consulter, comparer et simuler", "Créer le dossier brouillon", "Téléverser les justificatifs", "Soumettre le dossier", "Instruire la demande", "Accepter et réserver le stock", "Rejeter avec motif", "Suivre la décision", "Notifier la décision"],
            "associations": [("Client", "main"), ("Administrateur Banque", "Instruire la demande"), ("Concessionnaire", "Suivre la décision"), ("Service e-mail", "Notifier la décision")],
            "includes": [("main", "Créer et vérifier le compte"), ("main", "Consulter, comparer et simuler"), ("main", "Créer le dossier brouillon"), ("Créer le dossier brouillon", "Téléverser les justificatifs"), ("Téléverser les justificatifs", "Soumettre le dossier"), ("Soumettre le dossier", "Instruire la demande")], "extends": [("Accepter et réserver le stock", "Instruire la demande"), ("Rejeter avec motif", "Instruire la demande")],
            "primary": "Client", "secondary": "Administrateur Banque, concessionnaire et service e-mail", "objective": "Constituer un dossier complet puis obtenir une décision de la banque concernée.",
            "preconditions": "Marketplace et store actifs ; client inscrit et connecté ; produit éligible.", "trigger": "Le client souhaite poursuivre après la consultation ou la simulation d’un produit.",
            "nominal": "1. Le client crée et vérifie son compte. 2. Il consulte les produits et, si disponibles, utilise le comparateur ou le simulateur. 3. Il crée un dossier brouillon. 4. Il joint les pièces obligatoires. 5. Il soumet le dossier à l’état en attente. 6. La banque consulte le dossier et les documents. 7. Elle accepte ou rejette. 8. Le client est notifié ; un stock est réservé pour un produit concessionnaire accepté.",
            "alternatives": "Le client complète un brouillon incomplet avant soumission. En cas de rejet, la banque saisit obligatoirement le motif. Le concessionnaire consulte uniquement le suivi et ne prend pas la décision bancaire.",
            "postconditions": "Le dossier est accepté ou rejeté avec une décision persistée et communiquée au client.",
        },
        {
            "id": "UC-10 — Utiliser les assistants conversationnels", "sprint": "S10 — Assistants conversationnels et validation",
            "title": "Assistant IA SaaS et chatbot public", "caption": "Figure 3.14 — Diagramme de cas d’utilisation détaillé des assistants conversationnels",
            "actors_left": ["Administrateur SaaS", "Visiteur"], "actors_right": ["Gemini"],
            "main": "Interroger l’assistant SaaS", "children": ["Formuler une question analytique", "Valider une requête en lecture seule", "Restituer une réponse métier", "Envoyer un message au chatbot", "Détecter l’intention", "Consulter les données publiques", "Orienter vers comparateur ou simulateur"],
            "associations": [("Administrateur SaaS", "main"), ("Visiteur", "Envoyer un message au chatbot"), ("Gemini", "Formuler une question analytique")],
            "includes": [("main", "Formuler une question analytique"), ("main", "Valider une requête en lecture seule"), ("main", "Restituer une réponse métier"), ("Envoyer un message au chatbot", "Détecter l’intention"), ("Détecter l’intention", "Consulter les données publiques")], "extends": [("Orienter vers comparateur ou simulateur", "Envoyer un message au chatbot")],
            "primary": "Administrateur SaaS ou visiteur", "secondary": "Gemini pour l’assistant SaaS", "objective": "Faciliter la consultation analytique du SaaS et l’orientation des visiteurs dans une marketplace publique.",
            "preconditions": "Administrateur SaaS authentifié pour l’assistant analytique ; module Chatbot actif pour le chatbot public.", "trigger": "L’utilisateur formule une question ou envoie un message.",
            "nominal": "1. L’Administrateur SaaS pose une question. 2. Le système construit un contexte de données autorisé. 3. Gemini propose une requête. 4. Le validateur autorise uniquement une lecture contrôlée. 5. Le système restitue une réponse métier. Pour le chatbot, le visiteur envoie un message ; le système détecte une intention et répond uniquement à partir des données publiques du tenant.",
            "alternatives": "Une requête analytique non conforme est refusée sans exécution. Le chatbot peut orienter vers le comparateur ou le simulateur si ces modules sont actifs.",
            "postconditions": "Une réponse sécurisée est affichée sans modification de données ni exposition d’informations privées ou inter-tenants.",
        },
    ]


def main():
    specs = make_specs()
    for index, spec in enumerate(specs, start=1):
        spec["path"] = DIAGRAM_DIR / f"uc_sprint_{index:02d}.png"
        build_diagram(spec["path"], spec["title"], spec["actors_left"], spec["actors_right"], spec["main"], spec["children"], spec["associations"], spec["includes"], spec["extends"])

    document = Document(INPUT)
    chapter4_anchor = next(p for p in document.paragraphs if p.text.strip().startswith("Chapitre 4"))
    chapter3_heading = next(p for p in document.paragraphs if p.text.strip().startswith("Chapitre 3"))
    section_template = next(p for p in document.paragraphs if p.text.strip().startswith("2. Environnement"))
    subsection_template = next(p for p in document.paragraphs if p.text.strip().startswith("2.1."))
    body_template = next(p for p in document.paragraphs if p.text.strip().startswith("Pour développer notre solution"))
    caption_template = next(p for p in document.paragraphs if p.style.name == "Caption")

    styled_insert_before(chapter4_anchor, "3. Modélisation UML des besoins fonctionnels", section_template)
    styled_insert_before(chapter4_anchor, "3.1. Cas d’utilisation détaillés alignés sur le Product Backlog", subsection_template)
    styled_insert_before(chapter4_anchor, "Les cas d’utilisation détaillés ci-après sont structurés à partir des dix sprints du Product Backlog et des fonctionnalités réalisées dans les chapitres 4 et 5. Chaque diagramme représente un processus métier cohérent. Les acteurs sont positionnés à l’extérieur de la frontière du système Matchia, tandis que les objectifs sont représentés par des cas d’utilisation. Les relations <<include>> désignent des sous-fonctions obligatoires ; les relations <<extend>> représentent des traitements conditionnels, notamment les décisions de rejet ou les activations dépendantes d’un paiement confirmé.", body_template)

    for table_number, spec in enumerate(specs, start=1):
        styled_insert_before(chapter4_anchor, spec["id"], subsection_template)
        styled_insert_before(chapter4_anchor, f"Ce cas d’utilisation détaille le processus couvert par le sprint {spec['sprint'].split(' — ')[0]}. La figure et le tableau associés précisent les responsabilités des acteurs, les contrôles métier et les résultats attendus.", body_template)
        add_figure(document, chapter4_anchor, spec["path"], spec["caption"], caption_template)
        styled_insert_before(chapter4_anchor, "Le tableau suivant formalise le scénario du cas d’utilisation conformément aux éléments de base d’une description UML.", body_template)
        styled_insert_before(chapter4_anchor, f"Tableau 3.{table_number} — Description du scénario {spec['id'].split(' — ')[0]}", caption_template, WD_ALIGN_PARAGRAPH.CENTER)
        add_scenario_table(document, chapter4_anchor, scenario_rows(spec))
        styled_insert_before(chapter4_anchor, "", body_template)

    styled_insert_before(chapter4_anchor, "Cette modélisation relie ainsi la planification Agile, les besoins fonctionnels et les réalisations présentées dans les chapitres suivants. Elle constitue une référence pour la validation des processus métier et la recette de la plateforme.", body_template)
    document.core_properties.title = "Rapport PFE Matchia avec cas d’utilisation détaillés par sprint"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
