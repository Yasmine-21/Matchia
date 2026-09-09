from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\PFE M2\Platforme SaaS")
SOURCE = ROOT / "document_work" / "Master Report - diagrammes UML complets avec séquences.docx"
OUTPUT = ROOT / "document_work" / "Master Report - architecture et conception complétées.docx"

LOGICAL = ROOT / "document_work" / "architecture_diagrams_v2" / "architecture_logique_matchia.png"
MULTITENANT = ROOT / "document_work" / "full_report_diagrams" / "architecture_multitenant.png"
PHYSICAL = ROOT / "document_work" / "architecture_diagrams_v2" / "architecture_physique_matchia.png"


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def format_body(paragraph, keep_with_next=False):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_with_next = keep_with_next
    for run in paragraph.runs:
        set_run_font(run, size=11)


def add_body(anchor, text, keep_with_next=False):
    p = anchor.insert_paragraph_before(text, style="Normal")
    format_body(p, keep_with_next=keep_with_next)
    return p


def add_heading(anchor, text, level, page_break_before=False):
    p = anchor.insert_paragraph_before(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    return p


def add_figure(anchor, path, caption, alt_text):
    p = anchor.insert_paragraph_before(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(6.15))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", caption)

    cp = anchor.insert_paragraph_before(caption, style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(1)
    cp.paragraph_format.space_after = Pt(8)
    cp.paragraph_format.keep_with_next = False
    for run in cp.runs:
        set_run_font(run, size=9, italic=True, color=(49, 73, 97))
    return p, cp


def replace_paragraph_text(paragraph, new_text):
    paragraph.text = new_text
    if paragraph.style and paragraph.style.name == "Caption":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, size=9, italic=True, color=(49, 73, 97))


for required in (SOURCE, LOGICAL, MULTITENANT, PHYSICAL):
    if not required.exists():
        raise FileNotFoundError(required)

doc = Document(str(SOURCE))

# Localiser précisément le chapitre 3 et la section fonctionnelle existante.
chapter3_start = None
chapter4_start = None
functional_anchor = None
for idx, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text.startswith("Chapitre 3"):
        chapter3_start = idx
    elif chapter3_start is not None and text.startswith("Chapitre 4"):
        chapter4_start = idx
        break

if chapter3_start is None or chapter4_start is None:
    raise RuntimeError("Les limites du chapitre 3 n'ont pas été trouvées.")

chapter3_paragraphs = doc.paragraphs[chapter3_start:chapter4_start]
for p in chapter3_paragraphs:
    if p.text.strip().startswith("3. Modélisation fonctionnelle par cas d’utilisation"):
        functional_anchor = p
        break
if functional_anchor is None:
    raise RuntimeError("La section de modélisation fonctionnelle n'a pas été trouvée.")

# Les trois nouvelles figures d'architecture deviennent les figures 3.1 à 3.3.
# Les figures fonctionnelles existantes sont donc décalées de trois positions.
figure_pattern = re.compile(r"Figure 3\.(\d+)")
for p in chapter3_paragraphs:
    if figure_pattern.search(p.text):
        updated = figure_pattern.sub(lambda m: f"Figure 3.{int(m.group(1)) + 3}", p.text)
        replace_paragraph_text(p, updated)

# La nouvelle partie architecture devient la section 3 ; la modélisation fonctionnelle devient la section 4.
for p in chapter3_paragraphs:
    text = p.text.strip()
    updated = None
    if text.startswith("3. Modélisation fonctionnelle par cas d’utilisation"):
        updated = text.replace("3. Modélisation", "4. Modélisation", 1)
    elif text.startswith("3.1. Diagramme global"):
        updated = text.replace("3.1.", "4.1.", 1)
    elif text.startswith("3.2. Diagrammes et scénarios"):
        updated = text.replace("3.2.", "4.2.", 1)
    elif re.match(r"^3\.2\.\d+\.", text):
        updated = re.sub(r"^3\.2\.(\d+)\.", r"4.2.\1.", text)
    if updated is not None:
        replace_paragraph_text(p, updated)

functional_anchor.paragraph_format.page_break_before = True

# --- Nouvelle section d'architecture ---
add_heading(functional_anchor, "3. Architecture de la solution", 2, page_break_before=True)
add_body(
    functional_anchor,
    "L’architecture de Matchia a été conçue pour mutualiser un socle technique commun tout en "
    "préservant l’autonomie fonctionnelle et l’isolation logique de chaque banque. Elle adopte "
    "une architecture client-serveur en couches, mise en œuvre sous la forme d’un monolithe "
    "modulaire. Ce choix répond aux besoins du projet : centraliser l’administration SaaS, "
    "déployer plusieurs marketplaces bancaires à partir d’une même application et faire évoluer "
    "séparément les principaux domaines métier."
)

add_heading(functional_anchor, "3.1. Vue d’ensemble et choix architectural", 3)
add_body(
    functional_anchor,
    "La plateforme repose sur une application web monopage et une API REST sécurisée. Les espaces "
    "public, Client, Concessionnaire, Administrateur Banque et Administrateur SaaS partagent le même "
    "socle applicatif ; leurs routes, leurs données visibles et leurs actions autorisées varient "
    "cependant selon le rôle et, lorsque cela s’applique, selon la banque de rattachement. Cette "
    "séparation limite le couplage entre l’interface, les règles métier et la persistance."
)
add_body(
    functional_anchor,
    "Le monolithe modulaire a été retenu afin de conserver un déploiement simple et des transactions "
    "cohérentes, sans renoncer à une organisation claire par domaines : authentification, banques et "
    "marketplaces, concessionnaires, catalogue et véhicules, clients, financements, paiements, "
    "notifications et assistants intelligents. Cette modularité facilite les tests, la maintenance "
    "et une éventuelle extraction future de services indépendants."
)

add_heading(functional_anchor, "3.2. Architecture logique", 3)
add_body(
    functional_anchor,
    "L’architecture logique s’organise en cinq niveaux complémentaires. La couche de présentation, "
    "réalisée avec React et TypeScript, gère les vues, les formulaires et la navigation. La couche API "
    "expose les contrôleurs REST et les objets d’échange, puis valide les requêtes avant de les "
    "transmettre aux services métier. Spring Security et le filtre JWT assurent transversalement "
    "l’authentification, l’autorisation par rôle et la prise en compte du contexte multi-tenant."
)
add_body(
    functional_anchor,
    "La couche métier porte les règles propres à chaque domaine fonctionnel. Elle orchestre notamment "
    "la création et le traitement des demandes de marketplace, l’inscription des concessionnaires et "
    "des clients, la gestion du catalogue, les demandes de financement, les paiements et les échanges "
    "avec les assistants. Enfin, la couche de persistance s’appuie sur Spring Data JPA et PostgreSQL ; "
    "les dépôts encapsulent les accès aux données et appliquent les critères de filtrage nécessaires."
)
add_figure(
    functional_anchor,
    LOGICAL,
    "Figure 3.1 — Architecture logique en couches de la plateforme Matchia",
    "Schéma de l’architecture logique en couches de Matchia : présentation React, API REST, sécurité, services métier modulaires et persistance PostgreSQL."
)

add_heading(functional_anchor, "3.3. Architecture multi-tenant et isolation des données", 3)
add_body(
    functional_anchor,
    "Matchia utilise une application et une base de données partagées par plusieurs banques. Chaque "
    "banque constitue un tenant disposant de sa propre marketplace, de ses paramètres, de ses "
    "administrateurs et de ses données métier. Les ressources globales du fournisseur SaaS sont "
    "distinguées des ressources rattachées à une banque, ce qui permet de mutualiser l’infrastructure "
    "sans dupliquer l’application."
)
add_body(
    functional_anchor,
    "Le contexte du tenant est déterminé à partir du compte authentifié pour les espaces privés et à "
    "partir de l’identifiant public de la marketplace, tel que le slug ou l’hôte, pour les parcours "
    "publics. Les entités concernées conservent une référence vers la banque. Les contrôleurs, les "
    "services et les dépôts vérifient cette référence et filtrent systématiquement les recherches par "
    "banque. Ainsi, une ressource appartenant à un autre tenant n’est ni exposée ni modifiable, même "
    "lorsque son identifiant technique est connu."
)
add_figure(
    functional_anchor,
    MULTITENANT,
    "Figure 3.2 — Isolation logique des données dans l’architecture multi-tenant",
    "Schéma montrant le socle Matchia partagé et l’isolation logique des espaces et données propres aux banques A, B et C."
)

add_heading(functional_anchor, "3.4. Sécurité et contrôle d’accès", 3)
add_body(
    functional_anchor,
    "La sécurité repose sur Spring Security et une authentification stateless par jetons JWT. Après "
    "validation des identifiants et de l’état du compte, l’API délivre un jeton d’accès ; le mécanisme "
    "de renouvellement permet de maintenir la session selon les règles définies. Chaque requête "
    "protégée est interceptée afin d’identifier l’utilisateur, son rôle et son contexte d’accès."
)
add_body(
    functional_anchor,
    "Le contrôle d’accès combine les autorisations liées aux rôles — Administrateur SaaS, "
    "Administrateur Banque, Concessionnaire et Client — avec les vérifications liées au tenant. Les "
    "opérations sensibles sont contrôlées dans la couche métier, tandis que les dépôts limitent les "
    "lectures aux données autorisées. Les événements significatifs peuvent être consignés dans les "
    "journaux d’audit. Les secrets techniques et les clés des services externes sont fournis par la "
    "configuration de l’environnement de déploiement."
)

add_heading(functional_anchor, "3.5. Architecture physique et intégrations externes", 3)
add_body(
    functional_anchor,
    "Du point de vue physique, l’utilisateur accède à Matchia depuis un navigateur web. Le frontend "
    "React, construit avec Vite et servi par Nginx, communique en HTTPS avec l’API Spring Boot. Le "
    "backend, exécuté avec Java 17, applique les règles de sécurité et les traitements métier avant "
    "d’accéder à PostgreSQL par l’intermédiaire de JPA. Les documents fonctionnels, notamment ceux "
    "liés aux demandes de financement, sont gérés par le backend et associés aux enregistrements "
    "métier correspondants."
)
add_body(
    functional_anchor,
    "Trois intégrations complètent le système. Stripe prend en charge le paiement associé au "
    "traitement d’une demande et retourne au backend l’état de la transaction. Gemini alimente "
    "l’assistant SaaS et le chatbot de marketplace, dans le respect du contexte transmis par "
    "l’application. Enfin, le serveur SMTP assure l’envoi des codes de vérification et des "
    "notifications. Les composants déployés, leur hébergement et le stockage persistant sont détaillés "
    "dans le chapitre consacré au déploiement."
)
add_figure(
    functional_anchor,
    PHYSICAL,
    "Figure 3.3 — Architecture physique et intégrations externes de Matchia",
    "Schéma de déploiement : navigateur, frontend React servi par Nginx, backend Spring Boot, PostgreSQL et services externes Stripe, Gemini et SMTP."
)

add_heading(functional_anchor, "3.6. Synthèse des choix d’architecture", 3)
add_body(
    functional_anchor,
    "Cette architecture concilie simplicité de déploiement, séparation des responsabilités et "
    "évolutivité. Le monolithe modulaire facilite la maintenance, tandis que la mutualisation réduit "
    "les coûts. Chaque accès aux données tenantisées demeure néanmoins associé au contexte de la "
    "banque et soumis aux contrôles d’autorisation."
)

doc.save(str(OUTPUT))

# Contrôles structurels minimaux avant le rendu visuel.
check = Document(str(OUTPUT))
texts = [p.text.strip() for p in check.paragraphs]
assert "3. Architecture de la solution" in texts
assert "4. Modélisation fonctionnelle par cas d’utilisation" in texts
assert any(t.startswith("Figure 3.21") for t in texts), "La dernière figure du chapitre 3 n'est pas renumérotée."
assert len(check.inline_shapes) == len(Document(str(SOURCE)).inline_shapes) + 3
assert len(check.tables) == len(Document(str(SOURCE)).tables)

print(f"Rapport créé : {OUTPUT}")
print(f"Paragraphes : {len(check.paragraphs)}")
print(f"Tableaux conservés : {len(check.tables)}")
print(f"Illustrations : {len(check.inline_shapes)}")
