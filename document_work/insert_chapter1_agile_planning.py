from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches


INPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - chapitres 4 et 5 réorganisés.docx")
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - rapport mis à jour avec planification Agile.docx")


def first_formatted_run(paragraph):
    return next((run for run in paragraph.runs if run.text), None)


def clone_paragraph_format(target, template):
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if template._p.pPr is not None:
        target._p.insert(0, deepcopy(template._p.pPr))


def insert_paragraph_before(anchor, text, template):
    paragraph = anchor.insert_paragraph_before()
    clone_paragraph_format(paragraph, template)
    run = paragraph.add_run(text)
    source_run = first_formatted_run(template)
    if source_run is not None and source_run._r.rPr is not None:
        run._r.insert(0, deepcopy(source_run._r.rPr))
    return paragraph


def set_cell_text(cell, text, bold=False):
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    run = paragraph.add_run(text)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def insert_backlog_table_before(document, anchor, rows):
    table = document.add_table(rows=1, cols=2)
    table.style = "Grid Table 1 Light Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(1.2), Inches(5.2))
    for index, header in enumerate(("Sprint", "Éléments du Product Backlog")):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        set_cell_text(cell, header, bold=True)
    for sprint, backlog in rows:
        cells = table.add_row().cells
        for index, value in enumerate((sprint, backlog)):
            cells[index].width = widths[index]
            set_cell_text(cells[index], value)
    anchor._p.addprevious(table._tbl)


def main():
    document = Document(INPUT)
    conclusion = next(p for p in document.paragraphs if p.text.strip() == "Conclusion")
    section_template = next(p for p in document.paragraphs if p.text.strip().startswith("Méthodologie de Gestion de Projet"))
    subsection_template = next(p for p in document.paragraphs if p.text.strip().startswith("Le cycle de vie de Scrum"))
    body_template = next(p for p in document.paragraphs if p.text.strip().startswith("Dans le cadre de notre projet"))

    insert_paragraph_before(conclusion, "Planification Agile du projet", section_template)
    insert_paragraph_before(
        conclusion,
        "Afin d’assurer une réalisation progressive et maîtrisée de la plateforme, le projet a été planifié selon l’approche Agile Scrum. Le Product Backlog regroupe l’ensemble des besoins fonctionnels identifiés et est priorisé selon leur valeur métier et leurs dépendances techniques. Chaque sprint, prévu sur une durée de deux semaines, permet de produire, tester et valider un incrément fonctionnel de la plateforme.",
        body_template,
    )
    insert_paragraph_before(
        conclusion,
        "La planification retenue est organisée en deux releases correspondant aux deux grandes parties de réalisation du projet. La première release couvre le socle SaaS multi-tenant et le cycle de vie d’une marketplace bancaire. La seconde se concentre sur les parcours métier des concessionnaires et des clients, ainsi que sur les fonctionnalités conversationnelles.",
        body_template,
    )

    insert_paragraph_before(conclusion, "Release 1 — SaaS bancaire, onboarding et activation des marketplaces", subsection_template)
    insert_paragraph_before(
        conclusion,
        "Cette release vise à rendre opérationnel le parcours complet d’une banque, depuis sa demande de création de marketplace jusqu’à l’activation de son espace et de ses services après paiement.",
        body_template,
    )
    insert_backlog_table_before(
        document,
        conclusion,
        [
            ("Sprint 1", "Mise en place du socle applicatif : authentification, gestion des rôles, autorisation, gestion des sessions et isolation des données dans un contexte multi-tenant."),
            ("Sprint 2", "Développement de la demande de création d’une marketplace : informations de la banque et de l’Administrateur Banque, vérification de l’adresse électronique, configuration initiale, sélection des stores et modules, récapitulatif et soumission."),
            ("Sprint 3", "Réalisation du Back-office SaaS : tableau de bord, gestion des banques, demandes, stores et modules, utilisateurs et rôles, concessionnaires, marketplaces, contenus, audit et paramètres."),
            ("Sprint 4", "Mise en place du traitement des demandes : approbation ou rejet, notifications, e-mails, intégration du paiement Stripe, vérification de la transaction et activation de la banque, de la marketplace, de l’abonnement et du compte Administrateur Banque."),
            ("Sprint 5", "Réalisation du Back-office Banque et de la marketplace publique : personnalisation, contenus, stores et modules, abonnement, produits, comparateur et simulateur."),
        ],
    )

    insert_paragraph_before(conclusion, "Release 2 — Parcours métier des concessionnaires et des clients", subsection_template)
    insert_paragraph_before(
        conclusion,
        "Cette release couvre les fonctionnalités permettant aux concessionnaires de gérer leur activité multi-banque et aux clients de consulter les offres puis de déposer une demande de financement.",
        body_template,
    )
    insert_backlog_table_before(
        document,
        conclusion,
        [
            ("Sprint 6", "Développement du parcours d’inscription du concessionnaire : saisie des informations, dépôt des documents justificatifs, traitement de la demande par le SaaS, création du compte et envoi des identifiants."),
            ("Sprint 7", "Mise en œuvre des partenariats et de la gestion des produits : demande initiée par le concessionnaire ou la banque, traitement de la demande, contrats, catalogue centralisé, stock et publications multi-banques."),
            ("Sprint 8", "Réalisation du parcours client et de la demande de financement : consultation des produits, comparaison, simulation, inscription et vérification du compte, constitution du dossier, dépôt des documents et transmission à la banque."),
            ("Sprint 9", "Intégration des fonctionnalités conversationnelles et finalisation : assistant IA du Back-office SaaS, chatbot public, contrôle des accès aux données, tests fonctionnels, corrections et validation globale des parcours."),
        ],
    )
    insert_paragraph_before(
        conclusion,
        "À l’issue de chaque sprint, les fonctionnalités développées font l’objet de tests techniques et fonctionnels. Elles sont ensuite présentées lors de la revue de sprint afin de recueillir les retours nécessaires et d’ajuster, si besoin, les priorités du Product Backlog. Cette organisation favorise une progression continue du projet et une livraison progressive de fonctionnalités utilisables.",
        body_template,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
