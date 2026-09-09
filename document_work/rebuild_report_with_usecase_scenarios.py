from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import rebuild_report_with_usecase_diagrams as base


SOURCE = base.SOURCE
WORKSPACE = base.WORKSPACE
OUTPUT = (
    WORKSPACE
    / "document_work"
    / "Master Report - diagrammes et tableaux de scenarios use case.docx"
)


SCENARIOS = [
    {
        "code": "UC-01",
        "title": "S’authentifier",
        "primary": "Utilisateur : Administrateur SaaS, Administrateur Banque, Concessionnaire ou Client.",
        "secondary": "Aucun acteur secondaire direct.",
        "objective": "Permettre à un utilisateur d’accéder à l’espace correspondant à son rôle, à ses autorisations et, si nécessaire, au tenant auquel il est rattaché.",
        "pre": "L’utilisateur possède un compte enregistré et accède au formulaire de connexion. Le service d’authentification est disponible.",
        "trigger": "L’utilisateur saisit ses identifiants et soumet le formulaire de connexion.",
        "success": "Une session authentifiée est créée. Le rôle, les autorisations et, le cas échéant, le contexte du tenant sont associés à la session. L’utilisateur est redirigé vers l’espace approprié.",
        "failure": "Aucune session valide n’est créée ou conservée. L’accès à la ressource demandée est refusé et un message adapté est affiché.",
        "steps": [
            "L’utilisateur accède au formulaire de connexion.",
            "Il saisit son adresse e-mail et son mot de passe.",
            "Il soumet le formulaire de connexion.",
            "Le système vérifie les identifiants fournis et l’état du compte.",
            "Le système récupère le rôle et contrôle les autorisations.",
            "Si l’utilisateur est rattaché à une banque ou à une marketplace, le système détermine le contexte du tenant.",
            "Le système crée la session authentifiée.",
            "L’utilisateur est redirigé vers l’espace correspondant à son rôle.",
        ],
        "alts": [
            "Identifiants incorrects : l’authentification est refusée et un message d’erreur est affiché.",
            "Compte inactif : l’accès est refusé.",
            "Contexte du tenant invalide ou accès hors périmètre : la ressource est refusée.",
            "Session ou jeton expiré : l’utilisateur doit renouveler la session ou se reconnecter.",
        ],
    },
    {
        "code": "UC-02",
        "title": "Soumettre une demande de marketplace bancaire",
        "primary": "Internaute agissant comme représentant de la banque ou futur Administrateur Banque.",
        "secondary": "Service d’e-mail.",
        "objective": "Enregistrer une demande complète de création de marketplace bancaire afin qu’elle puisse être examinée par l’Administrateur SaaS.",
        "pre": "Le formulaire public est accessible. Le demandeur dispose d’une adresse e-mail valide et les stores, modules et offres proposés sont disponibles.",
        "trigger": "L’Internaute démarre une nouvelle demande de marketplace.",
        "success": "La demande est enregistrée avec son récapitulatif et son statut initial. L’Administrateur SaaS est notifié et le demandeur reçoit une confirmation.",
        "failure": "Aucune demande incomplète ou non vérifiée n’est transmise pour traitement.",
        "steps": [
            "L’Internaute renseigne les informations de la banque et du futur Administrateur Banque.",
            "Le système envoie un code de vérification à l’adresse e-mail indiquée.",
            "L’Internaute saisit le code reçu et le système valide l’adresse.",
            "L’Internaute configure le slug, l’identité visuelle, les stores, les modules et l’offre souhaités.",
            "Le système affiche un récapitulatif et contrôle la complétude des données.",
            "L’Internaute confirme et soumet la demande.",
            "Le système enregistre la demande, crée la notification et envoie l’e-mail de confirmation.",
        ],
        "alts": [
            "Données invalides ou incomplètes : le système signale les champs à corriger.",
            "Code incorrect ou expiré : la vérification est refusée et un nouveau code peut être demandé.",
            "Slug indisponible : le demandeur doit choisir une autre valeur.",
            "Abandon avant soumission : aucune demande définitive n’est créée.",
        ],
    },
    {
        "code": "UC-03",
        "title": "Administrer et superviser la plateforme SaaS",
        "primary": "Administrateur SaaS.",
        "secondary": "Aucun acteur secondaire direct.",
        "objective": "Piloter l’écosystème Matchia et administrer les ressources globales, les tenants, les offres, les abonnements et les paramètres de la plateforme.",
        "pre": "L’Administrateur SaaS possède un compte actif, est authentifié et dispose des autorisations d’administration globale.",
        "trigger": "L’Administrateur SaaS ouvre le back-office SaaS ou sélectionne une fonction d’administration.",
        "success": "Les consultations et modifications autorisées sont exécutées, persistées dans le périmètre approprié et tracées dans l’audit.",
        "failure": "Aucune modification non autorisée ou incohérente n’est enregistrée.",
        "steps": [
            "L’Administrateur SaaS s’authentifie et accède au tableau de bord global.",
            "Il choisit le domaine à administrer : stores, modules, banques, utilisateurs, rôles, concessionnaires, marketplaces, contenus, offres, abonnements, paiements ou paramètres.",
            "Le système affiche les données et indicateurs correspondants.",
            "L’administrateur consulte, crée ou met à jour les ressources autorisées.",
            "Le système valide les données et les transitions de statut.",
            "Le système enregistre l’action et ajoute une trace d’audit.",
        ],
        "alts": [
            "Données invalides : l’enregistrement est refusé avec indication des erreurs.",
            "Transition de statut interdite ou ressource introuvable : l’opération est annulée.",
            "Droit insuffisant : l’accès à la fonction est refusé.",
        ],
    },
    {
        "code": "UC-04",
        "title": "Traiter une demande de marketplace",
        "primary": "Administrateur SaaS.",
        "secondary": "Demandeur bancaire et service d’e-mail.",
        "objective": "Examiner une demande de marketplace, enregistrer une décision motivée et déclencher le paiement lorsqu’elle est approuvée.",
        "pre": "L’Administrateur SaaS est authentifié. Une demande complète existe au statut en attente de traitement.",
        "trigger": "L’Administrateur SaaS ouvre une demande depuis la liste des demandes reçues.",
        "success": "La décision est enregistrée. Une approbation déclenche le parcours de paiement ; un rejet conserve le motif et le communique au demandeur.",
        "failure": "La demande reste dans son état précédent et aucun paiement ni service n’est activé.",
        "steps": [
            "L’administrateur consulte la liste des demandes.",
            "Il ouvre le détail de la demande sélectionnée.",
            "Le système affiche les informations de la banque, du futur administrateur et de la configuration souhaitée.",
            "L’administrateur vérifie la cohérence et la complétude du dossier.",
            "Il approuve la demande ou la rejette en indiquant un motif.",
            "Le système enregistre la décision.",
            "En cas d’approbation, le système déclenche le paiement ; en cas de rejet, il notifie le demandeur.",
        ],
        "alts": [
            "Demande approuvée : un lien de paiement est généré et transmis.",
            "Demande rejetée : le motif est obligatoire et envoyé au demandeur.",
            "Dossier incomplet ou statut incompatible : la décision est bloquée.",
            "Service d’e-mail indisponible : la décision reste enregistrée et l’envoi est signalé en erreur.",
        ],
    },
    {
        "code": "UC-05",
        "title": "Payer et activer la marketplace",
        "primary": "Internaute représentant la banque et destinataire du lien de paiement.",
        "secondary": "Stripe et service d’e-mail.",
        "objective": "Finaliser le paiement d’une demande approuvée puis activer la banque, la marketplace, l’abonnement et le compte Administrateur Banque.",
        "pre": "La demande de marketplace est approuvée, le lien de paiement est valide et les ressources à activer sont encore inactives.",
        "trigger": "Le représentant de la banque ouvre le lien de paiement reçu.",
        "success": "Le paiement est confirmé. Les ressources souscrites et le compte Administrateur Banque sont activés, puis les informations de première connexion sont envoyées.",
        "failure": "Les ressources restent inactives et aucun accès privé à la marketplace n’est accordé.",
        "steps": [
            "L’Internaute ouvre le lien de paiement reçu par e-mail.",
            "Le système crée une session de paiement Stripe.",
            "L’Internaute saisit les informations requises et confirme la transaction.",
            "Stripe traite le paiement et renvoie son statut.",
            "Le backend vérifie la confirmation Stripe et met à jour le paiement.",
            "Le système active la banque, la marketplace, l’abonnement, les stores, les modules et le compte Administrateur Banque.",
            "Le service d’e-mail transmet les informations de première connexion.",
        ],
        "alts": [
            "Paiement en attente, annulé ou échoué : les ressources restent inactives.",
            "Transaction échouée : le demandeur peut relancer le paiement.",
            "Confirmation Stripe absente ou incohérente : l’activation est bloquée.",
            "Nouvelle notification de paiement déjà traitée : le système évite une seconde activation.",
        ],
    },
    {
        "code": "UC-06",
        "title": "Administrer la marketplace bancaire",
        "primary": "Administrateur Banque.",
        "secondary": "Aucun acteur secondaire direct.",
        "objective": "Configurer et exploiter la marketplace de la banque exclusivement dans le tenant auquel l’administrateur est rattaché.",
        "pre": "Le compte Administrateur Banque, la banque, la marketplace et l’abonnement sont actifs. L’utilisateur est authentifié dans le bon tenant.",
        "trigger": "L’Administrateur Banque ouvre son espace d’administration ou sélectionne une fonction de gestion.",
        "success": "Les informations et décisions autorisées sont enregistrées dans le tenant de la banque et deviennent visibles dans les espaces concernés.",
        "failure": "Aucune ressource d’un autre tenant ou aucun service non souscrit n’est modifié.",
        "steps": [
            "L’Administrateur Banque s’authentifie et accède à son tableau de bord.",
            "Il consulte son profil, les notifications et les indicateurs de sa marketplace.",
            "Il personnalise l’identité visuelle et les contenus.",
            "Il configure les stores, modules et paramètres autorisés par l’abonnement.",
            "Il gère les produits, clients, concessionnaires, partenariats et contrats de son tenant.",
            "Il traite les demandes de financement et consulte l’état de l’abonnement.",
            "Le système enregistre et trace les actions réalisées.",
        ],
        "alts": [
            "Accès à une ressource d’un autre tenant : l’opération est refusée.",
            "Module ou service non actif : la fonction n’est pas accessible.",
            "Donnée invalide ou transition interdite : l’enregistrement est annulé.",
        ],
    },
    {
        "code": "UC-07",
        "title": "Renouveler ou ajouter des services",
        "primary": "Administrateur Banque.",
        "secondary": "Administrateur SaaS, Stripe et service d’e-mail.",
        "objective": "Renouveler l’abonnement d’une banque ou ajouter des stores et modules, après approbation et paiement confirmé.",
        "pre": "L’Administrateur Banque est authentifié et un abonnement existe pour sa marketplace.",
        "trigger": "L’Administrateur Banque sélectionne l’action de renouvellement ou d’extension de services.",
        "success": "L’abonnement est mis à jour et les services approuvés et payés sont activés pour la marketplace.",
        "failure": "L’abonnement et les services existants restent inchangés.",
        "steps": [
            "L’Administrateur Banque consulte l’abonnement et son historique.",
            "Il choisit un renouvellement ou une extension.",
            "Il sélectionne les stores et modules souhaités, puis soumet la demande.",
            "L’Administrateur SaaS examine et approuve la demande.",
            "Le système transmet le lien de paiement.",
            "L’Administrateur Banque effectue le paiement via Stripe.",
            "Le backend confirme le paiement, met à jour l’abonnement et active les services.",
            "Les acteurs concernés reçoivent une notification.",
        ],
        "alts": [
            "Demande non approuvée : l’Administrateur SaaS enregistre un motif de rejet.",
            "Paiement échoué ou en attente : les services restent inchangés.",
            "Store ou module indisponible : la sélection doit être corrigée.",
            "Demande annulée avant paiement : aucune modification n’est appliquée.",
        ],
    },
    {
        "code": "UC-08",
        "title": "Consulter la marketplace, comparer et simuler",
        "primary": "Internaute ou Client.",
        "secondary": "Aucun acteur secondaire direct.",
        "objective": "Permettre la consultation publique du catalogue d’une banque et proposer la comparaison ou la simulation lorsque les modules correspondants sont actifs.",
        "pre": "Le slug identifie une marketplace active disposant d’au moins un store accessible publiquement.",
        "trigger": "L’utilisateur accède à l’adresse publique de la marketplace ou ouvre une fiche produit.",
        "success": "Les contenus publics, les produits et, le cas échéant, les résultats de comparaison ou de simulation sont affichés sans exposer de données privées.",
        "failure": "Aucune donnée d’un tenant inexistant, inactif ou non autorisé n’est affichée.",
        "steps": [
            "Le système identifie la marketplace à partir de son slug.",
            "Il affiche l’identité visuelle et les contenus publics de la banque.",
            "L’utilisateur consulte les stores disponibles et parcourt les produits publiés.",
            "Il ouvre la fiche détaillée d’un produit.",
            "Si le module Comparateur est actif, il sélectionne plusieurs produits et les compare.",
            "Si le module Simulateur est actif, il saisit les paramètres et consulte le résultat de simulation.",
            "S’il souhaite poursuivre vers un financement, le système l’oriente vers l’inscription ou l’espace Client.",
        ],
        "alts": [
            "Module Comparateur ou Simulateur inactif : l’option correspondante n’est pas proposée.",
            "Produit retiré, stock indisponible ou publication suspendue : la fiche n’est plus accessible à la sélection.",
            "Slug inconnu ou marketplace inactive : une page d’indisponibilité est affichée.",
        ],
    },
    {
        "code": "UC-09",
        "title": "Déposer une demande d’inscription comme Concessionnaire",
        "primary": "Internaute représentant le concessionnaire.",
        "secondary": "Administrateur SaaS.",
        "objective": "Transmettre un dossier d’inscription complet afin d’obtenir ultérieurement un compte Concessionnaire après validation.",
        "pre": "Le formulaire public est accessible et le représentant dispose des informations et pièces justificatives requises.",
        "trigger": "L’Internaute démarre une demande d’inscription Concessionnaire.",
        "success": "La demande et ses documents sont enregistrés au statut en attente et l’Administrateur SaaS est notifié.",
        "failure": "Aucune demande incomplète n’est transmise pour validation.",
        "steps": [
            "L’Internaute renseigne l’entreprise et ses coordonnées.",
            "Il sélectionne le store correspondant à son activité.",
            "Il complète les informations d’activité demandées.",
            "Il téléverse les pièces justificatives obligatoires.",
            "Le système vérifie la complétude du dossier.",
            "L’Internaute confirme et transmet la demande.",
            "Le système enregistre le dossier et notifie l’Administrateur SaaS.",
        ],
        "alts": [
            "Informations ou pièces manquantes : le dossier doit être corrigé ou complété.",
            "Entreprise déjà enregistrée : le système refuse la duplication.",
            "Échec du téléversement : le document concerné doit être chargé de nouveau.",
            "Abandon avant soumission : aucun dossier définitif n’est transmis.",
        ],
    },
    {
        "code": "UC-10",
        "title": "Traiter une demande d’inscription Concessionnaire",
        "primary": "Administrateur SaaS.",
        "secondary": "Demandeur concessionnaire et service d’e-mail.",
        "objective": "Contrôler l’identité, la conformité et l’éligibilité du demandeur, puis approuver son inscription ou la rejeter avec un motif.",
        "pre": "L’Administrateur SaaS est authentifié et une demande Concessionnaire existe au statut en attente.",
        "trigger": "L’Administrateur SaaS ouvre le détail d’une demande d’inscription.",
        "success": "Le dossier est approuvé et un compte Concessionnaire actif est créé, ou la demande est rejetée avec une décision tracée.",
        "failure": "La demande conserve son statut précédent et aucun compte n’est créé.",
        "steps": [
            "L’administrateur consulte les demandes reçues.",
            "Il ouvre les informations et les documents du dossier sélectionné.",
            "Il vérifie l’identité, la conformité des pièces et l’éligibilité de l’entreprise.",
            "Il enregistre sa décision.",
            "Si le dossier est admissible, le système crée le compte Concessionnaire.",
            "Le système transmet la décision et, en cas d’approbation, les identifiants de première connexion.",
        ],
        "alts": [
            "Dossier non admissible : l’administrateur saisit un motif de rejet.",
            "Pièce illisible ou information insuffisante : la décision est différée jusqu’à correction.",
            "Compte ou entreprise déjà existant : la création du compte est bloquée.",
            "Service d’e-mail indisponible : la décision reste enregistrée et l’échec d’envoi est signalé.",
        ],
    },
    {
        "code": "UC-11",
        "title": "Gérer les partenariats et les contrats",
        "primary": "Concessionnaire.",
        "secondary": "Administrateur Banque.",
        "objective": "Établir et administrer une relation commerciale permettant au concessionnaire de publier ses produits dans la marketplace d’une banque.",
        "pre": "Les deux acteurs sont authentifiés. Le concessionnaire est approuvé, la banque est active et le store sélectionné est compatible.",
        "trigger": "Le Concessionnaire initie une demande de partenariat ou l’Administrateur Banque ouvre une demande à traiter.",
        "success": "Le partenariat et le contrat sont activés, ou une décision de rejet, d’annulation, de résiliation ou d’expiration est enregistrée.",
        "failure": "Aucune publication n’est autorisée en l’absence d’un partenariat et d’un contrat valides.",
        "steps": [
            "Le Concessionnaire s’authentifie et initie une demande de partenariat.",
            "Il sélectionne la banque et le store concernés.",
            "L’Administrateur Banque consulte et traite la demande.",
            "La banque accepte la demande.",
            "Les conditions du contrat sont définies et le contrat est transmis au Concessionnaire.",
            "Le Concessionnaire accepte le contrat.",
            "Le système active le partenariat et le contrat.",
        ],
        "alts": [
            "Demande rejetée : le motif est enregistré et la relation reste inactive.",
            "Contrat refusé ou annulé : le partenariat ne permet pas la publication.",
            "Résiliation autorisée : le contrat et les droits associés sont désactivés.",
            "Date de fin atteinte : le contrat expire automatiquement.",
        ],
    },
    {
        "code": "UC-12",
        "title": "Gérer les produits, les stocks et les publications",
        "primary": "Concessionnaire.",
        "secondary": "Administrateur Banque.",
        "objective": "Centraliser le catalogue et le stock du concessionnaire, puis publier les produits dans plusieurs marketplaces après validation de chaque banque.",
        "pre": "Le Concessionnaire est authentifié et actif. Pour chaque banque ciblée, un partenariat et un contrat valides existent.",
        "trigger": "Le Concessionnaire crée ou modifie un produit, met à jour son stock ou demande une publication.",
        "success": "Le produit et le stock sont mis à jour. Chaque publication approuvée devient visible dans la marketplace autorisée.",
        "failure": "Une publication non conforme reste invisible et aucune donnée d’une banque non autorisée n’est modifiée.",
        "steps": [
            "Le Concessionnaire s’authentifie et crée ou modifie la fiche produit.",
            "Il renseigne le stock centralisé et les documents nécessaires.",
            "Il sélectionne les banques partenaires visées.",
            "Il soumet une demande de publication pour chaque banque.",
            "L’Administrateur Banque examine et valide la publication.",
            "Le système publie le produit dans les marketplaces autorisées.",
            "Les mises à jour ultérieures du produit et du stock sont propagées aux publications actives.",
        ],
        "alts": [
            "Publication refusée : le motif est enregistré et le produit reste non visible dans la banque concernée.",
            "Partenariat, contrat, store ou produit invalide : la soumission est bloquée.",
            "Demande de financement acceptée pour un produit concessionnaire : le stock correspondant est réservé.",
            "Stock insuffisant : la réservation ou la disponibilité publique est ajustée.",
        ],
    },
    {
        "code": "UC-13",
        "title": "S’inscrire comme Client",
        "primary": "Internaute.",
        "secondary": "Service d’e-mail.",
        "objective": "Créer et vérifier un compte Client rattaché à la marketplace bancaire depuis laquelle l’inscription est lancée.",
        "pre": "La marketplace est active, l’inscription Client est disponible et l’adresse e-mail fournie n’est pas déjà utilisée dans le contexte concerné.",
        "trigger": "L’Internaute sélectionne l’inscription Client ou poursuit un parcours de financement nécessitant un compte.",
        "success": "Le compte Client est vérifié, activé et rattaché au tenant de la marketplace. Le parcours initial peut être repris.",
        "failure": "Aucun compte actif n’est créé tant que l’adresse e-mail n’est pas vérifiée.",
        "steps": [
            "Le système identifie le tenant à partir de la marketplace courante.",
            "L’Internaute renseigne ses informations personnelles et son adresse e-mail.",
            "Le système valide le formulaire et envoie un code de vérification.",
            "L’Internaute saisit le code reçu.",
            "Le système vérifie le code et active le compte Client.",
            "Le Client est redirigé vers son espace ou vers la reprise du parcours de financement.",
        ],
        "alts": [
            "Code incorrect ou expiré : la vérification est refusée.",
            "Code non reçu ou expiré : un nouveau code peut être envoyé.",
            "Adresse e-mail déjà utilisée : le système propose la connexion plutôt qu’une nouvelle inscription.",
            "Tenant ou marketplace inactif : l’inscription est indisponible.",
        ],
    },
    {
        "code": "UC-14",
        "title": "Constituer et soumettre une demande de financement",
        "primary": "Client.",
        "secondary": "Administrateur Banque, destinataire du dossier.",
        "objective": "Permettre au Client de préparer un dossier complet, de joindre les pièces requises et de le soumettre à la banque concernée.",
        "pre": "Le Client est authentifié. La marketplace, le store et le produit sélectionné sont actifs et les exigences documentaires sont disponibles.",
        "trigger": "Le Client choisit de poursuivre un produit ou une simulation par une demande de financement.",
        "success": "Le dossier est soumis à la banque avec le statut en attente et reste consultable par le Client pour le suivi.",
        "failure": "Le dossier reste en brouillon et aucune instruction bancaire ne commence.",
        "steps": [
            "Le Client s’authentifie et sélectionne un produit ou une simulation.",
            "Le système crée un dossier à l’état brouillon.",
            "Le Client renseigne les informations financières et personnelles requises.",
            "Il téléverse les documents justificatifs demandés.",
            "Le système vérifie la présence des pièces obligatoires.",
            "Le Client confirme et soumet le dossier à la banque.",
            "Le système place le dossier au statut en attente et le rend accessible au suivi.",
        ],
        "alts": [
            "Dossier non finalisé : il est conservé en brouillon.",
            "Pièces obligatoires manquantes : la soumission est bloquée.",
            "Produit retiré ou store inactif avant soumission : le parcours ne peut pas être finalisé.",
            "Le Client peut corriger les informations tant que le dossier n’est pas soumis.",
        ],
    },
    {
        "code": "UC-15",
        "title": "Traiter une demande de financement",
        "primary": "Administrateur Banque.",
        "secondary": "Client, Concessionnaire et service d’e-mail.",
        "objective": "Instruire une demande relevant du tenant bancaire, enregistrer la décision et informer les acteurs concernés.",
        "pre": "L’Administrateur Banque est authentifié dans son tenant et une demande existe au statut en attente.",
        "trigger": "L’Administrateur Banque ouvre une demande de financement depuis la liste de son tenant.",
        "success": "Le dossier est accepté ou rejeté, la décision est persistée et le Client est notifié. Le stock est réservé lorsque les conditions sont réunies.",
        "failure": "Le statut du dossier et le stock restent inchangés.",
        "steps": [
            "L’Administrateur Banque consulte les demandes de son tenant.",
            "Il ouvre le dossier et les documents associés.",
            "Il vérifie la complétude et les informations fournies.",
            "Il accepte ou rejette la demande.",
            "Le système enregistre la décision et met à jour le statut.",
            "Le service d’e-mail notifie le Client.",
            "Si la demande acceptée concerne un produit concessionnaire, le système réserve le stock et informe le Concessionnaire.",
        ],
        "alts": [
            "Demande rejetée : la saisie du motif est obligatoire avant validation.",
            "Produit concessionnaire et stock insuffisant : la réservation échoue et l’acceptation doit être réévaluée.",
            "Dossier appartenant à un autre tenant : l’accès est refusé.",
            "Statut déjà modifié : le système empêche une décision concurrente.",
        ],
    },
    {
        "code": "UC-16",
        "title": "Interroger l’assistant IA du SaaS",
        "primary": "Administrateur SaaS.",
        "secondary": "Gemini.",
        "objective": "Obtenir une réponse analytique à partir d’une question en langage naturel, au moyen d’une requête SQL validée et exécutée uniquement en lecture.",
        "pre": "L’Administrateur SaaS est authentifié, l’assistant est disponible et le schéma de données autorisé peut être construit.",
        "trigger": "L’Administrateur SaaS formule une question dans l’interface de l’assistant.",
        "success": "Une réponse métier est affichée à partir d’une requête autorisée, sans modification des données ni exposition hors périmètre.",
        "failure": "Aucune requête non autorisée n’est exécutée et aucune donnée protégée n’est divulguée.",
        "steps": [
            "L’Administrateur SaaS s’authentifie et formule une question en langage naturel.",
            "Le système construit le schéma limité aux données autorisées.",
            "Il transmet la question et le schéma à Gemini.",
            "Gemini génère une proposition de requête SQL.",
            "Le système valide la requête de manière déterministe.",
            "La requête autorisée est exécutée en lecture seule avec une durée limitée.",
            "Le système reformule le résultat et affiche la réponse.",
        ],
        "alts": [
            "Requête non autorisée, non conforme ou hors périmètre : elle est refusée avant exécution.",
            "Temps maximal dépassé : l’exécution est interrompue.",
            "Gemini indisponible ou réponse inexploitable : un message d’erreur est affiché.",
            "Aucun résultat : le système l’indique sans inventer de données.",
        ],
    },
    {
        "code": "UC-17",
        "title": "Utiliser le chatbot public",
        "primary": "Internaute ou Client.",
        "secondary": "Aucun acteur secondaire direct.",
        "objective": "Orienter l’utilisateur dans une marketplace bancaire à partir des seules données publiques du tenant et des modules disponibles.",
        "pre": "La marketplace et le store sont identifiables et le module Chatbot est actif.",
        "trigger": "L’utilisateur envoie un message dans l’interface du chatbot public.",
        "success": "Une réponse publique et contextualisée est affichée, ou l’utilisateur est orienté vers les produits, le comparateur ou le simulateur.",
        "failure": "Aucune donnée privée, non autorisée ou appartenant à un autre tenant n’est révélée.",
        "steps": [
            "Le système identifie la marketplace et le store courants.",
            "Il vérifie que le module Chatbot est actif.",
            "L’utilisateur saisit son message.",
            "Le système normalise le message et détecte l’intention métier.",
            "Il recherche uniquement dans les données publiques du tenant.",
            "Il affiche la réponse ou oriente l’utilisateur vers un produit, le comparateur ou le simulateur.",
        ],
        "alts": [
            "Module Chatbot désactivé : l’interface est rendue indisponible.",
            "Demande portant sur une donnée privée ou une intention interdite : le chatbot refuse la demande.",
            "Intention non reconnue : le chatbot demande une reformulation.",
            "Aucun résultat public pertinent : le chatbot l’indique et propose une orientation générale.",
        ],
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size=9.0, bold=False, color="000000") -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")


def fill_cell(cell, value, *, bold=False, size=9.0, color="000000", numbered=False) -> None:
    cell.text = ""
    values = value if isinstance(value, list) else [value]
    for index, item in enumerate(values, start=1):
        p = cell.paragraphs[0] if index == 1 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if not numbered else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5 if numbered else 0)
        p.paragraph_format.line_spacing = 1.02
        prefix = f"{index}. " if numbered else ""
        set_run_font(p.add_run(prefix + item), size=size, bold=bold, color=color)


def add_table_caption_before(doc, anchor, text: str):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = True
    run.italic = True
    run.font.color.rgb = RGBColor(31, 78, 121)
    anchor._p.addprevious(p._p)
    return p


def add_scenario_table_before(doc, anchor, scenario) -> None:
    rows = [
        ("Cas d’utilisation", f"{scenario['code']} — {scenario['title']}", False),
        ("Acteur principal", scenario["primary"], False),
        ("Acteurs secondaires", scenario["secondary"], False),
        ("Objectif", scenario["objective"], False),
        ("Préconditions", scenario["pre"], False),
        ("Déclencheur", scenario["trigger"], False),
        ("Postconditions en cas de succès", scenario["success"], False),
        ("Postconditions en cas d’échec", scenario["failure"], False),
        ("Scénario nominal", scenario["steps"], True),
        ("Alternatives et exceptions", scenario["alts"], True),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    header = table.rows[0]
    set_repeat_header(header)
    prevent_row_split(header)
    for index, value in enumerate(("Élément", "Description")):
        cell = header.cells[index]
        set_cell_width(cell, 2350 if index == 0 else 6500)
        set_cell_margins(cell, top=110, bottom=110)
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        fill_cell(cell, value, bold=True, size=9.5, color="FFFFFF")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for row_index, (label, value, numbered) in enumerate(rows, start=1):
        row = table.add_row()
        prevent_row_split(row)
        label_cell, value_cell = row.cells
        set_cell_width(label_cell, 2350)
        set_cell_width(value_cell, 6500)
        set_cell_margins(label_cell)
        set_cell_margins(value_cell)
        set_cell_shading(label_cell, "DCE6F1")
        set_cell_shading(value_cell, "FFFFFF" if row_index % 2 else "F7FAFC")
        label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        value_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        fill_cell(label_cell, label, bold=True, size=9.0)
        label_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        fill_cell(value_cell, value, size=8.8, numbered=numbered)

    anchor._p.addprevious(table._tbl)
    spacer = doc.add_paragraph(style="Normal")
    spacer.paragraph_format.space_after = Pt(0)
    anchor._p.addprevious(spacer._p)


def find_section_break_before(doc, paragraph_index: int):
    for p in reversed(doc.paragraphs[:paragraph_index]):
        ppr = p._p.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            return p
    raise RuntimeError("Section break not found")


def clear_reserved_area(doc, section_break) -> None:
    all_paragraphs = doc.paragraphs
    break_index = next(i for i, p in enumerate(all_paragraphs) if p._p is section_break._p)
    previous_break_index = max(
        i
        for i, p in enumerate(all_paragraphs[:break_index])
        if p._p.find(qn("w:pPr")) is not None
        and p._p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None
    )
    placeholders = all_paragraphs[previous_break_index + 1 : break_index]
    for p in placeholders:
        if p.text.strip() or p._p.xpath(".//w:drawing | .//w:pict"):
            raise RuntimeError("The reserved Chapter 3 area is not empty")
    for p in placeholders:
        base.remove_paragraph(p)


def main() -> None:
    doc = Document(str(SOURCE))
    chapter4_index, chapter4 = next(
        (i, p)
        for i, p in enumerate(doc.paragraphs)
        if p.text.strip().startswith("Chapitre 4")
    )
    section_break = find_section_break_before(doc, chapter4_index)
    clear_reserved_area(doc, section_break)

    global_image = base.crop_caption(
        base.GLOBAL_SOURCE,
        base.ASSET_DIR / "global_sans_legende.png",
        bottom_pixels=260,
    )
    detail_images = []
    for uc_id, title, relative in base.USE_CASES:
        source = base.DETAIL_ROOT / relative
        target = base.ASSET_DIR / f"{uc_id}_sans_legende.png"
        detail_images.append((uc_id, title, base.crop_caption(source, target, 145)))

    by_code = {scenario["code"]: scenario for scenario in SCENARIOS}
    if set(by_code) != {code for code, _, _ in detail_images}:
        raise RuntimeError("Scenario list and diagram list do not match")

    base.add_heading_before(
        doc,
        section_break,
        "3. Modélisation fonctionnelle par cas d’utilisation",
        "Heading 2",
    )
    base.add_body_before(
        doc,
        section_break,
        "Cette section présente les interactions entre les acteurs de Matchia et les principales fonctionnalités de la plateforme. Les acteurs humains sont placés à gauche, les systèmes externes à droite et les relations UML « include » et « extend » distinguent les comportements obligatoires des comportements conditionnels.",
    )
    base.add_heading_before(doc, section_break, "3.1. Diagramme global", "Heading 3")
    base.add_body_before(
        doc,
        section_break,
        "Le diagramme global regroupe les parcours de l’Internaute, du Client, du Concessionnaire, de l’Administrateur Banque et de l’Administrateur SaaS, ainsi que les interactions avec les services externes.",
    )
    base.add_figure_before(
        doc,
        section_break,
        global_image,
        "Diagramme global des cas d’utilisation de Matchia",
        6.15,
    )
    base.add_caption_before(
        doc,
        section_break,
        "Figure 3.1 — Diagramme global des cas d’utilisation de Matchia",
    )

    base.add_heading_before(
        doc,
        section_break,
        "3.2. Diagrammes et scénarios de cas d’utilisation détaillés",
        "Heading 3",
        page_break=True,
    )
    base.add_body_before(
        doc,
        section_break,
        "Chaque diagramme est suivi d’un tableau de scénario précisant les acteurs, les conditions d’exécution, le déroulement nominal et les principales alternatives. Les cas sont regroupés par objectif métier et l’authentification n’est incluse que pour les parcours protégés.",
    )

    for index, (uc_id, title, image_path) in enumerate(detail_images, start=1):
        scenario = by_code[uc_id]
        base.add_heading_before(
            doc,
            section_break,
            f"3.2.{index}. {uc_id} — {title}",
            "Heading 4",
            page_break=index > 1,
        )
        base.add_body_before(
            doc,
            section_break,
            f"Ce cas d’utilisation décrit le processus « {title} ». La figure présente les relations UML et le tableau formalise le scénario fonctionnel correspondant.",
        )
        base.add_figure_before(
            doc,
            section_break,
            image_path,
            f"{uc_id} — {title}",
            6.15,
        )
        base.add_caption_before(
            doc,
            section_break,
            f"Figure 3.{index + 1} — Diagramme de cas d’utilisation détaillé : {title}",
        )
        add_table_caption_before(
            doc,
            section_break,
            f"Tableau 3.{index} — Description du scénario {uc_id}",
        )
        add_scenario_table_before(doc, section_break, scenario)

    chapter4.paragraph_format.page_break_before = True
    doc.core_properties.title = "Master Report Matchia avec diagrammes et scénarios de cas d’utilisation"
    doc.core_properties.subject = "Modélisation UML et description des scénarios fonctionnels"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)
    print(f"Diagrams inserted: {1 + len(detail_images)}")
    print(f"Scenario tables inserted: {len(SCENARIOS)}")
    print(f"Total tables: {len(doc.tables)}")


if __name__ == "__main__":
    main()
