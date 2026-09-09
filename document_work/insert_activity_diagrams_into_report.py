from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\PFE M2\Platforme SaaS")
SOURCE = ROOT / "document_work" / "Master Report - diagrammes et tableaux de scenarios use case.docx"
OUTPUT = ROOT / "document_work" / "Master Report - diagrammes UML complets.docx"
DIAGRAMS = ROOT / "document_work" / "diagrammes_activite_matchia"


def find_paragraph(document, predicate):
    for paragraph in document.paragraphs:
        if predicate(paragraph.text.strip()):
            return paragraph
    raise ValueError("Paragraph not found")


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def add_text_before(document, anchor, text: str):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.add_run(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_with_next = True
    anchor._p.addprevious(paragraph._p)
    return paragraph


def add_picture_before(document, anchor, image: Path, alt_text: str, width=5.35):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    shape = paragraph.add_run().add_picture(str(image), width=Inches(width))
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def add_caption_before(document, anchor, text: str):
    paragraph = document.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.bold = True
    run.italic = True
    run.font.color.rgb = RGBColor(31, 78, 121)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def format_existing_caption(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.style = "Caption"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.bold = True
    run.italic = True
    run.font.color.rgb = RGBColor(31, 78, 121)


def renumber_chapter4_figures(document) -> None:
    chapter4 = find_paragraph(document, lambda t: t.startswith("Chapitre 4"))
    chapter5 = find_paragraph(document, lambda t: t.startswith("Chapitre 5"))
    active = False
    pattern = re.compile(r"Figure 4\.(\d+)")
    for paragraph in document.paragraphs:
        if paragraph._p is chapter4._p:
            active = True
        if paragraph._p is chapter5._p:
            break
        if not active or "Figure 4." not in paragraph.text:
            continue

        def replacement(match):
            number = int(match.group(1))
            return f"Figure 4.{number + 1}" if number >= 2 else match.group(0)

        updated = pattern.sub(replacement, paragraph.text)
        if updated != paragraph.text:
            paragraph.text = updated


def remove_existing_creation_activity(document, creation_anchor) -> None:
    current = creation_anchor._p.getprevious()
    inspected = 0
    while current is not None and inspected < 6:
        if current.xpath(".//w:drawing | .//w:pict"):
            from docx.text.paragraph import Paragraph

            remove_paragraph(Paragraph(current, creation_anchor._parent))
            return
        current = current.getprevious()
        inspected += 1
    raise RuntimeError("Existing marketplace activity diagram was not found")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    required = {
        "auth": DIAGRAMS / "AD-01_authentification.png",
        "creation": DIAGRAMS / "AD-02_creation_demande_marketplace.png",
        "treatment": DIAGRAMS / "AD-03_traitement_demande_marketplace.png",
        "dealer": DIAGRAMS / "AD-04_parcours_concessionnaire.png",
        "client": DIAGRAMS / "AD-05_parcours_client.png",
    }
    for image in required.values():
        if not image.exists():
            raise FileNotFoundError(image)

    document = Document(str(SOURCE))
    renumber_chapter4_figures(document)

    auth_anchor = find_paragraph(
        document, lambda t: t.startswith("3. Demande de création d’une marketplace bancaire")
    )
    creation_anchor = find_paragraph(
        document, lambda t: t.startswith("3.1. Renseignement des informations de la banque")
    )
    treatment_caption = find_paragraph(
        document, lambda t: "Diagramme d’activité du traitement d’une demande" in t
    )
    dealer_caption = find_paragraph(
        document, lambda t: "Diagramme d’activité du parcours d’inscription et de validation d’un concessionnaire" in t
    )
    client_caption = find_paragraph(
        document, lambda t: "Diagramme d’activité du parcours de demande de financement" in t
    )
    treatment_followup_heading = find_paragraph(
        document, lambda t: t.startswith("4.3.2. Approbation et déclenchement du paiement")
    )

    remove_existing_creation_activity(document, creation_anchor)

    add_text_before(
        document,
        auth_anchor,
        "Le diagramme d’activité suivant synthétise le contrôle des identifiants et de l’état du compte, le chargement du rôle et des autorisations ainsi que la détermination du contexte du tenant lorsque l’utilisateur est rattaché à une banque ou à une marketplace.",
    )
    add_picture_before(
        document,
        auth_anchor,
        required["auth"],
        "Diagramme d’activité de l’authentification et de la détermination du contexte d’accès",
    )
    add_caption_before(
        document,
        auth_anchor,
        "Figure 4.1 — Diagramme d’activité de l’authentification et du contexte d’accès",
    )

    add_picture_before(
        document,
        creation_anchor,
        required["creation"],
        "Diagramme d’activité de création et de soumission d’une demande de marketplace bancaire",
    )
    add_caption_before(
        document,
        creation_anchor,
        "Figure 4.2 — Diagramme d’activité de création d’une demande de marketplace bancaire",
    )
    creation_anchor.paragraph_format.page_break_before = True

    add_picture_before(
        document,
        treatment_caption,
        required["treatment"],
        "Diagramme d’activité du traitement d’une demande de marketplace par l’Administrateur SaaS",
    )
    format_existing_caption(
        treatment_caption,
        "Figure 4.12 — Diagramme d’activité du traitement d’une demande de marketplace",
    )
    treatment_followup_heading.paragraph_format.page_break_before = True

    add_picture_before(
        document,
        dealer_caption,
        required["dealer"],
        "Diagramme d’activité du parcours métier du Concessionnaire",
    )
    format_existing_caption(
        dealer_caption,
        "Figure 5.1 — Diagramme d’activité du parcours métier du Concessionnaire",
    )

    add_picture_before(
        document,
        client_caption,
        required["client"],
        "Diagramme d’activité du parcours Client et de la demande de financement",
    )
    format_existing_caption(
        client_caption,
        "Figure 5.4 — Diagramme d’activité du parcours Client et de la demande de financement",
    )

    document.core_properties.title = "Master Report Matchia avec diagrammes UML complets"
    document.core_properties.subject = "Cas d’utilisation, scénarios et diagrammes d’activité"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    print(OUTPUT)
    print(f"inline_shapes={len(document.inline_shapes)}")
    print(f"tables={len(document.tables)}")


if __name__ == "__main__":
    main()
