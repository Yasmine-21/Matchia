from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont


INPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - rapport mis à jour - backlog lisible.docx")
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - chapitre 6 DevOps et Azure.docx")
ROOT = Path(r"D:\PFE M2\Platforme SaaS")
FIGURES = ROOT / "figure"
DIAGRAMS = ROOT / "document_work" / "diagrams_matchia"
GENERATED = ROOT / "document_work" / "chapter6_figures"


def source_run(paragraph):
    return next((run for run in paragraph.runs if run.text), None)


def copy_paragraph_properties(target, template):
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if template._p.pPr is not None:
        target._p.insert(0, deepcopy(template._p.pPr))


def insert_text(anchor, text, template):
    paragraph = anchor.insert_paragraph_before()
    copy_paragraph_properties(paragraph, template)
    run = paragraph.add_run(text)
    ref = source_run(template)
    if ref is not None and ref._r.rPr is not None:
        run._r.insert(0, deepcopy(ref._r.rPr))
    return paragraph


def insert_figure(document, anchor, image_path, caption, caption_template):
    image = Image.open(image_path)
    ratio = image.width / image.height
    width = Inches(5.85 if ratio >= 1.25 else 4.25)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=width)
    anchor._p.addprevious(paragraph._p)

    caption_paragraph = insert_text(anchor, caption, caption_template)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return caption_paragraph


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)
    paragraph._p = paragraph._element = None


def create_repository_figure():
    GENERATED.mkdir(parents=True, exist_ok=True)
    path = GENERATED / "fig_6_2_organisation_depot.png"
    canvas = Image.new("RGB", (1800, 950), "white")
    draw = ImageDraw.Draw(canvas)
    regular = ImageFont.truetype(r"C:\Windows\Fonts\consola.ttf", 43)
    title = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", 48)
    draw.text((100, 70), "Organisation du dépôt GitHub Matchia", fill="black", font=title)
    lines = [
        "Matchia/",
        "├── MatchiaBackend/     API Spring Boot et tests Maven",
        "├── MatchiaFrontend/    application React TypeScript",
        "├── docker-compose.yml  orchestration locale des conteneurs",
        "└── Jenkinsfile         pipeline CI/CD versionné",
    ]
    y = 220
    for index, line in enumerate(lines):
        draw.text((150, y), line, fill=(0, 0, 0), font=regular)
        y += 120 if index == 0 else 110
    canvas.save(path)
    return path


def main():
    document = Document(INPUT)
    paragraphs = document.paragraphs
    chapter_start = next(i for i, p in enumerate(paragraphs) if p.text.strip().startswith("Chapitre 6"))
    old_chapter = paragraphs[chapter_start:]
    anchor = paragraphs[chapter_start]

    chapter_template = paragraphs[chapter_start]
    body_template = paragraphs[chapter_start + 4]
    section_template = paragraphs[chapter_start + 5]
    subsection_template = paragraphs[chapter_start + 8]
    caption_template = next(p for p in paragraphs if p.style.name == "Caption")
    repo_figure = create_repository_figure()

    def heading(text):
        insert_text(anchor, text, section_template)

    def subheading(text):
        insert_text(anchor, text, subsection_template)

    def paragraph(text):
        insert_text(anchor, text, body_template)

    def figure(path, caption):
        insert_figure(document, anchor, path, caption, caption_template)

    insert_text(anchor, "Chapitre 6 : Mise en place de la démarche DevOps, du pipeline CI/CD et déploiement de Matchia sur Microsoft Azure", chapter_template)

    heading("1. Introduction")
    paragraph("Après la conception et le développement de la plateforme, une démarche de livraison et de déploiement a été mise en place afin que chaque nouvelle version puisse être construite, vérifiée et déployée de manière reproductible. L’objectif ne se limite pas à rendre l’application accessible dans le Cloud : il consiste également à automatiser les contrôles qui précèdent la mise en production.")
    paragraph("La solution repose sur GitHub pour la centralisation du code, Docker pour la conteneurisation, Jenkins pour l’orchestration du pipeline, SonarQube pour l’analyse de la qualité, Docker Hub pour la distribution des images et Microsoft Azure pour l’hébergement. Azure CLI permet enfin au pipeline d’appliquer les mises à jour sur les services Cloud.")
    paragraph("La figure suivante présente la chaîne globale allant du push du développeur jusqu’au déploiement des conteneurs dans Microsoft Azure.")
    figure(DIAGRAMS / "fig_6_1_cicd_pipeline.png", "Figure 6.1 — Vue générale de la chaîne DevOps et CI/CD de Matchia")

    heading("2. Démarche DevOps et principes de la CI/CD")
    subheading("2.1. Démarche DevOps")
    paragraph("Le DevOps est une démarche qui associe étroitement les activités de développement et d’exploitation afin d’améliorer la collaboration, la qualité des livraisons et la fiabilité des déploiements. Dans ce projet, il se traduit par un processus automatisé qui enchaîne la construction, les tests, l’analyse du code, la création d’images Docker et le déploiement des versions validées.")
    subheading("2.2. Intégration et déploiement continus")
    paragraph("L’intégration continue soumet chaque modification à des opérations automatiques de compilation, de test et de contrôle de qualité. Le déploiement continu prolonge cette vérification en construisant les images Docker, en les publiant dans un registre puis en mettant à jour les services Azure. Une étape critique en échec interrompt le pipeline et empêche le déploiement d’une version non validée.")

    heading("3. Préparation et validation de la conteneurisation")
    paragraph("Avant d’automatiser la livraison, les différentes étapes techniques ont été mises en place et validées séparément. Les constructions du backend et du frontend, la création des images Docker et l’exécution locale des conteneurs ont d’abord été réalisées manuellement. Cette phase a permis de vérifier les Dockerfiles, les variables d’environnement, les ports exposés et la communication entre les services avant de confier leur enchaînement à Jenkins.")
    subheading("3.1. Gestion des versions avec Git et GitHub")
    paragraph("Git assure le suivi local des évolutions du projet tandis que GitHub centralise le dépôt distant. Le dépôt rassemble le backend Spring Boot, le frontend React, le fichier docker-compose.yml et le Jenkinsfile. Il devient ainsi le point d’entrée du pipeline : un push sur la branche suivie par Jenkins déclenche la récupération de la nouvelle version du code.")
    paragraph("La figure 6.2 met en évidence l’organisation du dépôt et les principaux fichiers nécessaires à la construction et au déploiement de l’application.")
    figure(repo_figure, "Figure 6.2 — Organisation du dépôt GitHub de la plateforme Matchia")

    subheading("3.2. Conteneurisation du backend Spring Boot")
    paragraph("Le backend est développé avec Spring Boot et Java 17. Son Dockerfile suit une construction en plusieurs étapes : Maven compile d’abord l’application, puis une image Java d’exécution plus légère ne conserve que le fichier JAR produit. Le service expose l’API REST sur le port 8081. L’image a été construite et exécutée manuellement lors de la phase initiale afin de valider le démarrage de l’API avant son intégration au pipeline.")
    figure(FIGURES / "imageBackend.png", "Figure 6.3 — Dockerfile utilisé pour la conteneurisation du backend Matchia")

    subheading("3.3. Conteneurisation du frontend React")
    paragraph("Le frontend React, TypeScript et Vite est également construit en plusieurs étapes. Node.js installe les dépendances et génère les fichiers statiques de production ; ceux-ci sont ensuite servis par Nginx. La variable VITE_API_URL est fournie au moment du build afin d’associer le frontend à l’URL du backend correspondant à l’environnement cible. La construction et l’exécution de cette image ont été vérifiées manuellement avant leur automatisation.")
    figure(FIGURES / "imageFrontend.png", "Figure 6.4 — Dockerfile utilisé pour la conteneurisation du frontend Matchia")

    subheading("3.4. Orchestration locale avec Docker Compose")
    paragraph("Docker Compose décrit les services frontend et backend, les ports exposés, les variables d’environnement, le volume des fichiers téléversés et le réseau de communication. La commande docker compose up -d permet de démarrer un environnement local cohérent. Cette exécution manuelle a permis de vérifier que les deux conteneurs communiquent correctement et que l’application est fonctionnelle avant l’automatisation de la chaîne CI/CD.")
    figure(FIGURES / "dockercompose.png", "Figure 6.5 — Configuration locale des services Matchia avec Docker Compose")

    heading("4. Mise en place du pipeline CI/CD avec Jenkins")
    paragraph("Une fois les étapes de construction et de conteneurisation validées, Jenkins a été configuré pour les enchaîner automatiquement. Le pipeline reprend donc un processus déjà testé, en garantissant qu’il est exécuté de manière identique à chaque nouvelle version du code.")
    subheading("4.1. Rôle de Jenkins")
    paragraph("Jenkins est le composant central de la chaîne CI/CD. Exécuté dans un conteneur Docker, il récupère le code depuis GitHub et coordonne les outils nécessaires au pipeline : Maven, Node.js, Docker, SonarQube, Docker Hub et Azure CLI. Les accès aux services externes sont gérés à l’aide de credentials configurés dans Jenkins, afin de ne pas intégrer les informations d’authentification dans le dépôt source.")
    figure(FIGURES / "jenkins.png", "Figure 6.6 — Interface du serveur Jenkins configuré pour le projet Matchia")

    subheading("4.2. Pipeline as Code avec le Jenkinsfile")
    paragraph("La définition du pipeline est versionnée dans le Jenkinsfile placé à la racine du dépôt. Cette approche, appelée Pipeline as Code, garantit que les évolutions de la chaîne CI/CD sont suivies avec les évolutions applicatives. Le pipeline comporte les étapes Checkout, Build Backend, Build Frontend, Tests Backend, analyses SonarQube, Docker Build, publication des images puis déploiements du backend et du frontend sur Azure.")
    figure(FIGURES / "CréationPipeline.png", "Figure 6.7 — Définition du pipeline CI/CD de Matchia dans Jenkins")

    subheading("4.3. Déclenchement automatique avec GitHub Webhook et Cloudflare Tunnel")
    paragraph("Un Webhook GitHub notifie Jenkins lorsqu’un push est réalisé sur le dépôt. Le chemin /github-webhook/ est configuré afin que Jenkins puisse recevoir cet événement et démarrer automatiquement le pipeline. Ce mécanisme évite le lancement manuel d’un build après chaque modification.")
    figure(FIGURES / "ConfigWebhookSurGithub.png", "Figure 6.8 — Configuration du Webhook GitHub déclenchant Jenkins")
    paragraph("Jenkins étant exécuté localement pendant la phase de mise en place, son adresse localhost n’était pas accessible depuis GitHub. Un Cloudflare Tunnel a donc été utilisé pour exposer temporairement le serveur Jenkins par une adresse HTTPS publique et rediriger les événements vers l’instance locale.")
    figure(FIGURES / "CréationTunnelCloudflareTempVersJenkins.png", "Figure 6.9 — Exposition temporaire de Jenkins au moyen de Cloudflare Tunnel")

    heading("5. Exécution automatisée du pipeline CI/CD")
    paragraph("Le pipeline est déclenché à chaque push reçu depuis GitHub. Il applique ensuite, dans un ordre contrôlé, les étapes suivantes : récupération du code, construction du backend et du frontend, exécution des tests, analyse de qualité, création des images Docker, publication dans Docker Hub et déploiement des nouvelles versions sur Microsoft Azure. En cas d’échec d’une étape critique, les étapes suivantes ne sont pas exécutées.")
    subheading("5.1. Récupération du code et construction des applications")
    paragraph("Après le checkout, Jenkins construit le backend avec Maven puis le frontend avec Node.js et npm. Le backend est préparé avec la commande Maven de construction définie dans le Jenkinsfile ; le frontend exécute npm ci puis npm run build. Lors du build de production, VITE_API_URL reçoit l’adresse publique du backend Azure afin que le frontend déployé communique avec l’API cible.")
    figure(FIGURES / "CommandeBackend.png", "Figure 6.10 — Exécution de la phase de construction du backend dans le pipeline")

    subheading("5.2. Exécution automatique des tests backend")
    paragraph("Les tests backend sont exécutés avant la création des images Docker. Ils constituent une barrière de qualité : si l’exécution échoue, Jenkins interrompt le pipeline et les opérations de publication et de déploiement ne sont pas réalisées. Les rapports de couverture générés avec JaCoCo sont également utilisables lors de l’analyse SonarQube.")
    figure(FIGURES / "testBackend.png", "Figure 6.11 — Résultat de l’exécution automatique des tests backend")

    subheading("5.3. Analyse statique avec SonarQube")
    paragraph("SonarQube complète les tests en réalisant une analyse statique distincte sur le backend et le frontend. Les indicateurs suivis portent notamment sur les bugs, les vulnérabilités, les duplications, les odeurs de code et la couverture disponible. Cette étape aide à détecter les défauts qui ne sont pas nécessairement révélés par les tests fonctionnels.")
    figure(FIGURES / "SonarQubeProjects.png", "Figure 6.12 — Projets Matchia analysés dans SonarQube")

    subheading("5.4. Construction et publication des images Docker")
    paragraph("Après la validation des builds, des tests et de l’analyse de qualité, Jenkins utilise les Dockerfiles du projet pour construire une image dédiée au backend et une autre au frontend. Les images sont étiquetées avec le numéro du build Jenkins et avec le tag latest, ce qui assure la traçabilité entre une exécution du pipeline et l’artefact produit.")
    paragraph("Jenkins s’authentifie ensuite auprès de Docker Hub et publie les images du projet. Docker Hub joue ainsi le rôle de registre distant entre l’environnement de construction et les Azure Container Apps, qui peuvent récupérer la version à déployer.")
    figure(FIGURES / "créationRepoDockerHub.png", "Figure 6.13 — Registre Docker Hub utilisé pour publier les images Matchia")

    heading("6. Déploiement de Matchia sur Microsoft Azure")
    subheading("6.1. Préparation de l’environnement Cloud")
    paragraph("Microsoft Azure a été retenu afin d’héberger les composants de la plateforme dans un environnement Cloud. L’architecture répartit les responsabilités entre Azure Container Apps pour les conteneurs applicatifs, Azure Database for PostgreSQL pour les données métier et Azure Files pour les fichiers persistants. Cette séparation préserve les données et les fichiers lors du remplacement d’une révision de conteneur.")
    figure(FIGURES / "azure.png", "Figure 6.14 — Ressources Azure mises en place pour l’hébergement de Matchia")

    subheading("6.2. Base de données PostgreSQL managée")
    paragraph("La base de données PostgreSQL est hébergée dans Azure Database for PostgreSQL Flexible Server. Le backend y accède au moyen d’une chaîne JDBC fournie par des variables d’environnement, sans inscrire les paramètres de connexion dans le code source. La migration initiale peut être réalisée avec pg_dump puis pg_restore afin de conserver la continuité des données entre l’environnement local et l’environnement Cloud.")
    figure(FIGURES / "ConfigurationDataBase.png", "Figure 6.15 — Configuration de la base PostgreSQL Azure de Matchia")

    subheading("6.3. Stockage persistant avec Azure Files")
    paragraph("Les conteneurs applicatifs étant remplaçables, les fichiers déposés dans leur système de fichiers interne ne doivent pas constituer le stockage permanent. Azure Files fournit donc un partage monté dans le backend, notamment pour conserver les documents justificatifs, les images de produits, les logos et les autres contenus téléversés. Le montage du partage garantit que ces fichiers demeurent disponibles après la création d’une nouvelle révision.")
    figure(FIGURES / "StockageAzureDataBase.png", "Figure 6.16 — Ressource Azure de stockage persistant associée à Matchia")

    subheading("6.4. Hébergement avec Azure Container Apps")
    paragraph("Deux Azure Container Apps indépendantes exécutent l’application : matchia-frontend sert les fichiers React compilés à travers Nginx, tandis que matchia-backend expose les API REST Spring Boot sur le port 8081. L’Ingress Azure rend les services accessibles en HTTPS et les deux composants peuvent être mis à jour indépendamment.")
    figure(FIGURES / "ConfigContainerFrontend.png", "Figure 6.17 — Configuration de l’Azure Container App du frontend")
    figure(FIGURES / "ConfigurationContainerBackend.png", "Figure 6.18 — Configuration de l’Azure Container App du backend")

    subheading("6.5. Déploiement automatisé avec Azure CLI")
    paragraph("Après la publication réussie des images dans Docker Hub, Azure CLI assure la dernière étape du pipeline. Jenkins s’authentifie à Azure, sélectionne l’environnement configuré puis exécute une commande az containerapp update pour mettre à jour successivement la Container App du backend et celle du frontend. Une nouvelle révision est créée avec le numéro du build ; cette révision fournit une version identifiable et permet de suivre le résultat du déploiement.")
    figure(FIGURES / "DéploiementRéussi.png", "Figure 6.19 — Création d’une nouvelle révision Azure Container Apps après déploiement")

    heading("7. Architecture complète et validation du processus")
    paragraph("L’architecture finale relie le poste du développeur, GitHub, le Webhook, Jenkins, SonarQube, Docker Hub et les ressources Azure. À chaque push, Jenkins récupère le code, construit les deux composants, exécute les tests et les analyses, publie les images puis demande à Azure Container Apps d’utiliser les nouvelles versions. Le backend déployé communique avec PostgreSQL Azure et avec le stockage persistant Azure Files.")
    figure(DIAGRAMS / "fig_6_2_azure_deployment_architecture.png", "Figure 6.20 — Architecture globale de la chaîne CI/CD et du déploiement Cloud de Matchia")

    subheading("7.1. Validation du processus complet")
    paragraph("La validation de la chaîne consiste à effectuer une modification contrôlée du code, à l’enregistrer puis à exécuter un git push sur la branche suivie. Le Webhook déclenche Jenkins, qui exécute successivement les étapes de construction, de test, d’analyse, de création d’images, de publication et de déploiement. La disponibilité d’une nouvelle révision des Container Apps après une exécution réussie confirme le fonctionnement de bout en bout du processus automatisé.")

    heading("8. Conclusion")
    paragraph("Ce chapitre a présenté la mise en place du pipeline CI/CD de Matchia, depuis la validation manuelle initiale de la conteneurisation jusqu’au déploiement automatisé sur Microsoft Azure. Jenkins orchestre désormais une chaîne complète, contrôlée et reproductible : récupération du code, construction, tests, analyse de qualité, publication des images Docker et mise à jour des services Cloud.")

    for old in old_chapter:
        remove_paragraph(old)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
