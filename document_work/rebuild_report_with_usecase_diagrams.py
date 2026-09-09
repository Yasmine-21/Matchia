from __future__ import annotations

from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\ASUS\Downloads\Master Report - backlog chapitre 4 détaillé.docx")
WORKSPACE = Path(r"D:\PFE M2\Platforme SaaS")
OUTPUT = WORKSPACE / "document_work" / "Master Report - nouveaux diagrammes use case.docx"
ASSET_DIR = WORKSPACE / "document_work" / "report_diagram_assets"
GLOBAL_SOURCE = (
    WORKSPACE
    / "document_work"
    / "diagramme_usecase_global_matchia"
    / "Diagramme_cas_utilisation_global_Matchia_resume_optimise.png"
)
DETAIL_ROOT = WORKSPACE / "document_work" / "diagrammes_usecase_detailles_matchia"


USE_CASES = [
    ("UC-01", "S’authentifier", "UC-01_authentification/UC-01_S_authentifier.png"),
    (
        "UC-02",
        "Soumettre une demande de marketplace bancaire",
        "UC-02_demande_marketplace/UC-02_Soumettre_une_demande_de_marketplace_bancaire.png",
    ),
    (
        "UC-03",
        "Administrer et superviser la plateforme SaaS",
        "UC-03_administration_saas/UC-03_Administrer_et_superviser_la_plateforme_SaaS.png",
    ),
    (
        "UC-04",
        "Traiter une demande de marketplace",
        "UC-04_traitement_demande_marketplace/UC-04_Traiter_une_demande_de_marketplace.png",
    ),
    (
        "UC-05",
        "Payer et activer la marketplace",
        "UC-05_paiement_activation_marketplace/UC-05_Payer_et_activer_la_marketplace.png",
    ),
    (
        "UC-06",
        "Administrer la marketplace bancaire",
        "UC-06_administration_banque/UC-06_Administrer_la_marketplace_bancaire.png",
    ),
    (
        "UC-07",
        "Renouveler ou ajouter des services",
        "UC-07_renouvellement_services/UC-07_Renouveler_ou_ajouter_des_services.png",
    ),
    (
        "UC-08",
        "Consulter la marketplace, comparer et simuler",
        "UC-08_marketplace_publique/UC-08_Consulter_la_marketplace_comparer_et_simuler.png",
    ),
    (
        "UC-09",
        "Déposer une demande d’inscription comme Concessionnaire",
        "UC-09_inscription_concessionnaire/UC-09_Deposer_une_demande_d_inscription_comme_Concessionnaire.png",
    ),
    (
        "UC-10",
        "Traiter une demande d’inscription Concessionnaire",
        "UC-10_validation_concessionnaire/UC-10_Traiter_une_demande_d_inscription_Concessionnaire.png",
    ),
    (
        "UC-11",
        "Gérer les partenariats et les contrats",
        "UC-11_partenariats_contrats/UC-11_Gerer_les_partenariats_et_les_contrats.png",
    ),
    (
        "UC-12",
        "Gérer les produits, les stocks et les publications",
        "UC-12_produits_publications_stock/UC-12_Gerer_les_produits_les_stocks_et_les_publications.png",
    ),
    (
        "UC-13",
        "S’inscrire comme Client",
        "UC-13_inscription_client/UC-13_S_inscrire_comme_Client.png",
    ),
    (
        "UC-14",
        "Constituer et soumettre une demande de financement",
        "UC-14_demande_financement_client/UC-14_Constituer_et_soumettre_une_demande_de_financement.png",
    ),
    (
        "UC-15",
        "Traiter une demande de financement",
        "UC-15_traitement_financement_banque/UC-15_Traiter_une_demande_de_financement.png",
    ),
    (
        "UC-16",
        "Interroger l’assistant IA du SaaS",
        "UC-16_assistant_ia_saas/UC-16_Interroger_l_assistant_IA_du_SaaS.png",
    ),
    (
        "UC-17",
        "Utiliser le chatbot public",
        "UC-17_chatbot_public/UC-17_Utiliser_le_chatbot_public.png",
    ),
]


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def crop_caption(source: Path, target: Path, bottom_pixels: int) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as im:
        cropped = im.crop((0, 0, im.width, im.height - bottom_pixels))
        cropped.save(target, optimize=True)
    return target


def move_before(paragraph, anchor) -> None:
    anchor._p.addprevious(paragraph._p)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def add_heading_before(doc, anchor, text: str, style: str, page_break: bool = False):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    p.paragraph_format.page_break_before = page_break
    p.paragraph_format.space_after = Pt(6)
    set_keep_with_next(p)
    move_before(p, anchor)
    return p


def add_body_before(doc, anchor, text: str):
    p = doc.add_paragraph(style="Normal")
    p.add_run(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    set_keep_with_next(p)
    move_before(p, anchor)
    return p


def add_figure_before(doc, anchor, image_path: Path, alt_text: str, width_inches: float):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_together = True
    set_keep_with_next(p)
    run = p.add_run()
    shape = run.add_picture(str(image_path), width=Inches(width_inches))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", alt_text)
    move_before(p, anchor)
    return p


def add_caption_before(doc, anchor, text: str):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    run.bold = True
    run.italic = True
    run.font.color.rgb = RGBColor(31, 78, 121)
    move_before(p, anchor)
    return p


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not GLOBAL_SOURCE.exists():
        raise FileNotFoundError(GLOBAL_SOURCE)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(SOURCE))

    chapter4_index, chapter4 = next(
        (i, p)
        for i, p in enumerate(doc.paragraphs)
        if p.text.strip().startswith("Chapitre 4")
    )

    # The paragraph immediately before Chapter 4 carries the section break.
    section_break = None
    for p in reversed(doc.paragraphs[:chapter4_index]):
        ppr = p._p.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            section_break = p
            break
    if section_break is None:
        raise RuntimeError("Section break before Chapter 4 not found")

    # Remove only the empty placeholders located between the technology section
    # and this section break; they were reserved for the missing conception part.
    all_paragraphs = doc.paragraphs
    break_index = next(i for i, p in enumerate(all_paragraphs) if p._p is section_break._p)
    previous_break_index = max(
        i
        for i, p in enumerate(all_paragraphs[:break_index])
        if p._p.find(qn("w:pPr")) is not None
        and p._p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None
    )
    placeholders = all_paragraphs[previous_break_index + 1 : break_index]
    for p in placeholders:
        if p.text.strip() or p._p.xpath(".//w:drawing | .//w:pict"):
            raise RuntimeError("A non-empty element was found in the reserved Chapter 3 area")
    for p in placeholders:
        remove_paragraph(p)

    global_image = crop_caption(
        GLOBAL_SOURCE, ASSET_DIR / "global_sans_legende.png", bottom_pixels=260
    )
    detail_images = []
    for uc_id, title, relative in USE_CASES:
        source = DETAIL_ROOT / relative
        if not source.exists():
            raise FileNotFoundError(source)
        target = ASSET_DIR / f"{uc_id}_sans_legende.png"
        detail_images.append((uc_id, title, crop_caption(source, target, 145)))

    add_heading_before(
        doc,
        section_break,
        "3. Modélisation fonctionnelle par cas d’utilisation",
        "Heading 2",
    )
    add_body_before(
        doc,
        section_break,
        "Cette section présente une vue synthétique des interactions entre les acteurs de Matchia et les principales fonctionnalités de la plateforme. Les acteurs humains sont placés à gauche, les systèmes externes à droite et les relations UML « include » et « extend » distinguent les comportements obligatoires des comportements conditionnels.",
    )
    add_heading_before(doc, section_break, "3.1. Diagramme global", "Heading 3")
    add_body_before(
        doc,
        section_break,
        "Le diagramme global regroupe les parcours de l’Internaute, du Client, du Concessionnaire, de l’Administrateur Banque et de l’Administrateur SaaS, ainsi que les interactions avec les services externes.",
    )
    add_figure_before(
        doc,
        section_break,
        global_image,
        "Diagramme global des cas d’utilisation de Matchia",
        6.15,
    )
    add_caption_before(
        doc,
        section_break,
        "Figure 3.1 — Diagramme global des cas d’utilisation de Matchia",
    )

    add_heading_before(
        doc,
        section_break,
        "3.2. Diagrammes de cas d’utilisation détaillés",
        "Heading 3",
        page_break=True,
    )
    add_body_before(
        doc,
        section_break,
        "Les diagrammes suivants détaillent les cas d’utilisation issus des besoins fonctionnels et les regroupent par objectif métier, sans référence aux sprints. L’authentification n’est incluse que lorsque l’accès à un espace protégé est requis.",
    )

    for idx, (uc_id, title, image_path) in enumerate(detail_images, start=1):
        add_heading_before(
            doc,
            section_break,
            f"3.2.{idx}. {uc_id} — {title}",
            "Heading 4",
            page_break=idx > 1,
        )
        add_body_before(
            doc,
            section_break,
            f"Ce diagramme présente les acteurs, les sous-fonctions obligatoires et les comportements conditionnels associés au cas d’utilisation « {title} ».",
        )
        add_figure_before(
            doc,
            section_break,
            image_path,
            f"{uc_id} — {title}",
            6.15,
        )
        add_caption_before(
            doc,
            section_break,
            f"Figure 3.{idx + 1} — Diagramme de cas d’utilisation détaillé : {title}",
        )

    # Start Chapter 4 on a fresh page even if the inherited section break is continuous.
    chapter4.paragraph_format.page_break_before = True
    doc.core_properties.title = "Master Report — Matchia"
    doc.core_properties.subject = "Rapport enrichi des diagrammes de cas d’utilisation UML"
    doc.save(str(OUTPUT))
    print(OUTPUT)
    print(f"Inserted diagrams: {1 + len(USE_CASES)}")
    print(f"Inline shapes after edit: {len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
