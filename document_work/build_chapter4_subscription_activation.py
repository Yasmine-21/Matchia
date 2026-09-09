from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw

from build_matchia_chapters_4_5 import (
    ROOT, ASSETS, font, center_text, box, arrow, setup, para, add_heading,
    add_caption, add_figure, add_placeholder, add_matrix, add_bullets,
    apply_font
)

OUT = ROOT / 'document_work' / 'Chapitre_4_Matchia_Abonnement_Activation.docx'
DIAGRAMS = ROOT / 'document_work' / 'generated_chapter4_subscription'
DIAGRAMS.mkdir(parents=True, exist_ok=True)

def base(name, h):
    image = Image.new('RGB', (2200, h), 'white')
    return image, ImageDraw.Draw(image), DIAGRAMS / name

def lifeline(draw, label, x, h):
    box(draw, (x-125, 45, x+125, 130), label, size=22, fill='F2F2F2')
    draw.line((x, 137, x, h-65), fill='#777777', width=2)

def generate_diagrams():
    # Onboarding sequence requested by the user.
    im, d, p = base('fig_4_1_sequence_adhesion_activation.png', 1440)
    actors = [('Admin SaaS', 170), ('Plateforme', 550), ('Admin Banque', 930), ('Paiement', 1320), ('Marketplace', 1710), ('Messagerie', 2050)]
    for label, x in actors: lifeline(d, label, x, 1440)
    events = [
        (930, 550, 205, 'demande d’adhésion'),
        (550, 170, 310, 'dossier à traiter'),
        (170, 550, 415, 'approbation'),
        (550, 2050, 520, 'notification système'),
        (550, 1320, 625, 'générer lien / session'),
        (1320, 2050, 730, 'email de paiement'),
        (2050, 930, 835, 'lien sécurisé'),
        (930, 1320, 940, 'transaction'),
        (1320, 550, 1045, 'confirmation du paiement'),
        (550, 1710, 1150, 'activer marketplace'),
        (550, 2050, 1255, 'email des identifiants'),
        (2050, 930, 1360, 'accès back-office Banque'),
    ]
    for x1, x2, y, label in events: arrow(d, (x1,y), (x2,y), label)
    im.save(p)

    # State/dependency diagram to explain request / subscription / payment / activation.
    im, d, p = base('fig_4_2_etats_et_dependances.png', 880)
    labels = [
        ('Demande\nEN_ATTENTE', (70, 210, 420, 390)),
        ('Demande\nAPPROUVÉE', (540, 210, 890, 390)),
        ('Abonnement\nPENDING_PAYMENT', (1010, 210, 1360, 390)),
        ('Paiement\nCONFIRMÉ', (1480, 210, 1830, 390)),
    ]
    for lab, rect in labels: box(d, rect, lab, size=28, fill='F6F6F6')
    for i in range(3): arrow(d, (labels[i][1][2],300), (labels[i+1][1][0]-12,300))
    box(d, (800, 555, 1400, 725), 'Marketplace ACTIVE\n+ accès Banque autorisé', size=30, fill='F2F2F2')
    arrow(d, (1655, 390), (1100, 550), 'activation automatique')
    d.rounded_rectangle((100, 765, 2100, 850), radius=14, fill='#FAFAFA', outline='#444444', width=2)
    center_text(d, (125, 772, 2075, 842), 'Avant confirmation : la marketplace est inactive et l’abonnement ne donne aucun droit d’exploitation du back-office.', 26, False)
    im.save(p)

    # Configuration after active subscription.
    im, d, p = base('fig_4_3_configuration_tenant.png', 820)
    steps = [('Abonnement actif', 55), ('Choisir les stores souscrits', 440), ('Activer les modules autorisés', 825), ('Renseigner paramètres et contenus', 1210), ('Publier la marketplace', 1595)]
    for label, x in steps:
        box(d, (x, 240, x+310, 415), label, size=25, fill='F7F7F7')
    for i in range(4): arrow(d, (steps[i][1]+310,327), (steps[i+1][1]-12,327))
    d.rounded_rectangle((155, 540, 2050, 705), radius=16, fill='#FAFAFA', outline='#444444', width=2)
    center_text(d, (190, 560, 2015, 685), 'Le tenant personnalise son identité, ses bannières, ses stores et ses modules sans modifier le socle commun React / Spring Boot. La disponibilité dépend des services effectivement souscrits.', 28, False)
    im.save(p)

    # Renewal/upgrade activity.
    im, d, p = base('fig_4_4_renouvellement_upgrade.png', 1000)
    nodes = [
        ('Demande de renouvellement ou upgrade', (560, 55, 1640, 170)),
        ('Analyse et validation SaaS', (650, 255, 1550, 370)),
        ('Création du paiement complémentaire', (630, 455, 1570, 570)),
        ('Confirmation du paiement', (690, 655, 1510, 770)),
        ('Mise à jour de l’abonnement et des services', (565, 855, 1635, 970)),
    ]
    for lab, rect in nodes: box(d, rect, lab, size=28, fill='F7F7F7')
    for i in range(4):
        a=nodes[i][1]; b=nodes[i+1][1]
        arrow(d, ((a[0]+a[2])//2,a[3]), ((b[0]+b[2])//2,b[1]-10))
    d.text((95, 452), 'Refus : la configuration existante reste inchangée.', font=font(23, False), fill='#202124')
    d.line((635,510,390,510), fill='#202124', width=3)
    im.save(p)

def callout(doc, title, body):
    from build_matchia_chapters_4_5 import set_table_geometry, set_cell_shading, set_cell_margins, set_repeat_header
    table = doc.add_table(rows=1, cols=1); set_table_geometry(table, [9360]); set_repeat_header(table.rows[0])
    cell=table.cell(0,0); set_cell_shading(cell, 'F4F6F9'); set_cell_margins(cell, top=120, start=180, bottom=120, end=180)
    p=cell.paragraphs[0]; r=p.add_run(title); apply_font(r, 11, '1F4D78', True)
    p=cell.add_paragraph(); r=p.add_run(body); apply_font(r, 10, '202124')
    doc.add_paragraph().paragraph_format.space_after=Pt(4)

def cover(doc):
    for _ in range(7): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('MÉMOIRE DE FIN D’ÉTUDES'); apply_font(r,11,'2E74B5',True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(16); p.paragraph_format.space_after=Pt(12)
    r=p.add_run('Chapitre 4 — Réalisation du SaaS bancaire : abonnement, paiement, activation et configuration multi-tenant'); apply_font(r,25,'0B2545',True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(18)
    r=p.add_run('Plateforme Matchia'); apply_font(r,15,'1F4D78')
    for _ in range(9): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Version réorganisée pour intégration au rapport PFE'); apply_font(r,10,'5B6573',italic=True)
    doc.add_page_break()

def write_chapter(doc):
    add_heading(doc, 'Chapitre 4 — Réalisation du SaaS bancaire : abonnement, paiement, activation et configuration multi-tenant', 1)
    para(doc, 'Ce chapitre décrit la réalisation du noyau SaaS de Matchia sous l’angle du cycle de vie commercial et technique d’une marketplace bancaire. L’objectif est d’expliquer comment une demande d’adhésion devient un environnement bancaire exploitable, sans intervention manuelle de déploiement ni duplication du socle applicatif. L’abonnement occupe une position centrale : il matérialise les services souscrits et conditionne l’activation de la marketplace ainsi que l’accès au back-office de la banque.')
    add_heading(doc, '4.1 Principes de réalisation et rôle de l’abonnement', 2)
    para(doc, 'Matchia adopte une organisation multi-tenant : le catalogue de stores et de modules est administré au niveau SaaS, alors que chaque banque dispose de sa propre configuration, de son identité visuelle et de ses ressources autorisées. Cette séparation permet d’offrir une marketplace spécifique à une banque tout en préservant un code source unique. La banque ne crée donc pas une nouvelle application ; elle active et paramètre une instance logique de la plateforme.')
    para(doc, 'Dans ce modèle, l’abonnement ne se limite pas à une information comptable. Il établit la liste des stores et modules achetés, la période de validité et l’état d’autorisation des services. Tant que le paiement associé n’est pas confirmé, l’abonnement demeure en attente de paiement, la marketplace reste inactive et l’administrateur Banque ne reçoit pas les droits opérationnels du back-office. Cette dépendance évite l’ouverture anticipée d’un tenant non réglé.')
    callout(doc, 'Principe de contrôle d’accès', 'L’état de l’abonnement est la référence fonctionnelle pour l’activation. Une demande approuvée prépare le parcours commercial ; seule la confirmation de paiement transforme ce parcours en accès effectif aux services de la marketplace.')
    add_figure(doc, DIAGRAMS/'fig_4_2_etats_et_dependances.png', 'Figure 4.1 — Relation entre demande, abonnement, paiement et activation de la marketplace.', 6.25)
    add_heading(doc, '4.2 Back-office SaaS : traitement de la demande d’adhésion', 2)
    para(doc, 'Le back-office SaaS centralise l’instruction des demandes des banques. L’administrateur SaaS y consulte les demandes reçues, contrôle les informations déclarées, la marketplace envisagée, les stores et modules demandés, puis décide de l’approbation ou du rejet. Les indicateurs du tableau de bord facilitent le suivi des banques actives, des demandes en attente, des souscriptions, des alertes d’échéance et des revenus associés.')
    add_figure(doc, ASSETS/'dashboard.png', 'Figure 4.2 — Tableau de bord SaaS : suivi des demandes, abonnements et alertes d’échéance.', 6.25)
    add_placeholder(doc, 'Capture 4.1', 'Liste des demandes d’adhésion bancaires', 'Insérer la vue SaaS présentant les demandes, leur statut et les actions de traitement. Anonymiser les coordonnées de contact.')
    add_placeholder(doc, 'Capture 4.2', 'Détail et approbation d’une demande bancaire', 'Insérer la fenêtre de détail avec les stores, les modules et l’action d’approbation ou de rejet.')
    add_heading(doc, '4.2.1 Approbation et notification système', 3)
    para(doc, 'L’approbation constitue la décision initiale du cycle. Lorsque l’administrateur SaaS valide une demande, la plateforme enregistre ce changement d’état et génère une notification système. Cette notification assure la traçabilité du traitement dans l’interface de pilotage et permet d’informer les parties concernées qu’une étape commerciale est ouverte. Une demande approuvée ne signifie toutefois pas que la marketplace est déjà accessible : elle rend possible la génération de l’abonnement et de l’opération de paiement.')
    para(doc, 'La notification dissocie ainsi deux responsabilités : le SaaS valide la cohérence de la demande, tandis que le règlement confirme l’engagement de la banque pour les services choisis. Cette dissociation est importante pour éviter qu’une validation administrative active par erreur une marketplace avant la conclusion du paiement.')
    add_placeholder(doc, 'Capture 4.3', 'Notification système après approbation', 'Insérer le centre de notifications SaaS indiquant qu’une demande bancaire vient d’être approuvée.')
    add_heading(doc, '4.3 Processus d’abonnement et de paiement', 2)
    para(doc, 'Après approbation, la plateforme crée ou prépare l’abonnement correspondant à la demande. Les stores, modules, prix et paramètres commerciaux sélectionnés sont conservés dans cette souscription afin de garantir la cohérence historique : une modification ultérieure du catalogue SaaS ne doit pas altérer le contenu d’un abonnement déjà émis. À ce stade, l’abonnement est à l’état PENDING_PAYMENT et les services concernés ne sont pas activés.')
    add_heading(doc, '4.3.1 Génération du paiement et envoi du lien sécurisé', 3)
    para(doc, 'Le back-end génère l’opération de paiement auprès du prestataire de paiement et obtient une référence ou une session de paiement. Cette référence est associée à l’abonnement et transmise à l’administrateur Banque par un email contenant un lien sécurisé. L’email guide le destinataire vers la page de transaction sans exposer les clés privées ni la logique de paiement de la plateforme. La notification système et l’email sont complémentaires : la première trace l’événement interne, tandis que le second déclenche l’action attendue du client bancaire.')
    add_placeholder(doc, 'Capture 4.4', 'Email de paiement envoyé à la banque', 'Insérer l’email contenant le lien de paiement après masquage de toute adresse, numéro de téléphone, identifiant de session ou donnée sensible.')
    add_placeholder(doc, 'Capture 4.5', 'Page de paiement sécurisée', 'Insérer l’écran de paiement hébergé ou le formulaire de transaction, sans afficher de clé API ni de token.')
    add_heading(doc, '4.3.2 Réalisation de la transaction et validation', 3)
    para(doc, 'L’administrateur Banque accède au lien, effectue la transaction et reçoit le résultat fourni par le prestataire. La validation ne repose pas uniquement sur le retour visuel du navigateur : le serveur vérifie le statut de paiement reçu, met à jour la transaction et applique la transition de l’abonnement de PENDING_PAYMENT vers ACTIVE uniquement lorsque la confirmation est obtenue. En cas d’échec, d’annulation ou d’absence de confirmation, l’abonnement reste non actif et aucun accès de production n’est ouvert.')
    para(doc, 'Ce traitement peut s’appuyer sur une session de paiement hébergée ou sur un Payment Intent selon le scénario choisi. Dans les deux cas, la confirmation côté serveur est nécessaire. En environnement de production, la réception de webhooks signés et le traitement idempotent des événements renforcent cette garantie : un événement dupliqué ne doit jamais activer deux fois une même souscription.')
    add_placeholder(doc, 'Capture 4.6', 'Statut de paiement et abonnement en attente', 'Insérer la fiche SaaS montrant le paiement et l’abonnement avant confirmation : PENDING_PAYMENT.')
    add_figure(doc, DIAGRAMS/'fig_4_1_sequence_adhesion_activation.png', 'Figure 4.3 — Diagramme de séquence : de la demande d’adhésion à l’accès au back-office Banque.', 6.25)
    add_heading(doc, '4.3.3 Activation de la marketplace et communication des accès', 3)
    para(doc, 'Dès que la confirmation de paiement est validée, la plateforme met à jour de manière synchronisée les objets liés au cycle de souscription : le paiement est marqué comme confirmé, l’abonnement devient actif et la marketplace de la banque passe à l’état ACTIVE. L’activation rend les stores et modules souscrits disponibles dans le périmètre du tenant. Elle constitue donc le point de bascule entre le parcours commercial et le parcours d’exploitation.')
    para(doc, 'Une fois l’activation terminée, un email de bienvenue est envoyé à l’administrateur Banque. Il contient les informations nécessaires à la première connexion au back-office, notamment le login et un mot de passe initial ou un mécanisme sécurisé de définition du mot de passe. Ces informations doivent être communiquées par un canal maîtrisé et ne doivent jamais être conservées en clair dans les journaux ou les bases de données applicatives. À la première connexion, la banque peut modifier son mot de passe et commencer la configuration de son espace.')
    add_placeholder(doc, 'Capture 4.7', 'Marketplace activée et abonnement actif', 'Insérer la fiche de souscription après confirmation, avec l’état ACTIVE et les services concernés.')
    add_placeholder(doc, 'Capture 4.8', 'Email de création des accès Back-office Banque', 'Insérer l’email de bienvenue avec les identifiants masqués ; ne jamais afficher le mot de passe réel dans le rapport.')
    add_heading(doc, '4.4 Back-office Administrateur Banque et prérequis d’accès', 2)
    callout(doc, 'Prérequis d’accès', 'Abonnement valide et marketplace active. L’administrateur Banque ne peut exploiter pleinement son back-office que lorsque les conditions d’abonnement et d’activation de la marketplace sont satisfaites.')
    para(doc, 'Le back-office Banque représente l’espace d’administration spécifique au tenant. Le contrôle de ses droits tient compte du rôle ADMIN_BANK, de l’appartenance à la banque et de l’état de la souscription. Cette logique évite qu’un administrateur disposant d’un compte préparé avant paiement puisse publier des contenus, activer des modules ou gérer les ressources commerciales alors que la marketplace n’est pas encore active.')
    add_figure(doc, ASSETS/'backofficeTenant.png', 'Figure 4.4 — Back-office Banque : tableau de bord tenant et services configurables.', 6.25)
    para(doc, 'Une fois les prérequis remplis, l’administrateur Banque accède à un tableau de bord contextualisé. Il y retrouve les utilisateurs de son organisation, les stores assignés, les modules disponibles, ses demandes, ses produits, ses concessionnaires, ses financements, son abonnement et ses paramètres de branding. Les données affichées sont filtrées par le contexte de la banque afin de préserver l’isolation entre tenants.')
    add_placeholder(doc, 'Capture 4.9', 'Page de connexion au Back-office Banque', 'Insérer l’écran de connexion puis la redirection vers le tableau de bord de la banque, sans afficher d’identifiants réels.')
    add_heading(doc, '4.5 Configuration de la marketplace bancaire', 2)
    para(doc, 'L’activation ne fige pas la marketplace : elle ouvre une phase de paramétrage autonome pour la banque. L’administrateur Banque configure sa vitrine sans modifier le socle commun de la plateforme SaaS. Le catalogue global reste maintenu par le SaaS, tandis que le tenant choisit parmi les stores et modules effectivement souscrits ceux qu’il souhaite exposer à ses visiteurs. La personnalisation porte alors sur le contenu, les éléments visuels, les bannières et les paramètres fonctionnels associés aux modules.')
    add_figure(doc, DIAGRAMS/'fig_4_3_configuration_tenant.png', 'Figure 4.5 — Activité de configuration d’une marketplace bancaire active.', 6.25)
    add_heading(doc, '4.5.1 Contenu, identité visuelle et bannières', 3)
    para(doc, 'La banque définit son identité dans la marketplace à travers le logo, les couleurs, les textes d’accueil, les contenus éditoriaux et les bannières. Ces éléments sont rattachés au tenant et peuvent être adaptés à chaque store. La marketplace publique conserve ainsi une apparence cohérente avec la communication de la banque, tout en reposant sur les mêmes composants techniques que les autres tenants.')
    add_placeholder(doc, 'Capture 4.10', 'Gestion du contenu et des bannières', 'Insérer l’écran de gestion des textes, bannières et contenus de la marketplace bancaire.')
    add_placeholder(doc, 'Capture 4.11', 'Personnalisation visuelle de la marketplace', 'Insérer l’écran de branding : logo, couleurs et éléments de l’identité de la banque.')
    add_heading(doc, '4.5.2 Stores, modules et paramètres fonctionnels', 3)
    para(doc, 'Un store représente une verticale métier, par exemple le véhicule, l’immobilier ou la santé. L’administrateur Banque ne crée pas une nouvelle implémentation du store : il sélectionne, parmi les services souscrits, les stores qu’il souhaite rendre visibles dans sa marketplace. À chaque store peuvent être associés des modules, tels qu’un comparateur ou un simulateur. Les modules peuvent être activés ou désactivés selon la stratégie de la banque, dans la limite du périmètre contractuel de l’abonnement.')
    para(doc, 'Certains modules disposent de paramètres propres : valeurs de simulation, champs affichés, règles de présentation, contenus ou comportements autorisés. Ces paramètres sont conservés au niveau de l’association entre la marketplace, le store et le module. Cette approche évite un modèle rigide où chaque type de produit imposerait une table dédiée ; elle rend également possible l’évolution progressive de l’offre sans modifier le code commun pour chaque banque.')
    add_placeholder(doc, 'Capture 4.12', 'Stores disponibles dans la marketplace', 'Insérer la liste des stores assignés au tenant et leur état de visibilité.')
    add_placeholder(doc, 'Capture 4.13', 'Activation des modules par store', 'Insérer la liste des modules associés à un store, avec les actions d’activation ou de désactivation.')
    add_placeholder(doc, 'Capture 4.14', 'Paramètres d’un module', 'Insérer un exemple de paramètres spécifiques à un simulateur ou à un comparateur.')
    add_heading(doc, '4.6 Renouvellement et évolution de l’abonnement', 2)
    para(doc, 'Le cycle d’abonnement ne s’arrête pas à l’activation initiale. Une banque peut souhaiter prolonger son abonnement, ajouter un store, ajouter de nouveaux modules ou faire évoluer son offre vers une formule supérieure. Ces demandes ne modifient pas directement les services actifs : elles sont soumises au même mécanisme de contrôle que la souscription initiale afin de maintenir une correspondance exacte entre les droits du tenant et les services payés.')
    para(doc, 'La banque formule une demande de renouvellement ou d’upgrade depuis son espace. L’administrateur SaaS l’examine, valide les services concernés et prépare le montant complémentaire ou la période de prolongation. Une nouvelle opération de paiement est alors générée. Après confirmation, l’abonnement est mis à jour : sa date de fin est prolongée dans le cas d’un renouvellement, ou ses associations de stores et modules sont complétées dans le cas d’un upgrade. Les seuls services nouvellement validés deviennent alors actifs.')
    add_figure(doc, DIAGRAMS/'fig_4_4_renouvellement_upgrade.png', 'Figure 4.6 — Workflow de renouvellement ou d’upgrade d’un abonnement bancaire.', 6.25)
    add_placeholder(doc, 'Capture 4.15', 'Demande de renouvellement ou d’upgrade', 'Insérer l’écran Banque de création d’une demande d’évolution : nouveau store, module ou renouvellement.')
    add_placeholder(doc, 'Capture 4.16', 'Validation SaaS et paiement complémentaire', 'Insérer la vue SaaS de validation, puis la référence de paiement associée à l’évolution demandée.')
    add_placeholder(doc, 'Capture 4.17', 'Abonnement mis à jour après confirmation', 'Insérer l’abonnement après mise à jour des dates, stores ou modules nouvellement activés.')
    add_heading(doc, '4.7 Traçabilité, sécurité et continuité des traitements', 2)
    para(doc, 'La robustesse du processus repose sur la persistance des états et sur leur traçabilité. Chaque étape doit laisser une information exploitable par le SaaS : demande reçue, décision d’approbation, notification émise, email de paiement envoyé, paiement en attente, confirmation reçue, abonnement activé et compte Banque communiqué. Cette chronologie facilite le support, l’audit et la résolution des incidents, notamment lorsqu’une transaction est abandonnée ou qu’un paiement est confirmé de manière asynchrone.')
    para(doc, 'Les notifications et les emails assurent l’automatisation de la communication, alors que les contrôles de statut assurent la sécurité fonctionnelle. Les secrets de paiement, identifiants de session et mots de passe initiaux ne doivent pas apparaître dans les captures du mémoire ni dans les logs applicatifs. Cette séparation entre automatisation et protection est essentielle dans le contexte bancaire : la plateforme accélère l’onboarding sans réduire les garanties de contrôle des accès.')
    add_heading(doc, '4.8 Conclusion', 2)
    para(doc, 'La réalisation du chapitre démontre un cycle cohérent allant de la demande d’adhésion à l’exploitation d’une marketplace bancaire configurée. L’approbation SaaS déclenche les notifications et la préparation du paiement ; la confirmation du paiement active l’abonnement, la marketplace et les accès du back-office Banque. Une fois active, la banque personnalise ses contenus, ses stores et ses modules sans modifier le socle SaaS. Enfin, le renouvellement et les upgrades appliquent le même principe de validation et de paiement avant toute extension des services, ce qui maintient l’abonnement au centre du contrôle d’accès et de l’évolution de la plateforme.')

def main():
    generate_diagrams()
    doc = Document(); setup(doc); cover(doc); write_chapter(doc)
    doc.core_properties.title = 'Chapitre 4 - Matchia : abonnement, paiement et activation'
    doc.core_properties.subject = 'Réalisation du SaaS bancaire et configuration multi-tenant'
    doc.core_properties.author = 'Matchia'
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    main()
