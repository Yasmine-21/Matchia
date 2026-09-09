from pathlib import Path

from docx import Document

from reformat_agile_release_sprint_alignment import set_cell, set_table_rows


INPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - planification Agile alignée.docx")
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - backlog chapitre 4 détaillé.docx")


def main():
    document = Document(INPUT)
    release1 = document.tables[1]
    correspondence = document.tables[3]

    # The release 1 backlog mirrors the actual functional sections of chapter 4.
    set_table_rows(release1, [
        (
            "S1",
            "Socle SaaS multi-tenant",
            "Authentification, gestion des sessions, autorisation par rôles et application du contexte multi-tenant afin d’isoler les données de chaque banque.",
        ),
        (
            "S2",
            "Demande de création de marketplace",
            "Informations de la banque, futur Administrateur Banque, vérification de l’e-mail, slug, identité visuelle, stores, modules, récapitulatif et soumission.",
        ),
        (
            "S3",
            "Back-office SaaS",
            "Tableau de bord, stores et modules, banques, utilisateurs et rôles, concessionnaires, marketplaces, contenus, offres, abonnements, audit, logs et paramètres.",
        ),
        (
            "S4",
            "Traitement, paiement et activation",
            "Consultation des demandes, approbation ou rejet avec motif, e-mails, lien Stripe, contrôle du paiement, activation de la banque, de la marketplace, de l’abonnement et du compte Administrateur Banque.",
        ),
        (
            "S5",
            "Back-office Banque et marketplace publique",
            "Personnalisation et contenus, configuration des stores et modules, renouvellement ou ajout de services, accueil public, produits, comparateur, simulateur, chatbot et demande de financement.",
        ),
    ])

    # Reinforce the same link in the correspondence table for the chapter 4 entries.
    replacements = [
        ("Chapitre 4 — Authentification, autorisation et contexte multi-tenant", "S1", "Accès sécurisé, rôles et cloisonnement des données par banque."),
        ("Chapitre 4 — Demande de création de marketplace bancaire", "S2", "Saisie, vérification e-mail, configuration, stores, modules et soumission."),
        ("Chapitre 4 — Back-office SaaS : administration et supervision", "S3", "Tableau de bord et gestion des acteurs, catalogues, contenus, abonnements, audit et paramètres."),
        ("Chapitre 4 — Traitement, approbation et paiement", "S4", "Décision, e-mails, paiement Stripe et activation conditionnelle des services."),
        ("Chapitre 4 — Back-office Banque et marketplace publique", "S5", "Configuration bancaire, services souscrits et parcours public jusqu’à la demande de financement."),
    ]
    for row, values in zip(correspondence.rows[1:6], replacements):
        for col, value in enumerate(values):
            set_cell(row.cells[col], value, correspondence.rows[0].cells[col])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
