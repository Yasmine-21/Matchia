from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\PFE M2\Platforme SaaS")
SOURCE = Path(r"D:\PFE - BRI Technology\Rapport\Master Report.docx")
OUTPUT = ROOT / "deliverables" / "Master Report - diagrammes UML complets corriges.docx"
ACTIVITY = ROOT / "document_work" / "diagrammes_activite_modele"
SEQUENCE = ROOT / "document_work" / "diagrammes_sequence_matchia"


def normalize(text: str) -> str:
    return " ".join(text.split())


def find_paragraph(doc: Document, *, startswith: str | None = None, contains: str | None = None):
    for paragraph in doc.paragraphs:
        text = normalize(paragraph.text)
        if startswith is not None and text.startswith(startswith):
            return paragraph
        if contains is not None and contains in text:
            return paragraph
    raise ValueError(f"Anchor not found: startswith={startswith!r}, contains={contains!r}")


def set_font(run, name="Times New Roman", size=9):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)


def style_intro(paragraph):
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        set_font(run, size=11)


def style_caption(paragraph, text: str | None = None):
    if text is not None:
        paragraph.clear()
        paragraph.add_run(text)
    paragraph.style = "Caption"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    for run in paragraph.runs:
        set_font(run, size=9)
        run.italic = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def add_intro_before(anchor, text: str):
    paragraph = anchor.insert_paragraph_before(text)
    style_intro(paragraph)
    return paragraph


def add_picture_before(anchor, image: Path, alt_text: str, width_inches: float):
    if not image.exists():
        raise FileNotFoundError(image)
    paragraph = anchor.insert_paragraph_before()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    shape = paragraph.add_run().add_picture(str(image), width=Inches(width_inches))
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)
    return paragraph


def add_caption_before(anchor, text: str):
    paragraph = anchor.insert_paragraph_before(text)
    style_caption(paragraph)
    return paragraph


def insert_new_figure(anchor, image: Path, caption: str, alt_text: str, width: float, intro: str | None = None):
    if intro:
        add_intro_before(anchor, intro)
    add_picture_before(anchor, image, alt_text, width)
    return add_caption_before(anchor, caption)


def insert_at_existing_caption(caption, image: Path, new_caption: str, alt_text: str, width: float):
    add_picture_before(caption, image, alt_text, width)
    style_caption(caption, new_caption)


def remove_nearby_picture_before(anchor, limit=8):
    current = anchor._p.getprevious()
    inspected = 0
    while current is not None and inspected < limit:
        if current.xpath(".//w:drawing | .//w:pict"):
            current.getparent().remove(current)
            return True
        current = current.getprevious()
        inspected += 1
    return False


def replace_caption_number(paragraph, chapter: int, number: int):
    text = normalize(paragraph.text)
    updated = re.sub(rf"^Figure\s+{chapter}\.(?:\d+|[xX])", f"Figure {chapter}.{number}", text)
    if updated == text:
        return
    style_caption(paragraph, updated)


def renumber_chapter_figures(doc: Document, chapter: int, start_heading: str, end_heading: str | None):
    active = False
    number = 0
    for paragraph in doc.paragraphs:
        text = normalize(paragraph.text)
        if text.startswith(start_heading):
            active = True
            continue
        if active and end_heading and text.startswith(end_heading):
            break
        if active and re.match(rf"^Figure\s+{chapter}\.(?:\d+|[xX])", text):
            number += 1
            replace_caption_number(paragraph, chapter, number)
    return number


def clean_report_placeholders(doc: Document):
    in_chapter5 = False
    for paragraph in doc.paragraphs:
        text = normalize(paragraph.text)
        if text.startswith("Chapitre 5"):
            in_chapter5 = True
        elif text.startswith("Chapitre 6"):
            in_chapter5 = False
        if "diagramme d’activité(NAJA)" in text:
            paragraph.clear()
            paragraph.add_run("Le processus décisionnel peut également être représenté à travers le diagramme d’activité suivant.")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if in_chapter5 and text == "5. Introduction":
            paragraph.clear()
            paragraph.add_run("5. Conclusion")


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    doc = Document(SOURCE)
    source_paragraphs = len(doc.paragraphs)
    source_tables = len(doc.tables)
    source_shapes = len(doc.inline_shapes)

    # Chapitre 4 - authentification et contexte multi-tenant.
    creation_section = find_paragraph(doc, startswith="3. Demande de création d’une marketplace bancaire")
    insert_new_figure(
        creation_section,
        ACTIVITY / "AD-01_authentification.png",
        "Figure 4.1 — Diagramme d’activité de l’authentification et du contexte d’accès",
        "Diagramme d’activité UML de l’authentification, de l’autorisation et de la résolution du contexte tenant.",
        5.25,
        "Le diagramme d’activité suivant synthétise le contrôle des identifiants et de l’état du compte, le chargement du rôle et des autorisations ainsi que la détermination du contexte du tenant.",
    )
    insert_new_figure(
        creation_section,
        SEQUENCE / "SD-01_authentification.png",
        "Figure 4.2 — Diagramme de séquence de l’authentification et du contexte multi-tenant",
        "Diagramme de séquence UML de l’authentification et de la résolution conditionnelle du contexte multi-tenant.",
        6.0,
        "Le diagramme de séquence précise les échanges entre l’interface, le service d’authentification, la génération des jetons et la base de données.",
    )

    # Chapitre 4 - création de la demande. Remplacement de l'ancien visuel.
    creation_details = find_paragraph(doc, startswith="3.1. Renseignement des informations de la banque")
    if not remove_nearby_picture_before(creation_details):
        raise RuntimeError("Le diagramme d'activité existant de création de marketplace est introuvable")
    insert_new_figure(
        creation_details,
        ACTIVITY / "AD-02_creation_demande_marketplace.png",
        "Figure 4.3 — Diagramme d’activité de création d’une demande de marketplace bancaire",
        "Diagramme d’activité UML du formulaire progressif de création d’une marketplace bancaire.",
        5.25,
    )
    insert_new_figure(
        creation_details,
        SEQUENCE / "SD-02_creation_demande_marketplace.png",
        "Figure 4.4 — Diagramme de séquence de la création d’une demande de marketplace bancaire",
        "Diagramme de séquence UML de la vérification de l’e-mail, de la configuration et de l’enregistrement de la demande.",
        5.9,
        "Le diagramme de séquence complète ce parcours en présentant la vérification de l’adresse e-mail, la validation de la configuration, l’enregistrement et les notifications.",
    )

    # Chapitre 4 - traitement, paiement Stripe et activation.
    treatment_activity_caption = find_paragraph(doc, contains="Diagramme d’activité du traitement d’une demande")
    insert_at_existing_caption(
        treatment_activity_caption,
        ACTIVITY / "AD-03_traitement_demande_marketplace.png",
        "Figure 4.13 — Diagramme d’activité du traitement d’une demande et du paiement Stripe",
        "Diagramme d’activité UML du traitement de la demande, du paiement Stripe et de l’activation de la marketplace.",
        5.25,
    )
    treatment_sequence_caption = find_paragraph(doc, contains="Diagramme de séquence de l’approbation, du paiement et de l’activation")
    insert_at_existing_caption(
        treatment_sequence_caption,
        SEQUENCE / "SD-03_traitement_demande_stripe.png",
        "Figure 4.14 — Diagramme de séquence de l’approbation, du paiement et de l’activation",
        "Diagramme de séquence UML de l’approbation, de la Checkout Session Stripe, du webhook et de l’activation.",
        5.75,
    )

    # Chapitre 4 - renouvellement et extension de l'abonnement.
    renewal_caption = find_paragraph(doc, contains="Diagramme d’activité du renouvellement et de l’extension d’un abonnement")
    insert_at_existing_caption(
        renewal_caption,
        ACTIVITY / "AD-04_renouvellement_extension_abonnement.png",
        "Figure 4.31 — Diagramme d’activité du renouvellement et de l’extension d’un abonnement",
        "Diagramme d’activité UML du renouvellement ou de l’extension d’un abonnement avec validation et paiement.",
        5.25,
    )

    # Chapitre 5 - parcours du concessionnaire.
    dealer_activity_caption = find_paragraph(doc, contains="Diagramme d’activité du parcours d’inscription et de validation d’un concessionnaire")
    insert_at_existing_caption(
        dealer_activity_caption,
        ACTIVITY / "AD-05_parcours_concessionnaire.png",
        "Figure 5.1 — Diagramme d’activité du parcours métier du Concessionnaire",
        "Diagramme d’activité UML du parcours métier du Concessionnaire.",
        5.15,
    )
    dealer_capture = find_paragraph(doc, startswith="Capture 5.1 — Formulaire d’inscription du concessionnaire")
    insert_new_figure(
        dealer_capture,
        SEQUENCE / "SD-04_inscription_concessionnaire.png",
        "Figure 5.2 — Diagramme de séquence de l’inscription et de la validation d’un concessionnaire",
        "Diagramme de séquence UML de l’inscription, de l’approbation ou du rejet d’un concessionnaire.",
        5.9,
        "Le diagramme de séquence suivant précise la validation du dossier, sa transmission à l’Administrateur SaaS et les deux issues du traitement.",
    )

    # Chapitre 5 - produit central et publication multi-banque.
    product_capture = find_paragraph(doc, startswith="Capture 5.4 — Création d’un produit et gestion du stock")
    insert_new_figure(
        product_capture,
        SEQUENCE / "SD-05_gestion_produit.png",
        "Figure 5.4 — Diagramme de séquence de la gestion et de la publication multi-banque d’un produit",
        "Diagramme de séquence UML de la création, du stock et de la publication bancaire d’un produit.",
        5.9,
        "Le diagramme de séquence distingue la gestion du produit central, le stockage de ses médias et la publication propre à chaque banque partenaire.",
    )

    # Chapitre 5 - inscription du client.
    client_registration_capture = find_paragraph(doc, startswith="Capture 5.7 — Inscription et vérification du compte client")
    insert_new_figure(
        client_registration_capture,
        SEQUENCE / "SD-06_inscription_client.png",
        "Figure 5.5 — Diagramme de séquence de l’inscription et de l’activation du compte Client",
        "Diagramme de séquence UML de l’inscription, de la vérification e-mail et de l’activation du Client dans son tenant.",
        5.9,
        "Le diagramme de séquence détaille la création initiale du compte inactif, l’envoi du code à six chiffres et l’activation du Client dans la marketplace consultée.",
    )

    # Chapitre 5 - parcours et financement du client.
    client_activity_caption = find_paragraph(doc, contains="Diagramme d’activité du parcours de demande de financement")
    insert_at_existing_caption(
        client_activity_caption,
        ACTIVITY / "AD-06_parcours_client.png",
        "Figure 5.6 — Diagramme d’activité du parcours Client et de la demande de financement",
        "Diagramme d’activité UML du parcours Client, de la simulation au traitement bancaire.",
        5.15,
    )
    financing_sequence_caption = find_paragraph(doc, contains="Diagramme de séquence du traitement d’une demande de financement")
    insert_at_existing_caption(
        financing_sequence_caption,
        SEQUENCE / "SD-07_demande_financement.png",
        "Figure 5.7 — Diagramme de séquence de la constitution et du traitement d’une demande de financement",
        "Diagramme de séquence UML de la création du dossier, des justificatifs et de la décision bancaire.",
        5.75,
    )

    # Chapitre 5 - assistants conversationnels.
    assistant_caption = find_paragraph(doc, contains="Diagramme de séquence de l’assistant IA Text-to-SQL sécurisé")
    insert_at_existing_caption(
        assistant_caption,
        SEQUENCE / "SD-08_assistant_saas.png",
        "Figure 5.8 — Diagramme de séquence de l’assistant IA Text-to-SQL sécurisé",
        "Diagramme de séquence UML de l’assistant SaaS avec schéma filtré, validation et lecture seule.",
        5.9,
    )
    chatbot_caption = find_paragraph(doc, contains="Fonctionnement du chatbot public dans une marketplace")
    insert_at_existing_caption(
        chatbot_caption,
        SEQUENCE / "SD-09_chatbot_marketplace.png",
        "Figure 5.9 — Diagramme de séquence du chatbot public d’une marketplace bancaire",
        "Diagramme de séquence UML du chatbot public, du contrôle du module et de la détection d’intention.",
        5.9,
    )

    clean_report_placeholders(doc)
    chapter4_figures = renumber_chapter_figures(doc, 4, "Chapitre 4", "Chapitre 5")
    chapter5_figures = renumber_chapter_figures(doc, 5, "Chapitre 5", "Chapitre 6")

    doc.core_properties.title = "Master Report Matchia avec diagrammes UML complets"
    doc.core_properties.subject = "Rapport complet avec diagrammes d’activité et de séquence UML"
    doc.core_properties.comments = "Diagrammes UML harmonisés et insérés aux emplacements fonctionnels correspondants."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)

    check = Document(OUTPUT)
    if len(check.inline_shapes) != source_shapes + 14:
        # 15 diagrammes ont été insérés et l'ancien diagramme de création a été remplacé.
        raise RuntimeError(f"Nombre d'images inattendu: {len(check.inline_shapes)} au lieu de {source_shapes + 14}")
    if len(check.tables) != source_tables:
        raise RuntimeError(f"Les tableaux ont changé: {source_tables} -> {len(check.tables)}")
    expected = [
        "Diagramme d’activité de l’authentification",
        "Diagramme de séquence de l’authentification",
        "Diagramme d’activité de création",
        "Diagramme de séquence de la création",
        "Diagramme d’activité du traitement",
        "Diagramme de séquence de l’approbation",
        "Diagramme d’activité du renouvellement",
        "Diagramme d’activité du parcours métier du Concessionnaire",
        "Diagramme de séquence de l’inscription et de la validation d’un concessionnaire",
        "Diagramme de séquence de la gestion et de la publication multi-banque",
        "Diagramme de séquence de l’inscription et de l’activation du compte Client",
        "Diagramme d’activité du parcours Client",
        "Diagramme de séquence de la constitution et du traitement d’une demande de financement",
        "Diagramme de séquence de l’assistant IA",
        "Diagramme de séquence du chatbot public",
    ]
    texts = [normalize(p.text) for p in check.paragraphs]
    missing = [item for item in expected if not any(item in text for text in texts)]
    if missing:
        raise RuntimeError(f"Légendes manquantes: {missing}")
    print(f"OUTPUT={OUTPUT}")
    print(f"PARAGRAPHS={source_paragraphs}->{len(check.paragraphs)} TABLES={len(check.tables)} SHAPES={source_shapes}->{len(check.inline_shapes)}")
    print(f"FIGURES_CH4={chapter4_figures} FIGURES_CH5={chapter5_figures}")


if __name__ == "__main__":
    main()
