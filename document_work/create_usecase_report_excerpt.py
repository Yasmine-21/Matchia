from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.oxml.ns import qn


DELIVERABLES = Path(r"D:\PFE M2\Platforme SaaS\deliverables")
SOURCE = next(DELIVERABLES.glob("Master Report - diagrammes use case Draw.io *.docx"))
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\rendered_usecases_report\report_usecase_excerpt.docx")


def element_text(element) -> str:
    return "".join(element.itertext()).replace("\xa0", " ")


def main() -> None:
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    body = doc._element.body
    children = list(body)

    start_candidates = []
    end = None
    for index, child in enumerate(children):
        text = " ".join(element_text(child).split())
        if "6. Spécification détaillée des cas d’utilisation" in text:
            start_candidates.append(index)
        if "Tableau 3.9" in text and "UC-06" in text:
            end = index

    if end is not None:
        for index in range(end + 1, len(children)):
            if children[index].tag == qn("w:tbl"):
                end = index
                break

    eligible_starts = [index for index in start_candidates if end is not None and index < end]
    start = max(eligible_starts) if eligible_starts else None
    if start is None or end is None or end < start:
        raise RuntimeError(f"Unable to identify excerpt range: start={start}, end={end}")

    for index, child in reversed(list(enumerate(children))):
        if child.tag == qn("w:sectPr"):
            continue
        if index < start or index > end:
            body.remove(child)

    doc.save(OUTPUT)
    check = Document(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Paragraphs: {len(check.paragraphs)}")
    print(f"Tables: {len(check.tables)}")
    print(f"Inline shapes: {len(check.inline_shapes)}")


if __name__ == "__main__":
    main()
