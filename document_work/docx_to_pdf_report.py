from pathlib import Path
from html import escape
from xml.etree import ElementTree as ET

from docx import Document
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(r"D:\PFE M2\Platforme SaaS")
SOURCE = ROOT / "document_work" / "Chapitres_1_2_6_Matchia.docx"
OUT = ROOT / "output" / "pdf" / "Rapport_Matchia_Chapitres_1_2_6.pdf"
TMP = ROOT / "tmp" / "pdfs" / "matchia_images"


def register_fonts():
    regular = Path(r"C:\Windows\Fonts\arial.ttf")
    bold = Path(r"C:\Windows\Fonts\arialbd.ttf")
    italic = Path(r"C:\Windows\Fonts\ariali.ttf")
    bold_italic = Path(r"C:\Windows\Fonts\arialbi.ttf")
    if all(p.exists() for p in (regular, bold, italic, bold_italic)):
        pdfmetrics.registerFont(TTFont("Arial", str(regular)))
        pdfmetrics.registerFont(TTFont("Arial-Bold", str(bold)))
        pdfmetrics.registerFont(TTFont("Arial-Italic", str(italic)))
        pdfmetrics.registerFont(TTFont("Arial-BoldItalic", str(bold_italic)))
        return "Arial", "Arial-Bold", "Arial-Italic", "Arial-BoldItalic"
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Helvetica-BoldOblique"


REGULAR, BOLD, ITALIC, BOLD_ITALIC = register_fonts()


def make_styles():
    styles = getSampleStyleSheet()
    return {
        "normal": ParagraphStyle("MatchiaNormal", parent=styles["BodyText"], fontName=REGULAR, fontSize=10.1,
                                 leading=14.2, alignment=TA_JUSTIFY, spaceAfter=7),
        "h1": ParagraphStyle("MatchiaH1", parent=styles["Heading1"], fontName=BOLD, fontSize=16,
                              leading=20, textColor=colors.HexColor("#2E74B5"), spaceBefore=15, spaceAfter=10,
                              keepWithNext=True),
        "h2": ParagraphStyle("MatchiaH2", parent=styles["Heading2"], fontName=BOLD, fontSize=13,
                              leading=16, textColor=colors.HexColor("#2E74B5"), spaceBefore=12, spaceAfter=7,
                              keepWithNext=True),
        "h3": ParagraphStyle("MatchiaH3", parent=styles["Heading3"], fontName=BOLD, fontSize=11.3,
                              leading=14, textColor=colors.HexColor("#1F4D78"), spaceBefore=9, spaceAfter=5,
                              keepWithNext=True),
        "caption": ParagraphStyle("MatchiaCaption", parent=styles["BodyText"], fontName=ITALIC, fontSize=8.7,
                                   leading=11, alignment=TA_CENTER, textColor=colors.HexColor("#666666"),
                                   spaceBefore=3, spaceAfter=8),
        "cover": ParagraphStyle("MatchiaCover", parent=styles["BodyText"], fontName=BOLD, fontSize=20,
                                 leading=25, alignment=TA_CENTER, textColor=colors.HexColor("#0B2545"),
                                 spaceAfter=12),
    }


def cell_text(cell):
    return "<br/>".join(escape(p.text).replace("\n", "<br/>") for p in cell.paragraphs if p.text.strip()) or " "


def table_flowable(table, styles):
    rows = []
    for row in table.rows:
        rows.append([Paragraph(cell_text(cell), ParagraphStyle("Cell", parent=styles["normal"], fontSize=7.6,
                                                                 leading=9.4, spaceAfter=0, alignment=TA_LEFT))
                     for cell in row.cells])
    if not rows:
        return Spacer(1, 1)
    width = 6.45 * inch
    lengths = [max(len(cell.text) for cell in col) for col in zip(*(row.cells for row in table.rows))]
    raw = [max(1, min(v, 80)) for v in lengths]
    total = sum(raw)
    widths = [width * v / total for v in raw]
    tbl = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F4F6F9")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1F4D78")),
        ("FONTNAME", (0, 0), (-1, 0), BOLD),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B9C4D0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return tbl


def extract_images(paragraph, document, index):
    images = []
    for blip in paragraph._p.xpath(".//a:blip"):
        rid = blip.get(qn("r:embed"))
        if not rid or rid not in document.part.related_parts:
            continue
        part = document.part.related_parts[rid]
        suffix = ".png"
        if part.content_type == "image/jpeg":
            suffix = ".jpg"
        elif part.content_type == "image/gif":
            suffix = ".gif"
        path = TMP / f"figure_{index}{suffix}"
        path.write_bytes(part.blob)
        images.append(path)
    return images


def image_flowable(path):
    reader = ImageReader(str(path))
    width, height = reader.getSize()
    max_width, max_height = 6.2 * inch, 4.7 * inch
    scale = min(max_width / width, max_height / height, 1.0)
    return Image(str(path), width=width * scale, height=height * scale, hAlign="CENTER")


def ordered_blocks(document):
    body = document.element.body
    table_map = {id(t._element): t for t in document.tables}
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield "paragraph", DocxParagraph(child, document._body)
        elif child.tag == qn("w:tbl"):
            yield "table", table_map.get(id(child), DocxTable(child, document._body))


def draw_page(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D8E1EA"))
    canvas.line(0.7 * inch, 0.57 * inch, 7.8 * inch, 0.57 * inch)
    canvas.setFont(REGULAR, 8)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawString(0.7 * inch, 0.38 * inch, "Matchia - Rapport PFE")
    canvas.drawRightString(7.8 * inch, 0.38 * inch, f"Page {doc.page}")
    canvas.restoreState()


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    TMP.mkdir(parents=True, exist_ok=True)
    source = Document(SOURCE)
    styles = make_styles()
    story = []
    image_index = 0
    first_h1 = True
    for kind, block in ordered_blocks(source):
        if kind == "table":
            story.extend([table_flowable(block, styles), Spacer(1, 8)])
            continue
        text = block.text.strip()
        images = extract_images(block, source, image_index)
        image_index += len(images)
        for path in images:
            story.extend([image_flowable(path), Spacer(1, 4)])
        if not text:
            continue
        style_name = block.style.name if block.style else "Normal"
        if style_name == "Heading 1":
            if not first_h1 and text.startswith("Chapitre"):
                story.append(PageBreak())
            first_h1 = False
            style = styles["h1"]
        elif style_name == "Heading 2":
            style = styles["h2"]
        elif style_name == "Heading 3":
            style = styles["h3"]
        elif text.startswith(("Figure ", "Tableau ")):
            style = styles["caption"]
        elif len(text) > 80 and "Conception, développement et déploiement" in text:
            style = styles["cover"]
        else:
            style = styles["normal"]
        story.append(Paragraph(escape(text).replace("\n", "<br/>"), style))
    pdf = SimpleDocTemplate(str(OUT), pagesize=letter, leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                            topMargin=0.68 * inch, bottomMargin=0.75 * inch, title="Rapport Matchia")
    pdf.build(story, onFirstPage=draw_page, onLaterPages=draw_page)
    print(OUT)


if __name__ == "__main__":
    main()
