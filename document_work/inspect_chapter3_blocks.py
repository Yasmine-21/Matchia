from docx import Document
from docx.oxml.ns import qn


path = r"D:\PFE - BRI Technology\Rapport\Master Report.docx"
doc = Document(path)
paragraph_index = {id(p._p): i for i, p in enumerate(doc.paragraphs)}
table_index = {id(t._tbl): i for i, t in enumerate(doc.tables)}

inside = False
for child_index, child in enumerate(doc.element.body.iterchildren()):
    if child.tag == qn("w:p"):
        i = paragraph_index.get(id(child))
        text = "".join(child.itertext()).strip()
        if text.startswith("Chapitre 3"):
            inside = True
        if inside:
            sect = child.find(".//" + qn("w:sectPr")) is not None
            print(f"B{child_index:03d} P{i:03d} sect={sect} text={text[:140]!r}")
        if text.startswith("Chapitre 4"):
            break
    elif child.tag == qn("w:tbl") and inside:
        ti = table_index.get(id(child))
        rows = []
        for row in child.findall(qn("w:tr"))[:3]:
            cells = []
            for cell in row.findall(qn("w:tc"))[:4]:
                cells.append(" ".join("".join(cell.itertext()).split()))
            rows.append(" | ".join(cells))
        print(f"B{child_index:03d} T{ti:03d} rows={len(child.findall(qn('w:tr')))} sample={rows!r}")
