from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


INPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - chapitre 6 DevOps et Azure.docx")
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - planification Agile alignée.docx")


def copy_ppr(target, source):
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if source._p.pPr is not None:
        target._p.insert(0, deepcopy(source._p.pPr))


def copy_run_format(target_run, source_run):
    if source_run is not None and source_run._r.rPr is not None:
        target_run._r.insert(0, deepcopy(source_run._r.rPr))


def set_paragraph(paragraph, text, template=None):
    if template is not None:
        copy_ppr(paragraph, template)
    ref = next((r for r in paragraph.runs if r.text), None)
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    copy_run_format(run, ref if ref is not None else (next((r for r in template.runs if r.text), None) if template else None))
    return paragraph


def set_cell(cell, text, template_cell=None, bold=False):
    paragraph = cell.paragraphs[0]
    template_paragraph = template_cell.paragraphs[0] if template_cell else None
    set_paragraph(paragraph, text, template_paragraph)
    if bold:
        for run in paragraph.runs:
            run.bold = True
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def copy_cell_width(target, source):
    src_tcpr = source._tc.get_or_add_tcPr()
    src_w = src_tcpr.find(qn('w:tcW'))
    if src_w is not None:
        trg_tcpr = target._tc.get_or_add_tcPr()
        old = trg_tcpr.find(qn('w:tcW'))
        if old is not None:
            trg_tcpr.remove(old)
        trg_tcpr.append(deepcopy(src_w))


def format_table(table, template_table, widths):
    table.style = template_table.style
    table.autofit = False
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            copy_cell_width(cell, template_table.rows[0].cells[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_rows(table, rows):
    while len(table.rows) > 1:
        row = table.rows[-1]
        row._element.getparent().remove(row._element)
    for row_values in rows:
        row = table.add_row()
        for col, value in enumerate(row_values):
            set_cell(row.cells[col], value, table.rows[0].cells[col])


def add_table_before(document, anchor, template_table, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table._element.getparent().remove(table._element)
    anchor._p.addprevious(table._element)
    format_table(table, template_table, widths)
    for col, value in enumerate(headers):
        set_cell(table.rows[0].cells[col], value, template_table.rows[0].cells[col], bold=True)
        table.rows[0].cells[col].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for values in rows:
        row = table.add_row()
        for col, value in enumerate(values):
            set_cell(row.cells[col], value, template_table.rows[0].cells[col])
            row.cells[col].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if col == 1 else WD_ALIGN_PARAGRAPH.LEFT
    return table


def insert_before(anchor, text, template, alignment=None):
    paragraph = anchor.insert_paragraph_before()
    set_paragraph(paragraph, text, template)
    if alignment is not None:
        paragraph.alignment = alignment
    return paragraph


def main():
    document = Document(INPUT)
    paragraphs = document.paragraphs
    p_intro = next(p for p in paragraphs if p.text.startswith("Afin d’assurer une réalisation progressive"))
    p_release_expl = next(p for p in paragraphs if p.text.startswith("La planification retenue est organisée"))
    p_release1 = next(p for p in paragraphs if p.text.startswith("Release 1"))
    p_release1_desc = next(p for p in paragraphs if p.text.startswith("Cette release vise à rendre opérationnel"))
    p_release2 = next(p for p in paragraphs if p.text.startswith("Release 2"))
    p_release2_desc = next(p for p in paragraphs if p.text.startswith("Cette release couvre les fonctionnalités"))
    p_conclusion = next(p for p in paragraphs if p.text.startswith("À l’issue de chaque sprint"))

    table_release1 = document.tables[1]
    table_release2 = document.tables[2]

    set_paragraph(
        p_intro,
        "La planification du projet s’appuie sur l’approche Agile Scrum. Le Product Backlog est organisé par incréments fonctionnels cohérents avec la structure des chapitres de réalisation. Chaque sprint, prévu sur deux semaines, produit un ensemble de fonctionnalités testables et directement rattachées à une section ou à un sous-titre du rapport.",
    )
    set_paragraph(
        p_release_expl,
        "Deux releases structurent cette planification. La première couvre la réalisation du socle SaaS et le cycle de vie complet de la marketplace bancaire. La seconde regroupe les parcours des concessionnaires et des clients, ainsi que les assistants conversationnels. Selon l’ampleur fonctionnelle, un titre de chapitre peut être réalisé par plusieurs sprints, tandis qu’un sous-titre ciblé peut correspondre à un sprint unique.",
    )

    set_paragraph(p_release1, "Release 1 — SaaS bancaire, onboarding et activation de la marketplace")
    set_paragraph(
        p_release1_desc,
        "Cette release correspond aux principales sections du chapitre 4. Elle permet de livrer progressivement le socle multi-tenant, la demande de création de marketplace, l’administration SaaS, le paiement puis les espaces bancaire et public.",
    )
    set_table_rows(table_release1, [
        ("S1", "Socle SaaS multi-tenant", "Authentification, rôles, gestion des sessions, isolation des données et mécanismes de sécurité communs."),
        ("S2", "Demande de création de marketplace", "Informations de la banque et de l’administrateur, vérification e-mail, configuration, stores, modules et soumission."),
        ("S3", "Back-office SaaS", "Tableau de bord, gestion des banques, utilisateurs, stores, modules, offres, abonnements et paramètres."),
        ("S4", "Traitement, paiement et activation", "Analyse des demandes, approbation ou rejet, e-mails, paiement Stripe et activation conditionnelle des services."),
        ("S5", "Back-office Banque et marketplace publique", "Personnalisation, contenus, configuration des stores et modules, abonnements, catalogue public, comparateur et simulateur."),
    ])

    set_paragraph(p_release2, "Release 2 — Parcours des concessionnaires, des clients et assistants conversationnels")
    set_paragraph(
        p_release2_desc,
        "Cette release correspond aux sections du chapitre 5. Elle traite les parcours métier qui complètent l’exploitation de la marketplace, depuis l’inscription des concessionnaires jusqu’au dépôt et au traitement des demandes de financement des clients.",
    )
    set_table_rows(table_release2, [
        ("S6", "Inscription et validation du concessionnaire", "Demande d’inscription, dépôt des pièces justificatives, contrôle SaaS, création du compte et e-mail des identifiants."),
        ("S7", "Partenariats et contrats", "Demandes initiées par la banque ou le concessionnaire, validation, cycle de vie du contrat et gestion des conditions de collaboration."),
        ("S8", "Produits, publications et suivi des dossiers", "Catalogue centralisé, stock, publications par marketplace et consultation des demandes de financement associées aux produits."),
        ("S9", "Parcours client et demande de financement", "Inscription et vérification du compte, consultation, comparaison, simulation, constitution et traitement du dossier."),
        ("S10", "Assistants conversationnels et validation", "Assistant IA du SaaS, chatbot public à règles métier, contrôle des accès, tests et validation fonctionnelle globale."),
    ])

    # The correspondence table is deliberately placed after the two releases and before the closing Scrum paragraph.
    insert_before(p_conclusion, "Correspondance entre les titres de réalisation et les sprints", p_release1)
    insert_before(
        p_conclusion,
        "Le tableau suivant explicite le rattachement entre la structure fonctionnelle du rapport et les incréments prévus dans le Product Backlog.",
        p_release1_desc,
    )
    add_table_before(
        document,
        p_conclusion,
        table_release1,
        ("Titre ou sous-titre de réalisation", "Sprint", "Incrément fonctionnel livré"),
        [
            ("Chapitre 4 — Socle applicatif SaaS et multi-tenance", "S1", "Services communs, sécurité et cloisonnement des tenants."),
            ("Chapitre 4 — Demande de création de marketplace", "S2", "Parcours de souscription de la banque jusqu’à la soumission."),
            ("Chapitre 4 — Back-office SaaS", "S3", "Fonctions d’administration et de supervision de la plateforme."),
            ("Chapitre 4 — Traitement, approbation et paiement", "S4", "Décision SaaS, paiement Stripe et activation après règlement."),
            ("Chapitre 4 — Back-office Banque et marketplace publique", "S5", "Configuration bancaire et consultation des services publics."),
            ("Chapitre 5 — Inscription et validation du concessionnaire", "S6", "Onboarding contrôlé du concessionnaire."),
            ("Chapitre 5 — Partenariats et contrats", "S7", "Collaboration banque-concessionnaire et contractualisation."),
            ("Chapitre 5 — Produits, publications et suivi des demandes", "S8", "Gestion centralisée du catalogue et des dossiers liés aux produits."),
            ("Chapitre 5 — Parcours client et demande de financement", "S9", "Du compte client jusqu’à la décision bancaire sur le dossier."),
            ("Chapitre 5 — Assistants conversationnels", "S10", "Assistant IA du SaaS, chatbot public et validation finale."),
        ],
        (Inches(2.7), Inches(0.75), Inches(3.05)),
    )

    set_paragraph(
        p_conclusion,
        "À l’issue de chaque sprint, les fonctionnalités réalisées font l’objet de tests techniques et fonctionnels avant leur présentation en Sprint Review. Les retours recueillis permettent d’ajuster les priorités du Product Backlog et de préparer le sprint suivant. Cette planification maintient ainsi une progression cohérente entre les titres du rapport, les besoins fonctionnels et les incréments effectivement livrés.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
