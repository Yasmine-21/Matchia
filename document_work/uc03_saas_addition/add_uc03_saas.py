from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\PFE M2\Platforme SaaS")
SOURCE = Path(r"D:\PFE - BRI Technology\Rapport\Master Report.docx")
OUT_DOCX = ROOT / "document_work" / "Master Report - UC03 supervision SaaS détaillée.docx"
OUT_DRAWIO = ROOT / "document_work" / "UC-03 Administration et supervision de la plateforme SaaS.drawio"
OUT_PNG = ROOT / "document_work" / "uc03_saas_addition" / "figure_3_7_uc03_administration_saas.png"


def font(size: int, bold: bool = False, italic: bool = False):
    suffix = "bd" if bold else ""
    if italic:
        suffix = "i" if not bold else "bi"
    name = {"": "arial.ttf", "bd": "arialbd.ttf", "i": "ariali.ttf", "bi": "arialbi.ttf"}[suffix]
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / name), size)


def centered(draw: ImageDraw.ImageDraw, xy, text, fnt, fill=(0, 0, 0)):
    box = draw.multiline_textbbox((0, 0), text, font=fnt, align="center", spacing=3)
    x, y = xy
    draw.multiline_text((x - (box[2] - box[0]) / 2, y - (box[3] - box[1]) / 2), text, font=fnt, fill=fill, align="center", spacing=3)


def dashed_line(draw, p1, p2, dash=10, gap=7, fill=(40, 40, 40), width=2):
    import math

    x1, y1 = p1
    x2, y2 = p2
    distance = math.hypot(x2 - x1, y2 - y1)
    if distance == 0:
        return
    dx, dy = (x2 - x1) / distance, (y2 - y1) / distance
    pos = 0
    while pos < distance:
        end = min(pos + dash, distance)
        draw.line((x1 + dx * pos, y1 + dy * pos, x1 + dx * end, y1 + dy * end), fill=fill, width=width)
        pos += dash + gap


def arrow_head(draw, start, end, fill=(40, 40, 40), width=2):
    import math

    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    for delta in (2.55, -2.55):
        draw.line((end[0], end[1], end[0] + 13 * math.cos(angle + delta), end[1] + 13 * math.sin(angle + delta)), fill=fill, width=width)


def draw_actor(draw, x, y, label):
    black = (35, 35, 35)
    draw.ellipse((x - 12, y, x + 12, y + 24), outline=black, width=2)
    draw.line((x, y + 24, x, y + 67), fill=black, width=2)
    draw.line((x - 27, y + 39, x + 27, y + 39), fill=black, width=2)
    draw.line((x, y + 67, x - 22, y + 94), fill=black, width=2)
    draw.line((x, y + 67, x + 22, y + 94), fill=black, width=2)
    centered(draw, (x, y + 120), label, font(16, bold=True), black)


def ellipse(draw, box, label, fill, outline, label_font):
    draw.ellipse(box, fill=fill, outline=outline, width=2)
    centered(draw, ((box[0] + box[2]) / 2, (box[1] + box[3]) / 2), label, label_font, (55, 55, 55))


def create_png():
    image = Image.new("RGB", (1900, 1350), "white")
    draw = ImageDraw.Draw(image)
    association = (90, 90, 90)
    primary_fill, primary_stroke = (218, 230, 250), (126, 166, 232)
    secondary_fill, secondary_stroke = (216, 242, 237), (88, 123, 120)
    include_color = (18, 174, 202)

    draw_actor(draw, 115, 570, "Administrateur\nSaaS")
    draw_actor(draw, 115, 1135, "Gemini")

    main = (320, 560, 950, 690)
    ellipse(draw, main, "Administrer et superviser\nla plateforme SaaS", primary_fill, primary_stroke, font(27, bold=True))
    nodes = [
        ((600, 60, 1080, 150), "Consulter le tableau\nde bord"),
        ((600, 210, 1080, 300), "Consulter les\nnotifications"),
        ((1250, 55, 1780, 145), "Gérer les demandes\nde marketplace"),
        ((1250, 200, 1780, 290), "Gérer les banques"),
        ((1250, 345, 1780, 435), "Gérer les utilisateurs\net les rôles"),
        ((1250, 490, 1780, 580), "Gérer les\nconcessionnaires"),
        ((1250, 635, 1780, 725), "Gérer les stores\net les modules"),
        ((1250, 780, 1780, 870), "Gérer les marketplaces\net leurs contenus"),
        ((1250, 925, 1780, 1015), "Gérer les offres, abonnements\net paiements"),
        ((600, 855, 1080, 945), "Consulter les journaux\nd’audit"),
        ((600, 1010, 1080, 1100), "Gérer les paramètres\nde la plateforme"),
        ((600, 1165, 1080, 1260), "Interroger l’assistant IA"),
    ]
    for box, label in nodes:
        ellipse(draw, box, label, secondary_fill, secondary_stroke, font(22, bold=True))

    # Associations: actor to principal use case, Gemini to the AI query use case.
    draw.line((170, 630, 320, 630), fill=association, width=2)
    draw.line((170, 1195, 600, 1212), fill=association, width=2)

    # Mandatory components of the SaaS administration process, drawn as UML «include» links.
    right_start = (950, 625)
    for box, _label in nodes[2:9]:
        endpoint = (box[0], (box[1] + box[3]) / 2)
        dashed_line(draw, right_start, endpoint, fill=include_color, width=2)
        arrow_head(draw, (endpoint[0] - 12, endpoint[1]), endpoint, fill=include_color, width=2)
        centered(draw, ((right_start[0] + endpoint[0]) / 2 + 12, (right_start[1] + endpoint[1]) / 2 - 10), "«include»", font(14), include_color)
    for box, _label in nodes[:2]:
        endpoint = ((box[0] + box[2]) / 2, box[3])
        dashed_line(draw, (650, 560), endpoint, fill=include_color, width=2)
        arrow_head(draw, (endpoint[0], endpoint[1] + 12), endpoint, fill=include_color, width=2)
        centered(draw, ((650 + endpoint[0]) / 2 - 20, (560 + endpoint[1]) / 2 - 10), "«include»", font(14), include_color)
    for box, _label in nodes[9:]:
        endpoint = ((box[0] + box[2]) / 2, box[1])
        dashed_line(draw, (650, 690), endpoint, fill=include_color, width=2)
        arrow_head(draw, (endpoint[0], endpoint[1] - 12), endpoint, fill=include_color, width=2)
    image.save(OUT_PNG, dpi=(220, 220))


def mx_cell(cell_id, value, style, x, y, w, h, parent="1", vertex=True):
    attrs = f'id="{cell_id}" value="{escape(value)}" style="{escape(style)}" parent="{parent}"'
    if vertex:
        return f'<mxCell {attrs} vertex="1"><mxGeometry x="{x}" y="{y}" width="{w}" height="{h}" as="geometry"/></mxCell>'
    return f'<mxCell {attrs} edge="1"><mxGeometry relative="1" as="geometry"/></mxCell>'


def mx_edge(cell_id, source, target, label="", dashed=False):
    style = "edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;endArrow=none;startArrow=none;"
    if dashed:
        style = "edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;endArrow=open;endFill=0;dashed=1;dashPattern=6 6;strokeColor=#12AECA;fontColor=#12AECA;"
    return f'<mxCell id="{cell_id}" value="{escape(label)}" style="{style}" parent="1" source="{source}" target="{target}" edge="1"><mxGeometry relative="1" as="geometry"/></mxCell>'


def create_drawio():
    cells = [
        '<mxCell id="0"/>', '<mxCell id="1" parent="0"/>',
        mx_cell("admin", "Administrateur&lt;br&gt;SaaS", "shape=umlActor;verticalLabelPosition=bottom;verticalAlign=top;html=1;fontSize=15;fontStyle=1;", 25, 470, 130, 125),
        mx_cell("gemini", "Gemini", "shape=umlActor;verticalLabelPosition=bottom;verticalAlign=top;html=1;fontSize=15;fontStyle=1;", 25, 930, 100, 120),
        mx_cell("main", "Administrer et superviser&lt;br&gt;la plateforme SaaS", "ellipse;whiteSpace=wrap;html=1;fontSize=20;fontStyle=1;fillColor=#DAE6FA;strokeColor=#7EA6E8;", 240, 470, 500, 110),
    ]
    nodes = [
        ("dashboard", "Consulter le tableau&lt;br&gt;de bord", 470, 25),
        ("notifications", "Consulter les&lt;br&gt;notifications", 470, 145),
        ("requests", "Gérer les demandes&lt;br&gt;de marketplace", 945, 20),
        ("banks", "Gérer les banques", 945, 135),
        ("users", "Gérer les utilisateurs&lt;br&gt;et les rôles", 945, 250),
        ("dealers", "Gérer les&lt;br&gt;concessionnaires", 945, 365),
        ("stores", "Gérer les stores&lt;br&gt;et les modules", 945, 480),
        ("marketplaces", "Gérer les marketplaces&lt;br&gt;et leurs contenus", 945, 595),
        ("billing", "Gérer les offres, abonnements&lt;br&gt;et paiements", 945, 710),
        ("audit", "Consulter les journaux&lt;br&gt;d’audit", 470, 710),
        ("settings", "Gérer les paramètres&lt;br&gt;de la plateforme", 470, 825),
        ("ai", "Interroger l’assistant IA", 470, 940),
    ]
    for node_id, label, x, y in nodes:
        cells.append(mx_cell(node_id, label, "ellipse;whiteSpace=wrap;html=1;fontSize=16;fontStyle=1;fillColor=#D8F2ED;strokeColor=#587B78;", x, y, 350, 75))
    cells.append(mx_edge("assoc_admin", "admin", "main"))
    cells.append(mx_edge("assoc_gemini", "gemini", "ai"))
    for node_id, _label, _x, _y in nodes:
        cells.append(mx_edge(f"inc_{node_id}", "main", node_id, "«include»", dashed=True))
    xml = """<mxfile host=\"app.diagrams.net\" modified=\"2026-09-10T00:00:00.000Z\" agent=\"Codex\" version=\"24.7.17\" type=\"device\"><diagram id=\"uc03Saas\" name=\"UC-03 Administration SaaS\"><mxGraphModel dx=\"1350\" dy=\"850\" grid=\"1\" gridSize=\"10\" guides=\"1\" tooltips=\"1\" connect=\"1\" arrows=\"1\" fold=\"1\" page=\"1\" pageScale=\"1\" pageWidth=\"1169\" pageHeight=\"827\" math=\"0\" shadow=\"0\"><root>""" + "".join(cells) + "</root></mxGraphModel></diagram></mxfile>"
    OUT_DRAWIO.write_text(xml, encoding="utf-8")


def copy_ppr(src, dst):
    src_ppr = src._p.pPr
    if src_ppr is not None:
        if dst._p.pPr is not None:
            dst._p.remove(dst._p.pPr)
        dst._p.insert(0, deepcopy(src_ppr))


def add_text_before(anchor, text, source_para, align=None, italic=False, color=None):
    paragraph = anchor.insert_paragraph_before()
    copy_ppr(source_para, paragraph)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(12)
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return paragraph


def set_cell(cell, text, bold=False):
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(11)
    run.bold = bold
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)


def make_table(doc, anchor):
    rows = [
        ("Élément", "Description"),
        ("Identifiant", "UC-03 — Administrer et superviser la plateforme SaaS"),
        ("Sprint associé", "S3 — Back-office SaaS"),
        ("Acteur principal", "Administrateur SaaS"),
        ("Acteurs secondaires", "Gemini, sollicité par l’assistant analytique pour formuler une réponse aux requêtes autorisées."),
        ("Objectif", "Superviser l’ensemble des ressources globales de la plateforme, administrer les tenants et assurer le suivi de leurs services."),
        ("Préconditions", "L’utilisateur est authentifié avec le rôle Administrateur SaaS et dispose des autorisations requises pour accéder au Back-office SaaS."),
        ("Déclencheur", "L’Administrateur SaaS accède au Back-office afin de consulter un indicateur ou d’administrer une ressource globale."),
        ("Scénario nominal", "1. L’Administrateur SaaS consulte le tableau de bord et les notifications reçues.\n2. Il sélectionne une rubrique du Back-office SaaS.\n3. Il consulte et traite les demandes de création de marketplace selon le processus dédié.\n4. Il gère les banques, les utilisateurs et les rôles, les concessionnaires ainsi que les marketplaces et leurs contenus.\n5. Il gère les stores, les modules, les offres, les abonnements et les opérations de paiement associées.\n6. Il consulte les journaux d’audit ou modifie les paramètres autorisés de la plateforme.\n7. Le système contrôle les données saisies, applique les autorisations et persiste les modifications.\n8. L’Administrateur SaaS peut interroger l’assistant analytique pour obtenir une réponse fondée exclusivement sur des données de consultation autorisées."),
        ("Alternatives et exceptions", "Les listes de banques, demandes, concessionnaires ou abonnements peuvent être recherchées et filtrées selon leur statut. Si une donnée est invalide, si une opération est interdite ou si une ressource est introuvable, le système bloque l’action et affiche un message explicatif. Lorsqu’une demande de marketplace nécessite une approbation, un rejet ou un paiement, le système poursuit le workflow spécialisé correspondant. L’assistant analytique ne peut ni modifier les données ni exécuter une requête non autorisée."),
        ("Postconditions", "Les modifications autorisées sont persistées, les actions sensibles sont tracées dans les journaux d’audit et les notifications ou statuts concernés sont mis à jour. Le périmètre et les données des autres tenants restent inchangés."),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Grid Table 1 Light Accent 5"
    table.autofit = False
    table.columns[0].width = Inches(2.06)
    table.columns[1].width = Inches(4.31)
    for row_idx, row in enumerate(rows):
        set_cell(table.cell(row_idx, 0), row[0], bold=True)
        set_cell(table.cell(row_idx, 1), row[1], bold=(row_idx == 0))
        for cell in table.rows[row_idx].cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            cell_mar = tc_pr.first_child_found_in("w:tcMar")
            if cell_mar is None:
                cell_mar = OxmlElement("w:tcMar")
                tc_pr.append(cell_mar)
            for side in ("top", "start", "bottom", "end"):
                element = cell_mar.find(qn(f"w:{side}"))
                if element is None:
                    element = OxmlElement(f"w:{side}")
                    cell_mar.append(element)
                element.set(qn("w:w"), "90")
                element.set(qn("w:type"), "dxa")
    anchor._p.addprevious(table._tbl)
    return table


def update_uc03():
    doc = Document(SOURCE)
    heading_index, heading = next((i, p) for i, p in enumerate(doc.paragraphs) if "6.3." in p.text and "Administrer la plateforme SaaS" in p.text)
    intro = doc.paragraphs[heading_index + 1]
    caption = next(p for p in doc.paragraphs[heading_index:heading_index + 6] if "Figure 3.7" in p.text)

    heading.text = "6.3. Cas d’utilisation détaillée : Administrer et superviser la plateforme SaaS"
    intro.text = (
        "Ce cas d’utilisation détaille le processus couvert par le sprint S3. Il présente les fonctions de supervision "
        "offertes à l’Administrateur SaaS, depuis la consultation du tableau de bord jusqu’à l’administration des "
        "ressources globales, au suivi des actions et à l’interrogation de l’assistant analytique."
    )
    picture = caption.insert_paragraph_before()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(OUT_PNG), width=Inches(6.30))
    caption.text = "Figure 3.7 — Diagramme de cas d’utilisation détaillé de l’administration et de la supervision SaaS"
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption.runs:
        run = caption.runs[0]
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.italic = True
        run.font.color.rgb = RGBColor(31, 78, 121)
    scenario_intro = next(p for p in doc.paragraphs[heading_index:heading_index + 8] if "scénario du cas d’utilisation UC-03" in p.text)
    scenario_intro.text = "Le tableau suivant présente le scénario du cas d’utilisation UC-03 — Administrer et superviser la plateforme SaaS."

    table = next(
        table for table in doc.tables
        if len(table.rows) >= 2 and table.cell(1, 0).text.strip() == "Acteur principal" and table.cell(1, 1).text.strip() == "Administrateur SaaS"
    )
    descriptions = {
        "Acteur secondaire": "Gemini, sollicité uniquement lors de l’utilisation de l’assistant analytique.",
        "Objectifs": "Permettre à l’Administrateur SaaS de superviser les ressources globales, les tenants et les services associés à la plateforme.",
        "Préconditions": "L’utilisateur est authentifié avec le rôle Administrateur SaaS et dispose des autorisations nécessaires pour accéder au Back-office SaaS.",
        "Postconditions": "Les modifications autorisées sont enregistrées, les notifications ou statuts concernés sont mis à jour et les actions sensibles sont tracées dans les journaux d’audit.",
        "Scénario nominal": "1. L’Administrateur SaaS consulte le tableau de bord et les notifications.\n2. Il sélectionne une rubrique du Back-office SaaS.\n3. Il consulte ou gère les demandes de marketplace, les banques, les utilisateurs et les rôles.\n4. Il gère les concessionnaires, les stores, les modules, les marketplaces et leurs contenus.\n5. Il gère les offres, les abonnements et les paiements associés.\n6. Il consulte les journaux d’audit ou adapte les paramètres autorisés de la plateforme.\n7. Le système applique les contrôles d’accès, vérifie les données et enregistre les modifications.\n8. L’Administrateur SaaS peut interroger l’assistant analytique, qui consulte uniquement les données autorisées avant de retourner une réponse.",
        "Alternatives et exceptions": "Les ressources peuvent être recherchées et filtrées selon leur statut. Si une donnée est invalide, si une opération est interdite ou si une ressource est introuvable, le système bloque l’action et affiche un message explicatif. Le traitement décisionnel d’une demande de marketplace ainsi que son paiement suivent le workflow spécialisé décrit dans le cas d’utilisation correspondant. L’assistant analytique ne peut ni modifier les données ni exécuter une requête non autorisée.",
    }
    for row in table.rows:
        label = row.cells[0].text.strip().replace("conditions", "conditions")
        normalized = label.replace("é", "e").replace("É", "E")
        for key, value in descriptions.items():
            if normalized.lower() == key.lower().replace("é", "e"):
                set_cell(row.cells[1], value)
                break

    doc.save(OUT_DOCX)


if __name__ == "__main__":
    OUT_PNG.parent.mkdir(parents=True, exist_ok=True)
    create_png()
    create_drawio()
    update_uc03()
    print(OUT_DOCX)
    print(OUT_DRAWIO)
    print(OUT_PNG)
