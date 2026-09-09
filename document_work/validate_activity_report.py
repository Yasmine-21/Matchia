from pathlib import Path

from docx import Document


path = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - diagrammes activités modèle harmonisé.docx")
document = Document(path)
captions = [
    "Figure 4.1 — Diagramme d’activité de l’authentification et du contexte d’accès",
    "Figure 4.2 — Diagramme d’activité de création d’une demande de marketplace bancaire",
    "Figure 4.12 — Diagramme d’activité du traitement d’une demande et du paiement Stripe",
    "Figure 5.1 — Diagramme d’activité du parcours métier du Concessionnaire",
    "Figure 5.4 — Diagramme d’activité du parcours Client et de la demande de financement",
]
paragraph_texts = [paragraph.text.strip() for paragraph in document.paragraphs]
print(
    f"exists={path.exists()} bytes={path.stat().st_size} "
    f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} "
    f"inline_shapes={len(document.inline_shapes)}"
)
for caption in captions:
    print(f"caption_count={paragraph_texts.count(caption)} :: {caption}")
activity_alt_texts = [
    shape._inline.docPr.get("descr")
    for shape in document.inline_shapes
    if "Diagramme d’activité" in (shape._inline.docPr.get("descr") or "")
]
print(f"activity_alt_texts={len(activity_alt_texts)}")
