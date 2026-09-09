from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\PFE M2\Platforme SaaS")
INPUT = ROOT / "document_work" / "Master Report - cas utilisation détaillés sprints.docx"
OUTPUT = ROOT / "document_work" / "Master Report - chapitre 6 CI CD final.docx"
FIGURES = ROOT / "figure"
GENERATED = ROOT / "document_work" / "chapter6_figures"


def source_run(paragraph):
    return next((run for run in paragraph.runs if run.text), None)


def copy_ppr(target, template):
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if template._p.pPr is not None:
        target._p.insert(0, deepcopy(template._p.pPr))


def insert_text(anchor, text, template, alignment=None):
    paragraph = anchor.insert_paragraph_before()
    copy_ppr(paragraph, template)
    run = paragraph.add_run(text)
    reference = source_run(template)
    if reference is not None and reference._r.rPr is not None:
        run._r.insert(0, deepcopy(reference._r.rPr))
    if alignment is not None:
        paragraph.alignment = alignment
    return paragraph


def insert_code(anchor, text, template):
    paragraph = insert_text(anchor, text, template)
    paragraph.paragraph_format.left_indent = Inches(0.35)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    for run in paragraph.runs:
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(9)
    return paragraph


def move_before(anchor, element):
    element.getparent().remove(element)
    anchor._p.addprevious(element)


def insert_figure(document, anchor, image_path, caption, caption_template):
    image = Image.open(image_path)
    ratio = image.width / image.height
    width = Inches(5.85 if ratio >= 1.25 else 4.55)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(image_path), width=width)
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    move_before(anchor, paragraph._p)
    return insert_text(anchor, caption, caption_template, WD_ALIGN_PARAGRAPH.CENTER)


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)
    paragraph._p = paragraph._element = None


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margin(cell, value=105):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False):
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(8.7)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margin(cell)


def insert_difficulties_table(document, anchor, caption_template):
    caption = insert_text(anchor, "Tableau 6.1 — Difficultés rencontrées lors de la mise en place du pipeline et solutions appliquées", caption_template, WD_ALIGN_PARAGRAPH.CENTER)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    move_before(anchor, table._tbl)
    headers = ("Élément", "Difficulté rencontrée", "Solution appliquée")
    widths = (Inches(1.35), Inches(2.6), Inches(2.55))
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        set_cell_shading(cell, "D9D9D9")
        set_cell_text(cell, value, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)
    rows = [
        ("Jenkinsfile", "Ordonnancement cohérent des étapes et partage sécurisé des variables.", "Pipeline déclaratif versionné ; utilisation de credentials Jenkins et de variables d’environnement."),
        ("Tests backend", "Séparer la compilation des tests afin de localiser clairement un échec.", "Stage Tests Backend distinct ; arrêt du pipeline lorsque la commande Maven échoue."),
        ("Build frontend", "Injection de l’URL publique de l’API dans le build Vite.", "Transmission de VITE_API_URL au build du frontend."),
        ("SonarQube", "Exploiter la couverture Java et analyser aussi le frontend.", "Rapport JaCoCo pour le backend et scanner Sonar dédié au frontend."),
        ("Webhook", "Jenkins local n’est pas joignable directement depuis GitHub.", "Configuration du webhook GitHub et exposition temporaire via Cloudflare Tunnel."),
        ("Azure CLI et MFA", "Authentifier l’environnement de déploiement tout en respectant la protection du compte Azure.", "Connexion initiale par az login --use-device-code, puis sélection de l’abonnement par le pipeline."),
        ("Azure Files", "Conserver les fichiers téléversés malgré le remplacement d’un conteneur.", "Montage d’un partage Azure Files dans le backend sur le répertoire /app/uploads."),
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].width = widths[index]
            set_cell_text(cells[index], value, bold=index == 0)
    return table


def font(size, bold=False):
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / ("arialbd.ttf" if bold else "arial.ttf")), size)


def box(draw, rect, text, fill="#F2F2F2"):
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, radius=22, fill=fill, outline="#202020", width=3)
    f = font(22, True)
    words, lines, current = text.split(), [], ""
    for word in words:
        candidate = (current + " " + word).strip()
        if not current or draw.textbbox((0, 0), candidate, font=f)[2] <= (x2 - x1 - 30):
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    yy = (y1 + y2) / 2 - len(lines) * 15
    for line in lines:
        width = draw.textbbox((0, 0), line, font=f)[2]
        draw.text(((x1 + x2 - width) / 2, yy), line, font=f, fill="#111111")
        yy += 31


def arrow(draw, start, end):
    import math
    draw.line((*start, *end), fill="#202020", width=4)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 16
    points = [end]
    for delta in (math.pi - math.pi / 7, math.pi + math.pi / 7):
        points.append((end[0] + size * math.cos(angle + delta), end[1] + size * math.sin(angle + delta)))
    draw.polygon(points, fill="#202020")


def create_final_architecture():
    GENERATED.mkdir(parents=True, exist_ok=True)
    path = GENERATED / "fig_6_17_architecture_globale_finale.png"
    image = Image.new("RGB", (2300, 1650), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), "Architecture globale du pipeline CI/CD et du déploiement Matchia", font=font(36, True), fill="#111111")
    draw.line((70, 100, 2230, 100), fill="#202020", width=2)
    # Upper trigger chain.
    top_boxes = [(80, 180, 390, 300, "Développeur git push"), (460, 180, 720, 300, "GitHub"), (790, 180, 1060, 300, "Webhook"), (1130, 180, 1430, 300, "Cloudflare Tunnel"), (1500, 180, 1820, 300, "Jenkins")]
    for rect, text in [(coords, text) for *coords, text in top_boxes]:
        box(draw, rect, text)
    for left, right in zip(top_boxes, top_boxes[1:]):
        arrow(draw, ((left[2] + 10), 240), ((right[0] - 10), 240))
    # Build branches.
    box(draw, (1310, 440, 1640, 570), "Build Backend\nMaven")
    box(draw, (1760, 440, 2090, 570), "Build Frontend\nNode.js Vite")
    arrow(draw, (1660, 300), (1475, 440))
    arrow(draw, (1710, 300), (1925, 440))
    box(draw, (1490, 680, 1820, 810), "Tests Backend")
    box(draw, (1490, 900, 1820, 1030), "SonarQube")
    box(draw, (1490, 1120, 1820, 1250), "Docker Build\nDocker Hub")
    box(draw, (1490, 1340, 1820, 1470), "Azure CLI")
    arrow(draw, (1475, 570), (1605, 680))
    arrow(draw, (1925, 570), (1710, 680))
    arrow(draw, (1655, 810), (1655, 900))
    arrow(draw, (1655, 1030), (1655, 1120))
    arrow(draw, (1655, 1250), (1655, 1340))
    # Azure targets and dependencies.
    draw.rounded_rectangle((80, 620, 1250, 1540), radius=28, outline="#202020", width=3, fill="#FAFAFA")
    draw.text((115, 650), "Microsoft Azure", font=font(27, True), fill="#111111")
    box(draw, (180, 790, 550, 930), "Container App\nFrontend")
    box(draw, (720, 790, 1090, 930), "Container App\nBackend")
    box(draw, (680, 1130, 1030, 1260), "PostgreSQL\nAzure")
    box(draw, (680, 1360, 1030, 1490), "Azure Files\n/app/uploads")
    # Azure CLI deploys a new revision for both application services.
    arrow(draw, (1490, 1385), (365, 930))
    arrow(draw, (1490, 1430), (905, 930))
    arrow(draw, (550, 860), (720, 860))
    arrow(draw, (905, 930), (855, 1130))
    arrow(draw, (905, 930), (855, 1360))
    draw.text((90, 1570), "Le frontend et le backend sont déployés comme deux services distincts. Le backend accède à PostgreSQL Azure et au stockage persistant Azure Files.", font=font(20), fill="#222222")
    image.save(path)
    return path


def main():
    document = Document(INPUT)
    paragraphs = document.paragraphs
    chapter_start_index = next(i for i, p in enumerate(paragraphs) if p.text.strip().startswith("Chapitre 6"))
    old_chapter = paragraphs[chapter_start_index:]
    anchor = paragraphs[chapter_start_index]

    chapter_template = anchor
    body_template = paragraphs[chapter_start_index + 2]
    section_template = paragraphs[chapter_start_index + 1]
    subsection_template = paragraphs[chapter_start_index + 8]
    caption_template = next(p for p in paragraphs if p.style.name == "Caption")
    final_architecture = create_final_architecture()

    def heading(value):
        insert_text(anchor, value, section_template)

    def subheading(value):
        insert_text(anchor, value, subsection_template)

    def paragraph(value):
        insert_text(anchor, value, body_template)

    def code(value):
        insert_code(anchor, value, body_template)

    def figure(filename, caption):
        insert_figure(document, anchor, filename, caption, caption_template)

    insert_text(anchor, "Chapitre 6 : Mise en place du pipeline CI/CD et déploiement de Matchia sur Microsoft Azure", chapter_template)

    heading("1. Introduction")
    paragraph("Ce chapitre présente la mise en place d’une chaîne permettant de construire, contrôler et déployer Matchia de manière automatisée. Il décrit les mécanismes retenus pour rendre la livraison de nouvelles versions plus fiable, reproductible et traçable, jusqu’à leur mise à disposition sur Microsoft Azure.")

    heading("2. Démarche DevOps et pipeline CI/CD")
    subheading("2.1. DevOps appliqué à Matchia")
    paragraph("DevOps est une démarche qui rapproche le développement et l’exploitation afin de fluidifier la livraison d’une application. Son objectif est d’automatiser les opérations répétitives tout en intégrant des contrôles de qualité avant la mise en production. Pour Matchia, cette démarche relie la gestion du code, la construction des applications, les tests, l’analyse de qualité, la conteneurisation et le déploiement Cloud.")
    subheading("2.2. Intégration et déploiement continus")
    paragraph("L’intégration continue, ou CI, exécute automatiquement les contrôles techniques à chaque évolution du code. Le déploiement continu, ou CD, prend ensuite le relais pour publier une version validée et mettre à jour l’environnement Cloud.")
    code("Code → Build → Tests → Analyse qualité")
    code("Version validée → Image Docker → Registry → Cloud")
    paragraph("Après avoir présenté les principes de la démarche DevOps et du CI/CD, nous détaillons dans la suite la mise en œuvre concrète du pipeline Matchia.")

    heading("3. Mise en place du pipeline CI/CD avec Jenkins")
    subheading("3.1. Jenkins et définition du pipeline")
    paragraph("Jenkins est un serveur d’automatisation capable d’exécuter des tâches de construction, de contrôle et de déploiement. Dans Matchia, il orchestre l’ensemble du pipeline. Il est exécuté dans un conteneur Docker, ce qui facilite son installation et la reproductibilité de son environnement d’exécution.")
    paragraph("Les extensions Jenkins mobilisées permettent notamment l’exécution de pipelines, l’intégration avec GitHub, l’utilisation de Node.js, l’analyse SonarQube et la gestion des credentials. Maven, Node.js, Docker et Azure CLI sont ainsi disponibles pour les étapes qui les nécessitent. Les secrets, tels que les accès au registre Docker Hub ou les jetons d’analyse, sont stockés dans les credentials Jenkins et ne sont pas inscrits dans le dépôt source.")
    figure(FIGURES / "jenkins1.png", "Figure 6.1 — Construction de l’image Docker utilisée pour exécuter Jenkins")
    paragraph("Le pipeline est défini dans le fichier Jenkinsfile placé à la racine du dépôt. Cette approche Pipeline as Code garantit que les étapes d’intégration et de déploiement évoluent avec le code applicatif.")
    figure(FIGURES / "CréationPipeline.png", "Figure 6.2 — Création d’un pipeline Jenkins pour le projet Matchia")

    subheading("3.2. Git, GitHub et étape Checkout")
    paragraph("Git assure le suivi des modifications locales, tandis que GitHub centralise la version partagée du projet. À chaque évolution validée, le développeur transmet la version concernée au dépôt distant. Jenkins récupère ensuite le code lors du stage Checkout.")
    code("Développeur → Git → GitHub → Jenkins Checkout")
    paragraph("Le stage Checkout exécute la récupération du dépôt configuré dans Jenkins. Il fournit ainsi au pipeline le backend Spring Boot, le frontend React, les Dockerfiles, le fichier Docker Compose et le Jenkinsfile correspondant à la version à livrer.")

    subheading("3.3. Docker et préparation locale des composants")
    paragraph("Docker est utilisé une seule fois comme technologie de conteneurisation commune au backend, au frontend et au serveur Jenkins. Il permet d’isoler chaque composant avec ses dépendances et de produire une image exécutable de façon homogène entre l’environnement local et Azure.")
    paragraph("Le Dockerfile du backend prépare l’API Spring Boot avec Java 17 et Maven. Il suit une construction en plusieurs étapes : Maven produit le fichier JAR, puis une image d’exécution conserve uniquement l’artefact nécessaire. Le backend expose ses services REST sur le port 8081.")
    figure(FIGURES / "imageBackend.png", "Figure 6.3 — Dockerfile du backend Spring Boot exposé sur le port 8081")
    paragraph("Le Dockerfile du frontend s’appuie sur Node.js pour installer les dépendances et produire les fichiers React avec Vite. Ces fichiers sont ensuite servis par Nginx sur le port 80. La variable VITE_API_URL est injectée au moment de la construction afin que le frontend déployé contacte l’URL publique du backend.")
    figure(FIGURES / "imageFrontend.png", "Figure 6.4 — Dockerfile du frontend React Vite servi par Nginx")
    paragraph("Avant l’intégration dans Jenkins, les deux composants ont été démarrés et vérifiés localement avec Docker Compose. Cette étape a permis de tester la communication frontend-backend, les variables d’environnement, les ports exposés et le volume destiné aux fichiers téléversés.")
    figure(FIGURES / "dockercompose.png", "Figure 6.5 — Configuration Docker Compose pour les tests locaux des composants Matchia")

    subheading("3.4. Construction et tests dans Jenkins")
    paragraph("Le pipeline entre ensuite dans ses stages de construction. Le stage Build Backend utilise Maven. Pour une vérification Maven complète, la commande suivante peut être exécutée :")
    code("./mvnw clean verify")
    paragraph("Dans le Jenkinsfile de la configuration finale, la compilation et les tests sont volontairement séparés : Build Backend exécute ./mvnw clean package -DskipTests, puis le stage Tests Backend exécute ./mvnw test. Cette séparation identifie clairement l’étape responsable d’un éventuel échec.")
    paragraph("Le stage Build Frontend installe les dépendances de manière reproductible avec npm ci, puis exécute npm run build après avoir fourni VITE_API_URL.")
    code("npm ci")
    code("VITE_API_URL=$AZURE_BACKEND_URL npm run build")
    figure(FIGURES / "CommandeBackend.png", "Figure 6.6 — Exécution du stage Build Backend dans Jenkins")
    paragraph("La logique de continuité du pipeline est la suivante :")
    code("Build → Tests Backend → succès : pipeline continue / échec : pipeline arrêté")
    paragraph("Les tests backend constituent une barrière de qualité avant la création des images. L’exécution validée de la suite de tests affiche notamment : Tests run: 338, Failures: 0, Errors: 0, Skipped: 0.")
    figure(FIGURES / "testBackend.png", "Figure 6.7 — Résultat de l’exécution des tests backend avant la livraison")

    subheading("3.5. Analyse de qualité avec SonarQube")
    paragraph("SonarQube est l’outil d’analyse statique utilisé dans le pipeline. Il complète les tests en contrôlant les défauts potentiels, les vulnérabilités, les duplications et les odeurs de code. Le stage SonarQube Backend exploite le rapport de couverture JaCoCo produit par Maven. La configuration finale comporte également un stage SonarQube Frontend exécuté avec le scanner dédié au projet React.")
    code("Checkout → Build Backend → Build Frontend → Tests Backend → SonarQube Backend → SonarQube Frontend")
    figure(FIGURES / "SonarQubeProjects.png", "Figure 6.8 — Projets backend et frontend analysés dans SonarQube")

    subheading("3.6. Construction des images et publication dans Docker Hub")
    paragraph("Après validation du code, Jenkins réutilise les Dockerfiles précédemment présentés afin de construire automatiquement les images du frontend et du backend. Cette étape ne redéfinit pas Docker : elle automatise la création des artefacts conteneurisés à partir de la version contrôlée du dépôt.")
    code("SonarQube → Docker Build → image backend + image frontend")
    paragraph("Chaque image reçoit le tag latest ainsi qu’un tag correspondant au BUILD_NUMBER Jenkins. Le tag latest simplifie le déploiement de la dernière version validée ; le tag de build assure la traçabilité entre une exécution Jenkins et une image précise.")
    paragraph("Docker Hub joue le rôle de registre distant. Jenkins s’y authentifie par credentials puis publie les images matchia-backend et matchia-frontend. Elles deviennent ainsi disponibles pour les Azure Container Apps.")
    code("Jenkins → Docker Build → Docker Hub")
    figure(FIGURES / "créationRepoDockerHub.png", "Figure 6.9 — Registre Docker Hub destiné aux images Matchia")

    subheading("3.7. Environnement Azure cible")
    paragraph("Microsoft Azure est retenu comme environnement cible afin d’héberger les composants conteneurisés, les données métier et les fichiers persistants. Cette séparation répond au caractère remplaçable des conteneurs : une nouvelle révision applicative ne doit ni supprimer les données PostgreSQL ni effacer les documents téléversés.")
    figure(FIGURES / "azure.png", "Figure 6.10 — Ressources Azure configurées pour l’environnement de déploiement")
    paragraph("La migration de la base PostgreSQL depuis l’environnement local vers Azure suit le principe suivant :")
    code("PostgreSQL local → pg_dump → pg_restore → PostgreSQL Azure")
    paragraph("Le backend utilise ensuite une chaîne de connexion JDBC fournie par des variables d’environnement, sans inscrire les informations de connexion dans le code source.")
    figure(FIGURES / "ConfigurationDataBase.png", "Figure 6.11 — Configuration du service PostgreSQL Azure utilisé par Matchia")
    paragraph("Le répertoire /app/uploads est associé à un partage Azure Files. Il permet de conserver les justificatifs, images et autres contenus téléversés même lorsqu’une nouvelle révision du backend est créée.")
    figure(FIGURES / "StockageAzureDataBase.png", "Figure 6.12 — Stockage persistant Azure Files associé au répertoire /app/uploads")
    paragraph("Le déploiement repose sur deux services distincts : matchia-frontend pour l’application React servie par Nginx et matchia-backend pour l’API Spring Boot. Le frontend contacte le backend, lequel accède à PostgreSQL Azure et au partage Azure Files.")
    code("Frontend Container App → Backend Container App → PostgreSQL Azure + Azure Files")
    figure(FIGURES / "ConfigContainerFrontend.png", "Figure 6.13 — Configuration de l’Azure Container App du frontend")
    figure(FIGURES / "ConfigurationContainerBackend.png", "Figure 6.14 — Configuration de l’Azure Container App du backend")

    heading("4. Automatisation du déploiement")
    subheading("4.1. Rôle de Azure CLI")
    paragraph("Azure CLI est l’interface en ligne de commande de Microsoft Azure. Jenkins orchestre le pipeline, tandis qu’Azure CLI lui permet d’exécuter les commandes nécessaires au déploiement sur Microsoft Azure. Une authentification initiale peut être réalisée avec le mécanisme de code appareil, adapté aux contraintes de MFA :")
    code("az login --use-device-code")
    paragraph("Après cette préparation, le pipeline sélectionne l’abonnement configuré puis met à jour les deux Container Apps. La commande suivante illustre la mise à jour du backend ; le même principe est appliqué au frontend.")
    code("az containerapp update --name matchia-backend --resource-group rg-matchia --image yassmine24/matchia-backend:latest")
    paragraph("Chaque mise à jour crée une nouvelle révision Azure Container Apps, identifiée par le numéro de build. Cette révision permet de suivre la version effectivement déployée et d’isoler les mises à jour du frontend et du backend.")

    heading("5. Automatisation du déclenchement")
    subheading("5.1. GitHub Webhook et Cloudflare Tunnel")
    paragraph("Le déclenchement automatique repose sur le webhook GitHub : un git push notifie Jenkins, qui démarre le pipeline sans intervention manuelle.")
    code("git push → GitHub → Webhook → Jenkins")
    figure(FIGURES / "ConfigWebhookSurGithub.png", "Figure 6.15 — Configuration du webhook GitHub déclenchant le pipeline Jenkins")
    paragraph("Pendant la mise en place, Jenkins est exécuté localement et son adresse localhost n’est pas accessible directement depuis GitHub. Cloudflare Tunnel a donc exposé temporairement une adresse HTTPS publique redirigée vers Jenkins local.")
    code("GitHub → Cloudflare Tunnel → Jenkins")
    figure(FIGURES / "CréationTunnelCloudflareTempVersJenkins.png", "Figure 6.16 — Exposition temporaire de Jenkins local au moyen de Cloudflare Tunnel")

    heading("6. Architecture globale")
    paragraph("Le schéma suivant récapitule le fonctionnement complet, depuis le push du développeur jusqu’au déploiement des deux services conteneurisés. Il présente une vue unique des flux de déclenchement, de contrôle, de publication et de déploiement.")
    figure(final_architecture, "Figure 6.17 — Architecture globale du pipeline CI/CD et du déploiement de Matchia")

    heading("7. Validation du processus complet")
    paragraph("La validation de bout en bout consiste à transmettre une modification vers la branche suivie par Jenkins avec la commande suivante :")
    code("git push origin main")
    paragraph("La chaîne attendue est alors : GitHub, webhook, Jenkins, build, tests, SonarQube, Docker, Docker Hub, Azure CLI puis Azure Container Apps. Une exécution réussie conduit à la publication des images et à la création des nouvelles révisions applicatives.")
    code("GitHub ✓ → Webhook ✓ → Jenkins ✓ → Build ✓ → Tests ✓ → SonarQube ✓ → Docker ✓ → Docker Hub ✓ → Azure CLI ✓ → Azure Container Apps ✓")
    paragraph("La capture suivante illustre l’accessibilité du Back-office SaaS depuis l’URL de l’application déployée. Elle constitue une vérification fonctionnelle complémentaire après la mise à jour des services Cloud.")
    figure(FIGURES / "DéploiementRéussi.png", "Figure 6.18 — Accès au Back-office SaaS après le déploiement de la plateforme")

    heading("8. Difficultés rencontrées et solutions")
    paragraph("La mise en place a nécessité plusieurs ajustements techniques. Le tableau suivant synthétise les difficultés réelles rencontrées et les solutions retenues afin de fiabiliser l’exécution automatisée.")
    insert_difficulties_table(document, anchor, caption_template)

    heading("9. Conclusion")
    paragraph("Ce chapitre a présenté une chaîne CI/CD complète pour Matchia. Jenkins automatise la récupération du code, les constructions, les tests, les analyses SonarQube, la création et la publication des images Docker, puis le déploiement sur Azure via Azure CLI. La plateforme peut ainsi être livrée de manière plus cohérente, contrôlée et reproductible à partir d’un simple push Git.")

    for old in old_chapter:
        remove_paragraph(old)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
