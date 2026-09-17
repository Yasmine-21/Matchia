from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def norm(text):
    return " ".join(text.split())


def main():
    path = Path(sys.argv[1])
    doc = Document(path)
    expected = [
        "Diagramme d’activité de l’authentification",
        "Diagramme de séquence de l’authentification",
        "Diagramme d’activité de création",
        "Diagramme de séquence de la création",
        "Diagramme d’activité du traitement",
        "Diagramme de séquence de l’approbation",
        "Diagramme d’activité du renouvellement",
        "Diagramme d’activité du parcours métier du Concessionnaire",
        "Diagramme de séquence de l’inscription et de la validation d’un concessionnaire",
        "Diagramme de séquence de la gestion et de la publication multi-banque",
        "Diagramme de séquence de l’inscription et de l’activation du compte Client",
        "Diagramme d’activité du parcours Client",
        "Diagramme de séquence de la constitution et du traitement d’une demande de financement",
        "Diagramme de séquence de l’assistant IA",
        "Diagramme de séquence du chatbot public",
    ]
    paragraphs = doc.paragraphs
    texts = [norm(p.text) for p in paragraphs]
    found_indices = []
    for label in expected:
        matches = [i for i, text in enumerate(texts) if label in text]
        assert len(matches) == 1, (label, matches)
        index = matches[0]
        found_indices.append(index)
        caption = paragraphs[index]
        assert caption.style.name == "Caption", (label, caption.style.name)
        assert caption.alignment == WD_ALIGN_PARAGRAPH.CENTER, label
        previous = caption._p.getprevious()
        assert previous is not None and previous.xpath(".//w:drawing | .//w:pict"), f"Image absente avant {label}"
        page_break = previous.find("./w:pPr/w:pageBreakBefore", previous.nsmap)
        assert page_break is not None, f"Saut de page absent avant {label}"

    assert found_indices == sorted(found_indices), found_indices

    diagram_shapes = []
    for shape in doc.inline_shapes:
        descr = shape._inline.docPr.get("descr", "")
        if descr.startswith("Diagramme"):
            diagram_shapes.append((descr, shape.width / 914400, shape.height / 914400))
    assert len(diagram_shapes) == 15, len(diagram_shapes)
    assert all(width <= 6.01 and height <= 8.6 for _, width, height in diagram_shapes), diagram_shapes

    chapter_numbers = {4: [], 5: []}
    active = None
    for text in texts:
        if text.startswith("Chapitre 4"):
            active = 4
        elif text.startswith("Chapitre 5"):
            active = 5
        elif text.startswith("Chapitre 6"):
            active = None
        if active in chapter_numbers:
            match = re.match(rf"^Figure {active}\.(\d+)", text)
            if match:
                chapter_numbers[active].append(int(match.group(1)))
    for chapter, numbers in chapter_numbers.items():
        assert numbers == list(range(1, len(numbers) + 1)), (chapter, numbers)

    assert not any("NAJA" in text for text in texts)
    assert "5. Conclusion" in texts
    print(f"REPORT OK: paragraphs={len(paragraphs)}, tables={len(doc.tables)}, images={len(doc.inline_shapes)}")
    print(f"DIAGRAMS OK: {len(diagram_shapes)} inserted, ordered, captioned, page-broken and within page bounds")
    print(f"FIGURES OK: chapter4={len(chapter_numbers[4])}, chapter5={len(chapter_numbers[5])}")
    for descr, width, height in diagram_shapes:
        print(f"{width:.2f}x{height:.2f} in | {descr}")


if __name__ == "__main__":
    main()
