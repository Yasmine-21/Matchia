from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml.ns import qn

SKILL_DIR = Path(r"C:\Users\ASUS\.codex\plugins\cache\openai-primary-runtime\documents\26.826.12353\skills\documents")
sys.path.insert(0, str(SKILL_DIR / "scripts"))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa


INPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - rapport mis à jour avec planification Agile.docx")
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - rapport mis à jour - backlog lisible.docx")


def set_cell(cell, text, *, bold=False, centered=False):
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def insert_compact_backlog_table(document, anchor, rows):
    table = document.add_table(rows=1, cols=3)
    table.style = "Grid Table 1 Light Accent 1"
    headers = ("Sprint", "Objectif", "Éléments clés du backlog")
    for index, label in enumerate(headers):
        set_cell(table.rows[0].cells[index], label, bold=True, centered=True)
    for sprint, objective, elements in rows:
        cells = table.add_row().cells
        set_cell(cells[0], sprint, bold=True, centered=True)
        set_cell(cells[1], objective, bold=True)
        set_cell(cells[2], elements)

    # The first report section governs chapter 1; use its actual usable width.
    usable_width = section_content_width_dxa(document.sections[0])
    widths = column_widths_from_weights((1.0, 2.1, 4.4), usable_width)
    apply_table_geometry(table, widths, table_width_dxa=sum(widths))
    anchor._p.addprevious(table._tbl)


def main():
    document = Document(INPUT)
    backlog_tables = [
        table for table in document.tables
        if len(table.columns) == 2
        and table.cell(0, 0).text.strip() == "Sprint"
        and table.cell(0, 1).text.strip() == "Éléments du Product Backlog"
    ]
    if len(backlog_tables) != 2:
        raise RuntimeError(f"Deux tableaux de backlog étaient attendus, {len(backlog_tables)} trouvés.")

    for table in backlog_tables:
        table._tbl.getparent().remove(table._tbl)

    release2 = next(p for p in document.paragraphs if p.text.strip().startswith("Release 2 —"))
    closing = next(p for p in document.paragraphs if p.text.strip().startswith("À l’issue de chaque sprint"))

    insert_compact_backlog_table(
        document,
        release2,
        [
            ("S1", "Socle SaaS", "Authentification, rôles, sessions et isolation multi-tenant."),
            ("S2", "Onboarding bancaire", "Demande, vérification e-mail, configuration, stores et modules."),
            ("S3", "Administration SaaS", "Tableau de bord, demandes, banques, catalogues, utilisateurs et audit."),
            ("S4", "Abonnement et paiement", "Approbation ou rejet, e-mails, Stripe et activation conditionnelle."),
            ("S5", "Espaces Banque et public", "Personnalisation, contenus, services souscrits, produits, comparateur et simulateur."),
        ],
    )
    insert_compact_backlog_table(
        document,
        closing,
        [
            ("S6", "Onboarding concessionnaire", "Inscription, documents, validation SaaS, compte et e-mail d’accès."),
            ("S7", "Partenariats et produits", "Demandes banque/concessionnaire, contrats, catalogue, stock et publications."),
            ("S8", "Parcours client", "Consultation, comparaison, simulation, compte, dossier et décision bancaire."),
            ("S9", "Assistants et qualité", "Assistant IA, chatbot public, sécurité, tests, corrections et validation globale."),
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
