from pathlib import Path

from docx import Document


INPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - planification Agile alignée.docx")
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - planification Agile détaillée.docx")


def replace_cell_text(cell, text):
    paragraph = cell.paragraphs[0]
    reference = next((run for run in paragraph.runs if run.text), None)
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    if reference is not None and reference._r.rPr is not None:
        from copy import deepcopy
        run._r.insert(0, deepcopy(reference._r.rPr))
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def main():
    document = Document(INPUT)
    release1 = document.tables[1]

    # Align the first release with the full functional content documented in Chapter 4.
    detailed_items = [
        "Connexion, déconnexion, renouvellement de session et réinitialisation du mot de passe ; gestion des rôles Administrateur SaaS, Administrateur Banque, Concessionnaire et Client ; rattachement au tenant et isolation des utilisateurs, contenus, produits, configurations, stores et modules de chaque banque.",
        "Accès au formulaire public ; informations de la banque (nom, e-mail, téléphone, site web, description et logo) ; informations du futur Administrateur Banque (nom, prénom, e-mail, téléphone et photo) ; vérification par code e-mail ; slug, couleurs et bannière ; sélection des stores et modules ; récapitulatif, soumission, notification SaaS et e-mail de confirmation.",
        "Tableau de bord et indicateurs ; catalogue des stores et modules ; gestion des banques, des utilisateurs et des rôles, des concessionnaires, des marketplaces et de leurs contenus ; gestion des offres et abonnements ; consultation des audits et logs ; paramètres généraux de la plateforme.",
        "Consultation détaillée des demandes ; approbation ou rejet avec motif et e-mail ; préparation de la banque, de la marketplace, de l’abonnement et du compte Administrateur Banque ; génération du lien de paiement Stripe ; vérification du paiement ; activation conditionnelle des services et envoi des identifiants du Back-office Banque.",
        "Back-office Banque : tableau de bord, personnalisation (couleurs, logo, bannières, textes et contenus), stores, modules et paramètres du simulateur, produits, concessionnaires, partenariats, clients, demandes de financement, abonnements, historique, demandes de renouvellement ou d’extension et notifications ; marketplace publique : page d’accueil personnalisée, stores et modules accessibles selon les services souscrits.",
    ]

    for row, details in zip(release1.rows[1:], detailed_items):
        replace_cell_text(row.cells[2], details)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
