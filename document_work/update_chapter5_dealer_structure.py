from pathlib import Path

from docx import Document


INPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - chapitre 4 réorganisé.docx")
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - chapitres 4 et 5 réorganisés.docx")


def replace_text(paragraph, text):
    """Replace paragraph content while keeping the existing paragraph/run styling."""
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def find_paragraph(document, starts_with):
    return next(paragraph for paragraph in document.paragraphs if paragraph.text.strip().startswith(starts_with))


def main():
    document = Document(INPUT)

    replacements = {
        "2. Gestion des concessionnaires": "2. Parcours Concessionnaires",
        "2.1. Inscription et validation du concessionnaire": "2.1. Inscription et validation du concessionnaire",
        "3. Gestion des partenariats et des contrats": "2.2. Partenariats et contrats",
        "4. Gestion des produits et publication multi-banque": "2.3. Produits et publications multi-banque",
        "5. Suivi des demandes de financement côté concessionnaire": "2.4. Suivi des demandes de financement",
        "6. Parcours client et demande de financement": "3. Parcours client et demande de financement",
        "6.1. Consultation, comparaison et simulation": "3.1. Consultation, comparaison et simulation",
        "6.2. Inscription du client et activation du compte": "3.2. Inscription du client et activation du compte",
        "6.3. Constitution et soumission du dossier": "3.3. Constitution et soumission du dossier",
        "6.4. Traitement de la demande par la banque": "3.4. Traitement de la demande par la banque",
        "7. Assistants conversationnels de la plateforme": "4. Assistants conversationnels de la plateforme",
        "7.1. Assistant IA pour le pilotage du SaaS": "4.1. Assistant IA pour le pilotage du SaaS",
        "7.2. Fonctionnement et sécurisation de l’assistant IA": "4.2. Fonctionnement et sécurisation de l’assistant IA",
        "7.3. Chatbot public des marketplaces": "4.3. Chatbot public des marketplaces",
        "8. Introduction": "5. Conclusion",
    }

    for old_start, new_text in replacements.items():
        replace_text(find_paragraph(document, old_start), new_text)

    partnership = find_paragraph(document, "Un partenariat peut être proposé")
    replace_text(
        partnership,
        "La demande de partenariat peut être initiée dans les deux sens. Depuis son espace, le concessionnaire sélectionne la banque et le store auxquels il souhaite proposer sa collaboration ; la banque destinataire peut ensuite consulter la demande et l’approuver ou la rejeter. Inversement, l’Administrateur Banque peut sélectionner un concessionnaire disponible et lui adresser une demande de partenariat. Le concessionnaire reçoit alors cette demande dans son espace et peut à son tour l’accepter ou la rejeter. Dans les deux cas, la décision et, le cas échéant, son motif sont enregistrés afin d’assurer la traçabilité du traitement.",
    )

    replace_text(
        find_paragraph(document, "Le concessionnaire dispose également une section"),
        "Le concessionnaire dispose également d’une section dédiée au suivi des demandes de financement associées à ses produits. Cette fonctionnalité prolonge le principe de compte unique : les demandes provenant de différentes marketplaces partenaires peuvent être consultées depuis le même espace.",
    )
    replace_text(
        find_paragraph(document, "Cependant, la publication d’un produit"),
        "Cependant, la publication d’un produit reste soumise à plusieurs conditions. Le concessionnaire doit être approuvé, le partenariat et le contrat doivent être valides, le store concerné doit être actif et la publication doit être approuvée par la banque.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
