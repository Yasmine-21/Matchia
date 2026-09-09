from copy import deepcopy
from pathlib import Path

from docx import Document


INPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - rapport mis à jour avec planification Agile.docx")
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\document_work\Master Report - planification Agile lisible.docx")


def first_formatted_run(paragraph):
    return next((run for run in paragraph.runs if run.text), None)


def clone_paragraph_format(target, template):
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if template._p.pPr is not None:
        target._p.insert(0, deepcopy(template._p.pPr))


def apply_run_format(target, source):
    if source is not None and source._r.rPr is not None:
        target._r.insert(0, deepcopy(source._r.rPr))


def insert_paragraph_before(anchor, text, template, label=None):
    paragraph = anchor.insert_paragraph_before()
    clone_paragraph_format(paragraph, template)
    source_run = first_formatted_run(template)
    if label:
        label_run = paragraph.add_run(label)
        apply_run_format(label_run, source_run)
        label_run.bold = True
    text_run = paragraph.add_run(text)
    apply_run_format(text_run, source_run)
    return paragraph


def remove_table(table):
    table._element.getparent().remove(table._element)


def add_sprint_blocks(document, anchor, title_template, body_template, sprints):
    for number, title, objective, backlog in sprints:
        insert_paragraph_before(anchor, f"Sprint {number} — {title}", title_template)
        insert_paragraph_before(anchor, objective, body_template, label="Objectif : ")
        insert_paragraph_before(anchor, backlog, body_template, label="Éléments du backlog : ")


def main():
    document = Document(INPUT)

    # The two dense backlog tables are deliberately replaced with short sprint
    # blocks, preserving the rest of the report unchanged.
    dense_tables = [
        table for table in document.tables
        if len(table.rows) and len(table.rows[0].cells) >= 2
        and table.cell(0, 0).text.strip() == "Sprint"
        and table.cell(0, 1).text.strip() == "Éléments du Product Backlog"
    ]
    for table in dense_tables:
        remove_table(table)

    release_1 = next(p for p in document.paragraphs if p.text.strip().startswith("Release 1 —"))
    release_2 = next(p for p in document.paragraphs if p.text.strip().startswith("Release 2 —"))
    conclusion = next(p for p in document.paragraphs if p.text.strip() == "Conclusion")
    closing_paragraph = next(p for p in document.paragraphs if p.text.strip().startswith("À l’issue de chaque sprint"))
    closing_text = closing_paragraph.text
    closing_paragraph._element.getparent().remove(closing_paragraph._element)
    title_template = release_1
    body_template = next(p for p in document.paragraphs if p.text.strip().startswith("Cette release vise"))

    add_sprint_blocks(
        document,
        release_2,
        title_template,
        body_template,
        [
            ("1", "Socle applicatif", "Mettre en place les fondations techniques et de sécurité de la plateforme.", "Authentification et gestion des sessions ; rôles et autorisations ; isolation des données dans un contexte multi-tenant."),
            ("2", "Demande de création d’une marketplace", "Permettre à une banque de soumettre une demande complète et vérifiée.", "Informations de la banque et de l’Administrateur Banque ; vérification de l’adresse électronique ; configuration initiale ; stores, modules, récapitulatif et soumission."),
            ("3", "Back-office SaaS", "Fournir à l’Administrateur SaaS les fonctions de supervision de la plateforme.", "Tableau de bord ; gestion des banques, demandes, stores et modules ; utilisateurs et rôles ; concessionnaires ; marketplaces ; contenu, audit et paramètres."),
            ("4", "Traitement, paiement et activation", "Automatiser le passage d’une demande approuvée à une marketplace active.", "Approbation ou rejet ; notifications et e-mails ; paiement Stripe ; vérification de la transaction ; activation de la banque, de la marketplace, de l’abonnement et du compte Administrateur Banque."),
            ("5", "Back-office Banque et marketplace publique", "Permettre à chaque banque d’administrer et de présenter ses services.", "Personnalisation et contenus ; stores et modules ; abonnement ; produits publics ; comparateur et simulateur."),
        ],
    )

    add_sprint_blocks(
        document,
        conclusion,
        title_template,
        body_template,
        [
            ("6", "Inscription du concessionnaire", "Mettre en place l’accès contrôlé des concessionnaires à la plateforme.", "Saisie des informations ; dépôt des documents justificatifs ; traitement de la demande par le SaaS ; création du compte et envoi des identifiants."),
            ("7", "Partenariats et produits", "Permettre la collaboration multi-banque et la gestion centralisée du catalogue.", "Demandes de partenariat initiées par la banque ou le concessionnaire ; contrats ; catalogue centralisé ; gestion du stock ; publications multi-banques."),
            ("8", "Parcours client et financement", "Permettre au client de passer de la découverte d’une offre à la soumission d’un dossier.", "Consultation, comparaison et simulation ; inscription et vérification du compte ; constitution du dossier ; dépôt des documents ; transmission à la banque."),
            ("9", "Assistants conversationnels et finalisation", "Faciliter l’accès à l’information et consolider la qualité de la solution.", "Assistant IA du Back-office SaaS ; chatbot public ; contrôle des accès aux données ; tests fonctionnels ; corrections et validation des parcours."),
        ],
    )
    insert_paragraph_before(conclusion, closing_text, body_template)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
