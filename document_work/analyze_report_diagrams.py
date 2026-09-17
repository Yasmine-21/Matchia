from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def paragraph_image_ids(paragraph):
    ids = []
    for blip in paragraph._p.xpath('.//a:blip'):
        rid = blip.get(qn('r:embed'))
        if rid:
            ids.append(rid)
    return ids


def main() -> None:
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    doc = Document(source)
    relationship_targets = {
        rid: rel.target_ref
        for rid, rel in doc.part.rels.items()
        if 'image' in rel.reltype
    }
    paras = []
    for idx, paragraph in enumerate(doc.paragraphs):
        text = ' '.join(paragraph.text.split())
        style = paragraph.style.name if paragraph.style else ''
        images = paragraph_image_ids(paragraph)
        paras.append({'index': idx, 'style': style, 'text': text, 'images': images})

    pattern = re.compile(r'(diagramme|activité|activite|séquence|sequence)', re.I)
    hits = []
    for item in paras:
        if item['text'] and pattern.search(item['text']):
            start = max(0, item['index'] - 4)
            end = min(len(paras), item['index'] + 5)
            hits.append({'hit': item, 'context': paras[start:end]})

    # Associate each image-containing paragraph with nearby textual context.
    image_contexts = []
    for item in paras:
        if item['images']:
            start = max(0, item['index'] - 5)
            end = min(len(paras), item['index'] + 6)
            ctx = [p for p in paras[start:end] if p['text']]
            image_contexts.append({'image_paragraph': item, 'context': ctx})

    data = {
        'source': str(source),
        'paragraph_count': len(paras),
        'paragraphs': paras,
        'table_count': len(doc.tables),
        'hits': hits,
        'image_contexts': image_contexts,
        'relationship_targets': relationship_targets,
    }
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'paragraphs={len(paras)} tables={len(doc.tables)} hits={len(hits)} images={len(image_contexts)}')


if __name__ == '__main__':
    main()
