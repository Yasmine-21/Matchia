from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "diagramme_usecase_authentification_matchia"
FIGURE = OUT_DIR / "Diagramme_cas_utilisation_detaille_Authentification_Matchia.png"
FIGURE_FOR_DOC = OUT_DIR / "Diagramme_cas_utilisation_detaille_Authentification_Matchia_sans_legende.png"
OUTPUT = OUT_DIR / "Section_6_1_UC_01_Authentification_Matchia.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_font(run, name="Aptos", size=10.5, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=4, line=1.08):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def clear_and_write(cell, text, bold=False, color="000000", size=10.2):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.08)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_numbered_lines(cell, lines):
    cell.text = ""
    for index, line in enumerate(lines, start=1):
        p = cell.paragraphs[0] if index == 1 else cell.add_paragraph()
        set_paragraph_spacing(p, after=3, line=1.08)
        run = p.add_run(f"{index}. {line}")
        set_font(run, size=10.0)


def add_bulleted_lines(cell, lines):
    cell.text = ""
    for index, line in enumerate(lines):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        set_paragraph_spacing(p, after=4, line=1.08)
        run = p.add_run(f"• {line}")
        set_font(run, size=10.0)


def add_row(table, label, content, shade=False, mode="text"):
    cells = table.add_row().cells
    set_cell_shading(cells[0], "EDF3F8" if not shade else "E3EDF5")
    set_cell_shading(cells[1], "F7FAFC" if shade else "FFFFFF")
    clear_and_write(cells[0], label, bold=True, color="1F2933", size=10.1)
    if mode == "numbered":
        add_numbered_lines(cells[1], content)
    elif mode == "bulleted":
        add_bulleted_lines(cells[1], content)
    else:
        clear_and_write(cells[1], content, size=10.1)
    for cell in cells:
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return cells


def build_document():
    with Image.open(FIGURE) as source:
        source.crop((0, 0, source.width, source.height - 250)).save(FIGURE_FOR_DOC)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_style.font.size = Pt(20)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(title, after=10)
    set_font(title.add_run("Authentification et contexte d accès"), name="Aptos Display", size=20, bold=True)

    heading = doc.add_paragraph(style="Heading 1")
    set_paragraph_spacing(heading, before=2, after=8)
    set_font(heading.add_run("6.1 Cas d’utilisation détaillé : S’authentifier"), name="Aptos Display", size=15, bold=True)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(intro, after=10, line=1.15)
    set_font(intro.add_run(
        "Ce cas d’utilisation décrit l’authentification commune aux différents rôles de la plateforme. "
        "Le système vérifie l’identité et l’état du compte, charge les autorisations puis, lorsque le rôle "
        "est rattaché à une banque ou à une marketplace, détermine le contexte du tenant avant de créer la session."
    ), size=10.7)

    figure_p = doc.add_paragraph()
    figure_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(figure_p, after=4)
    shape = figure_p.add_run().add_picture(str(FIGURE_FOR_DOC), width=Inches(7.0))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", "Diagramme détaillé de l’authentification et du contexte d’accès Matchia")

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(caption, after=10)
    set_font(caption.add_run("Figure 3.5 — Diagramme de cas d’utilisation détaillé de l’authentification"), size=9.2, italic=True, color="3F4A4A")

    doc.add_page_break()

    table_heading = doc.add_paragraph(style="Heading 2")
    set_paragraph_spacing(table_heading, after=7)
    set_font(table_heading.add_run("Tableau 3.1 — Description du scénario UC-01"), name="Aptos Display", size=12.5, bold=True)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.65)
    table.columns[1].width = Inches(5.05)
    table.style = "Table Grid"
    set_table_borders(table)

    header = table.rows[0].cells
    for index, text in enumerate(("Élément", "Description")):
        header[index].width = Inches(1.65 if index == 0 else 5.05)
        set_cell_shading(header[index], "315B7D")
        clear_and_write(header[index], text, bold=True, color="FFFFFF", size=10.5)
        header[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_margins(header[index], top=120, bottom=120)
    set_repeat_table_header(table.rows[0])

    add_row(table, "Cas d’utilisation", "UC-01 — S’authentifier", shade=False)
    add_row(table, "Acteur principal", "Utilisateur", shade=True)
    add_row(
        table,
        "Rôles concernés",
        "Administrateur SaaS, Administrateur Banque, Concessionnaire et Client.",
        shade=False,
    )
    add_row(table, "Acteur secondaire", "Aucun acteur secondaire direct.", shade=True)
    add_row(
        table,
        "Objectif",
        "Permettre à un utilisateur d’accéder à l’espace correspondant à son rôle, à ses autorisations et, le cas échéant, au tenant auquel il est rattaché.",
        shade=False,
    )
    add_row(
        table,
        "Préconditions",
        "L’utilisateur possède un compte enregistré et accède au formulaire de connexion. Le système d’authentification est disponible.",
        shade=True,
    )
    add_row(
        table,
        "Déclencheur",
        "L’utilisateur saisit ses identifiants et soumet le formulaire de connexion.",
        shade=False,
    )
    add_row(
        table,
        "Postconditions en cas de succès",
        "Une session authentifiée est créée. Le rôle, les autorisations et, si nécessaire, le contexte du tenant sont associés à la session. L’utilisateur est redirigé vers l’espace approprié.",
        shade=True,
    )
    add_row(
        table,
        "Postconditions en cas d’échec",
        "Aucune session valide n’est créée ou conservée. L’accès à la ressource demandée est refusé et un message adapté est affiché.",
        shade=False,
    )
    add_row(
        table,
        "Scénario nominal",
        [
            "L’utilisateur accède au formulaire de connexion.",
            "Il saisit son adresse e-mail et son mot de passe.",
            "Il soumet le formulaire de connexion.",
            "Le système vérifie les identifiants fournis.",
            "Le système vérifie que le compte est actif.",
            "Le système récupère le rôle et contrôle les autorisations de l’utilisateur.",
            "Si l’utilisateur est rattaché à une banque ou à une marketplace, le système détermine le contexte du tenant correspondant.",
            "Le système crée une session authentifiée et y associe le rôle, les autorisations et le contexte d’accès.",
            "L’utilisateur est redirigé vers l’espace correspondant à son rôle.",
        ],
        shade=True,
        mode="numbered",
    )
    add_row(
        table,
        "Alternatives et exceptions",
        [
            "A1 — Identifiants incorrects : à l’étape 4, le système refuse l’authentification et affiche un message d’erreur.",
            "A2 — Compte inactif : à l’étape 5, le système refuse l’authentification et indique que le compte ne permet pas l’accès.",
            "A3 — Contexte du tenant invalide : à l’étape 7, si aucun tenant autorisé ne correspond au rôle ou si l’utilisateur tente d’accéder à un autre tenant, le système refuse l’accès.",
            "A4 — Session ou jeton expiré : le système invalide la session et demande à l’utilisateur de la renouveler ou de se reconnecter.",
            "A5 — Ressource non autorisée : toute tentative d’accès à une fonctionnalité extérieure au périmètre du rôle est refusée.",
        ],
        shade=False,
        mode="bulleted",
    )
    for row in table.rows:
        row.cells[0].width = Inches(1.65)
        row.cells[1].width = Inches(5.05)

    doc.core_properties.title = "Authentification et contexte d’accès"
    doc.core_properties.subject = "UC-01 S’authentifier"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "UML, authentification, multi-tenant, Matchia"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
