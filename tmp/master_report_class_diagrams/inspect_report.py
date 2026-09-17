import json
import sys
from pathlib import Path

from docx import Document

source = Path(sys.argv[1])
output = Path(sys.argv[2])
doc = Document(source)

paragraphs = []
for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if text:
        paragraphs.append({
            "index": index,
            "style": paragraph.style.name if paragraph.style else "",
            "text": text,
        })

payload = {
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "section_count": len(doc.sections),
    "paragraphs": paragraphs,
}
output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
