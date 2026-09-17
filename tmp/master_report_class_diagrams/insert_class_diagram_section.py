from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"D:\PFE - BRI Technology\Rapport\Master Report.docx")
OUTPUT = Path(r"D:\PFE M2\Platforme SaaS\deliverables\rapport\Master Report_avec_diagrammes_classes_Drawio.docx")
DIAGRAMS = Path(r"D:\PFE M2\Platforme SaaS\tmp\master_report_class_diagrams\drawio_exports")
EXPECTED_SHA256 = "d8f6bb0fee7666dd072d5bcbfbbbde5a13cf54ddd734f8641477f5517521bc96"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def set_run_font(run, size: float, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_keep_with_next(paragraph, enabled: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = enabled


def add_heading_before(anchor, text: str, level: int, page_break: bool = False):
    style = "Heading 2" if level == 2 else "isselectedend"
    paragraph = anchor.insert_paragraph_before(style=style)
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    set_keep_with_next(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, 16 if level == 2 else 14, bold=True)
    return paragraph


def add_body_before(anchor, segments):
    paragraph = anchor.insert_paragraph_before(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(6)
    for text, bold in segments:
        run = paragraph.add_run(text)
        set_run_font(run, 12, bold=bold)
    return paragraph


def add_picture_before(anchor, image_path: Path, alt_text: str, width_inches: float = 6.25):
    paragraph = anchor.insert_paragraph_before(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    set_keep_with_next(paragraph)
    inline = paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    inline._inline.docPr.set("descr", alt_text)
    inline._inline.docPr.set("title", alt_text)
    return paragraph


def add_seq_field(paragraph, sequence_name: str) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f"SEQ {sequence_name} \\* ARABIC")
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_fonts = OxmlElement("w:rFonts")
    run_fonts.set(qn("w:ascii"), "Times New Roman")
    run_fonts.set(qn("w:hAnsi"), "Times New Roman")
    run_properties.append(run_fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    run_properties.append(size)
    italic = OxmlElement("w:i")
    run_properties.append(italic)
    run.append(run_properties)
    text = OxmlElement("w:t")
    text.text = "0"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def add_caption_before(anchor, title: str):
    paragraph = anchor.insert_paragraph_before(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("Figure ")
    set_run_font(run, 9, italic=True)
    add_seq_field(paragraph, "Figure")
    run = paragraph.add_run(f" : {title}")
    set_run_font(run, 9, italic=True)
    return paragraph


def add_entity_descriptions_before(anchor, entries):
    lead = anchor.insert_paragraph_before(style="Normal")
    lead.paragraph_format.space_before = Pt(5)
    lead.paragraph_format.space_after = Pt(3)
    set_keep_with_next(lead)
    run = lead.add_run("Tables principales du diagramme :")
    set_run_font(run, 11.5, bold=True)
    for index, (name, description) in enumerate(entries):
        paragraph = anchor.insert_paragraph_before(style="Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.paragraph_format.space_after = Pt(7 if index == len(entries) - 1 else 2)
        bullet = paragraph.add_run("• ")
        set_run_font(bullet, 11.5)
        entity = paragraph.add_run(name)
        set_run_font(entity, 11.5, bold=True)
        detail = paragraph.add_run(f" : {description}")
        set_run_font(detail, 11.5)


def ensure_update_fields_on_open(document: Document) -> None:
    settings = document.settings._element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def replace_paragraph_text(paragraph, text: str, size: float, bold: bool) -> None:
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    set_run_font(run, size, bold=bold)


def add_conclusion_after(heading):
    new_p = OxmlElement("w:p")
    heading._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    paragraph = Paragraph(new_p, heading._parent)
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(
        "Ce chapitre a présenté l’architecture retenue, l’environnement de développement et le modèle "
        "de classes de la plateforme Matchia. La vue globale et ses déclinaisons fonctionnelles clarifient "
        "la responsabilité des entités, la persistance des associations et les dépendances entre les "
        "principaux domaines. Ces éléments constituent la base de l’implémentation présentée dans le chapitre suivant."
    )
    set_run_font(run, 12)


def remove_empty_paragraphs_before_next_section(heading) -> None:
    """Remove source spacer paragraphs that would create a blank page before chapter 4."""
    current = heading._p.getnext()
    if current is not None:
        current = current.getnext()  # paragraph containing the regenerated conclusion

    while current is not None:
        if current.xpath("./w:pPr/w:sectPr"):
            return
        next_element = current.getnext()
        text = "".join(current.itertext()).strip()
        if current.tag == qn("w:p") and not text:
            current.getparent().remove(current)
            current = next_element
            continue
        return


def main() -> None:
    if sha256(SOURCE) != EXPECTED_SHA256:
        raise RuntimeError("Le document source ne correspond plus à la version analysée.")
    required_images = [
        "drawio_global.png",
        "drawio_saas.png",
        "drawio_catalogue.png",
        "drawio_dealer.png",
        "drawio_financement.png",
        "drawio_securite.png",
    ]
    missing = [name for name in required_images if not (DIAGRAMS / name).exists()]
    if missing:
        raise FileNotFoundError(f"Diagrammes manquants : {', '.join(missing)}")

    document = Document(str(SOURCE))
    candidates = [p for p in document.paragraphs if p.text.strip() == "5. Conclusion"]
    if len(candidates) != 1:
        raise RuntimeError(f"Point d’insertion ambigu : {len(candidates)} occurrence(s) de '5. Conclusion'.")
    conclusion = candidates[0]

    add_heading_before(conclusion, "5. Diagramme de classe", level=2, page_break=True)
    add_body_before(conclusion, [
        ("Le diagramme de classes formalise la structure statique de la plateforme Matchia. Il a été consolidé à partir des entités persistantes du backend et de leurs usages dans les parcours fonctionnels du frontend. Les associations pleines représentent les relations matérialisées dans le modèle de données ; les traits pointillés signalent des références logiques sans clé étrangère directe. Les cardinalités indiquent, pour chaque extrémité, le nombre minimal et maximal d’instances pouvant participer à la relation.", False),
    ])

    add_heading_before(conclusion, "5.1. Vue globale du modèle de classes", level=3)
    add_body_before(conclusion, [
        ("La vue globale relie les cinq domaines fonctionnels sans les isoler artificiellement. Les classes ", False),
        ("Bank", True),
        (", ", False),
        ("Marketplace", True),
        (", ", False),
        ("Store", True),
        (", ", False),
        ("DealerProduct", True),
        (", ", False),
        ("User", True),
        (" et ", False),
        ("FinancingRequest", True),
        (" assurent les principales jonctions. Ce positionnement permet de suivre le parcours depuis la configuration d’un marketplace et de ses offres jusqu’à la publication d’un produit, la demande de financement, la constitution du dossier et sa traçabilité.", False),
    ])
    add_picture_before(conclusion, DIAGRAMS / "drawio_global.png", "Vue globale du modèle de classes Matchia exportée depuis Draw.io", 6.30)
    add_caption_before(conclusion, "Vue globale du modèle de classes Matchia")
    add_entity_descriptions_before(conclusion, [
        ("Bank", "représente l’établissement bancaire et porte le contexte fonctionnel du tenant."),
        ("Marketplace", "centralise la configuration visuelle et commerciale du portail associé à une banque."),
        ("Request", "enregistre une demande SaaS et les stores et modules sélectionnés avant son activation."),
        ("User", "porte l’identité, le rôle et le rattachement éventuel à une banque ou à un concessionnaire."),
        ("DealerProduct", "constitue la source fonctionnelle des produits proposés par les concessionnaires."),
        ("FinancingRequest", "regroupe le dossier de financement, son client, son produit, sa banque et son état de traitement."),
    ])

    add_heading_before(conclusion, "5.2. Noyau SaaS, demandes, abonnements et paiements", level=3, page_break=True)
    add_body_before(conclusion, [
        ("Le noyau SaaS organise le cycle de souscription. ", False),
        ("Request", True),
        (" centralise une demande adressée à une banque et à un marketplace. ", False),
        ("RequestStoreSelection", True),
        (" et ", False),
        ("RequestModuleSelection", True),
        (" conservent les choix réalisés au moment de la demande, y compris les informations tarifaires utiles. Après validation, ", False),
        ("Subscription", True),
        (" matérialise l’abonnement obtenu, tandis que ", False),
        ("Payment", True),
        (" enregistre les transactions associées. Cette séparation évite de confondre la phase de configuration commerciale avec l’exécution de l’abonnement.", False),
    ])
    add_picture_before(conclusion, DIAGRAMS / "drawio_saas.png", "Noyau SaaS, demandes, abonnements et paiements exporté depuis Draw.io")
    add_caption_before(conclusion, "Noyau SaaS, demandes, abonnements et paiements")
    add_entity_descriptions_before(conclusion, [
        ("Request", "table centrale de la demande de création ou de renouvellement d’un service SaaS."),
        ("RequestStoreSelection", "conserve le snapshot des stores choisis dans la demande."),
        ("RequestModuleSelection", "conserve les modules sélectionnés et leurs paramètres au moment de la demande."),
        ("Subscription", "matérialise l’abonnement activé à la suite d’une demande validée."),
        ("Payment", "enregistre les paiements et les renouvellements liés à l’abonnement."),
    ])

    add_heading_before(conclusion, "5.3. Catalogue, stores, modules et personnalisation marketplace", level=3, page_break=True)
    add_body_before(conclusion, [
        ("Ce sous-ensemble décrit la configuration fonctionnelle et visuelle d’un marketplace. ", False),
        ("MarketplaceStore", True),
        (" matérialise l’affectation d’un ", False),
        ("Store", True),
        (" à un ", False),
        ("Marketplace", True),
        (", alors que ", False),
        ("ModuleStore", True),
        (" définit les modules activés et leur tarification contextualisée. Les valeurs spécifiques sont portées par ", False),
        ("ModuleStoreParameter", True),
        (". Les classes ", False),
        ("Content", True),
        (" et ", False),
        ("ContentVisibility", True),
        (" pilotent les contenus et leur visibilité par store. Enfin, ", False),
        ("ProductParameterDefinition", True),
        (" décrit les paramètres de catalogue réellement proposés dans cette configuration.", False),
    ])
    add_picture_before(conclusion, DIAGRAMS / "drawio_catalogue.png", "Catalogue, stores, modules et personnalisation marketplace exporté depuis Draw.io")
    add_caption_before(conclusion, "Catalogue, stores, modules et personnalisation marketplace")
    add_entity_descriptions_before(conclusion, [
        ("Marketplace", "décrit le portail bancaire, son identité visuelle et ses paramètres principaux."),
        ("MarketplaceStore", "associe un store au marketplace et porte sa visibilité ainsi que son ordre d’affichage."),
        ("ModuleStore", "définit les modules disponibles et leur configuration pour un store."),
        ("ContentVisibility", "contrôle la visibilité d’un contenu dans un marketplace donné."),
        ("ProductParameterDefinition", "définit le schéma des paramètres applicables aux produits d’un store."),
    ])

    add_heading_before(conclusion, "5.4. Concessionnaires, partenariats, contrats et publication produit", level=3, page_break=True)
    add_body_before(conclusion, [
        ("Le domaine concessionnaire repose sur ", False),
        ("Dealer", True),
        (", profil métier distinct du compte d’accès ", False),
        ("User", True),
        (". ", False),
        ("DealerBankPartnership", True),
        (" est l’association porteuse du partenariat entre un concessionnaire, une banque et un store. ", False),
        ("PartnershipContract", True),
        (" dépend uniquement de cette association : les liens directs redondants vers les trois parties sont ainsi supprimés. ", False),
        ("DealerProduct", True),
        (" constitue la source fonctionnelle unique des produits publiés ; ses documents, ses images et les demandes de publication sont modélisés par des classes dépendantes.", False),
    ])
    add_picture_before(conclusion, DIAGRAMS / "drawio_dealer.png", "Concessionnaires, partenariats, contrats et publication produit exporté depuis Draw.io")
    add_caption_before(conclusion, "Concessionnaires, partenariats, contrats et publication produit")
    add_entity_descriptions_before(conclusion, [
        ("Dealer", "porte le profil métier, les coordonnées et le statut du concessionnaire."),
        ("DealerBankPartnership", "matérialise l’association entre le concessionnaire, la banque et le store."),
        ("PartnershipContract", "conserve les versions contractuelles rattachées au partenariat, sans dupliquer les trois parties."),
        ("DealerProduct", "décrit le produit publié par le concessionnaire et ses informations commerciales."),
        ("ProductPublicationRequest", "trace la demande de validation et de publication d’un produit."),
    ])

    add_heading_before(conclusion, "5.5. Client, simulation et traitement d’une demande de financement", level=3, page_break=True)
    add_body_before(conclusion, [
        ("Le traitement du financement est centré sur ", False),
        ("FinancingRequest", True),
        (". Chaque demande est créée par un ", False),
        ("User", True),
        (" et porte sur un ", False),
        ("DealerProduct", True),
        (", avec la banque et le store concernés. Les données issues de la simulation alimentent les montants et la durée de la demande. Les documents déposés sont agrégés par ", False),
        ("FinancingRequestDocument", True),
        (" et les règles documentaires applicables sont décrites par ", False),
        ("RequiredFinancingDocument", True),
        (". Le statut de la demande et ses dates de traitement assurent le suivi du dossier.", False),
    ])
    add_picture_before(conclusion, DIAGRAMS / "drawio_financement.png", "Client, simulation et traitement d’une demande de financement exporté depuis Draw.io")
    add_caption_before(conclusion, "Client, simulation et traitement d’une demande de financement")
    add_entity_descriptions_before(conclusion, [
        ("FinancingRequest", "table centrale du dossier ; elle porte les montants, la durée, le statut et les références métier."),
        ("DealerProduct", "identifie l’unique produit concessionnaire faisant l’objet du financement."),
        ("RequiredFinancingDocument", "décrit les pièces exigées par la banque et le store pour un type de financement."),
        ("FinancingRequestDocument", "enregistre les pièces réellement déposées par le client dans son dossier."),
        ("Notification", "informe le destinataire des changements importants associés au traitement de la demande."),
    ])

    add_heading_before(conclusion, "5.6. Authentification, vérifications, notifications et traçabilité", level=3, page_break=True)
    add_body_before(conclusion, [
        ("La classe ", False),
        ("User", True),
        (" constitue le point d’entrée de la sécurité et de l’identité. ", False),
        ("RefreshToken", True),
        (" et ", False),
        ("PasswordResetToken", True),
        (" prennent en charge le renouvellement de session et la réinitialisation du mot de passe. ", False),
        ("JoinEmailVerification", True),
        (" et ", False),
        ("ClientRegistrationVerification", True),
        (" isolent les parcours de vérification. Les classes ", False),
        ("Notification", True),
        (" et ", False),
        ("AuditLog", True),
        (" complètent le dispositif par l’information de l’utilisateur et la conservation des traces. Certaines références métier y restent logiques afin de ne pas créer de dépendances persistantes inutiles.", False),
    ])
    add_picture_before(conclusion, DIAGRAMS / "drawio_securite.png", "Authentification, vérifications, notifications et traçabilité exporté depuis Draw.io")
    add_caption_before(conclusion, "Authentification, vérifications, notifications et traçabilité")
    add_entity_descriptions_before(conclusion, [
        ("User", "porte le compte applicatif, le rôle, le statut et les rattachements métier de l’utilisateur."),
        ("RefreshToken", "gère la continuité des sessions authentifiées et la révocation des jetons."),
        ("PasswordResetToken", "sécurise le workflow temporaire de réinitialisation du mot de passe."),
        ("JoinEmailVerification", "vérifie une adresse électronique avant l’association définitive au compte."),
        ("ClientRegistrationVerification", "conserve les données temporaires de vérification avant la création du client."),
        ("AuditLog", "trace les actions, les ressources concernées et les changements importants du système."),
    ])

    replace_paragraph_text(conclusion, "6. Conclusion", 14, True)
    conclusion.paragraph_format.page_break_before = False
    conclusion.paragraph_format.space_after = Pt(6)
    set_keep_with_next(conclusion)
    add_conclusion_after(conclusion)
    remove_empty_paragraphs_before_next_section(conclusion)

    ensure_update_fields_on_open(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    print(f"Created: {OUTPUT}")
    print(f"Size: {OUTPUT.stat().st_size}")


if __name__ == "__main__":
    main()
