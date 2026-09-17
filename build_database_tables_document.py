from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = r"D:\PFE M2\Platforme SaaS\Tables_et_roles_de_la_base_Matchia.docx"

GROUPS = [
    ("Identité et sécurité", [
        ("users", "Comptes SaaS, banques, clients et concessionnaires."),
        ("bank", "Informations et configuration des banques."),
        ("marketplace", "Configuration et identité visuelle de chaque marketplace."),
        ("refresh_tokens", "Sessions sécurisées des utilisateurs."),
        ("password_reset_tokens", "Jetons temporaires de réinitialisation des mots de passe."),
        ("join_email_verifications", "Vérification e-mail lors de l'inscription d'une banque ou marketplace."),
        ("client_registration_verifications", "Codes e-mail temporaires lors de la création d'un compte client."),
    ]),
    ("Catalogue et configuration des marketplaces", [
        ("store", "Catalogue global des stores : mobile, véhicule, médical, etc."),
        ("module", "Catalogue global des modules."),
        ("module_store", "Modules compatibles avec un store, avec prix et statut."),
        ("module_store_parameter", "Paramètres configurables d'un module pour un store."),
        ("marketplace_store", "Stores affectés à une marketplace."),
        ("marketplace_store_module", "Modules réellement activés et visibles pour un store d'une marketplace."),
        ("marketplace_store_banner", "Images multiples des bannières par store et marketplace."),
    ]),
    ("Demandes abonnements et paiements", [
        ("request", "Demandes d'inscription, de store, module, abonnement ou renouvellement."),
        ("request_store", "Stores sélectionnés dans une demande."),
        ("request_module", "Modules sélectionnés dans une demande."),
        ("request_store_selection", "Détails et prix des stores sélectionnés dans une demande."),
        ("request_module_selection", "Détails et prix des modules sélectionnés par store dans une demande."),
        ("subscription", "Cycle de vie des abonnements annuels."),
        ("payment", "Paiements initiaux et de renouvellement."),
    ]),
    ("Contenus notifications et suivi", [
        ("content", "Contenu global lié à un store, créé côté SaaS."),
        ("content_visibility", "Visibilité d'un contenu global dans chaque marketplace."),
        ("marketplace_content", "Contenu propre à une marketplace."),
        ("notifications", "Notifications SaaS, banque, client et concessionnaire."),
        ("audit_logs", "Historique des actions réalisées dans l'application."),
    ]),
    ("Concessionnaires et produits", [
        ("dealer", "Profil et informations du concessionnaire."),
        ("dealer_account_request", "Demandes de création de comptes concessionnaires."),
        ("dealer_request_document", "Pièces jointes d'une demande de compte concessionnaire."),
        ("dealer_bank_partnership", "Partenariats entre banques et concessionnaires."),
        ("partnership_contract", "Contrats PDF des partenariats banque-concessionnaire."),
        ("dealer_product", "Produits ajoutés par les concessionnaires."),
        ("dealer_product_catalog_image", "Galerie d'images des produits concessionnaires."),
        ("dealer_product_document", "Documents associés aux produits concessionnaires."),
        ("dealer_product_parameter_value", "Valeurs des caractéristiques des produits concessionnaires."),
        ("product_publication_request", "Demandes de publication de produits concessionnaires."),
    ]),
    ("Financement et ancien catalogue banque", [
        ("financing_request", "Demandes de financement envoyées par les clients."),
        ("financing_request_document", "Documents déposés pour une demande de financement."),
        ("required_financing_document", "Documents exigés selon la banque et le store pour un financement."),
        ("product", "Ancien catalogue de produits créés par une banque."),
        ("product_parameter_definition", "Définitions des caractéristiques d'un produit par store ; également utilisées par les produits concessionnaires."),
        ("product_parameter_value", "Valeurs des caractéristiques des anciens produits banque."),
    ]),
]

TECHNICAL = [("flyway_schema_history", "Historique technique des migrations SQL Flyway.")]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9E2F0", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=10.5, bold=False, color="1F2937", name="Aptos"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, space_after=0, alignment=None):
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = 1.08
    if alignment is not None:
        paragraph.alignment = alignment


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.style = "Table Grid"
    table.columns[0].width = Cm(5.0)
    table.columns[1].width = Cm(11.2)

    header = table.rows[0]
    set_repeat_table_header(header)
    headers = ("Table", "Rôle")
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        cell.width = Cm(5.0 if idx == 0 else 11.2)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "123A71")
        set_cell_border(cell, "D9E2F0")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        style_paragraph(p, alignment=WD_ALIGN_PARAGRAPH.LEFT)
        run = p.add_run(value)
        set_font(run, size=10.5, bold=True, color="FFFFFF")

    for index, (name, role) in enumerate(rows):
        row = table.add_row()
        fill = "F5F8FC" if index % 2 == 0 else "FFFFFF"
        values = (name, role)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.width = Cm(5.0 if idx == 0 else 11.2)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, fill)
            set_cell_border(cell, "D9E2F0")
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            style_paragraph(p, alignment=WD_ALIGN_PARAGRAPH.LEFT)
            run = p.add_run(value)
            set_font(run, size=9.6, bold=(idx == 0), color="123A71" if idx == 0 else "26354A", name="Aptos")

    doc.add_paragraph().paragraph_format.space_after = Pt(3)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(1.7)
section.right_margin = Cm(1.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title_run = title.add_run("Tables et rôles de la base Matchia")
set_font(title_run, size=22, bold=True, color="000000", name="Aptos Display")
style_paragraph(title, space_after=8)

intro = doc.add_paragraph()
style_paragraph(intro, space_after=12)
intro_run = intro.add_run(
    "Ce document présente les 42 tables métier actuellement mappées par l'application Matchia, "
    "ainsi que la table technique de migration Flyway. Il facilite l'identification des données à conserver, "
    "à nettoyer ou à faire évoluer."
)
set_font(intro_run, size=10.8, color="334155")

summary = doc.add_paragraph()
style_paragraph(summary, space_after=15)
summary_run = summary.add_run("Total attendu : 42 tables métier et 1 table technique.")
set_font(summary_run, size=10.8, bold=True, color="123A71")

for heading, rows in GROUPS:
    paragraph = doc.add_paragraph(style="Heading 1")
    style_paragraph(paragraph, space_after=6)
    run = paragraph.add_run(heading)
    set_font(run, size=14, bold=True, color="000000")
    add_table(doc, rows)

paragraph = doc.add_paragraph(style="Heading 1")
style_paragraph(paragraph, space_after=6)
run = paragraph.add_run("Table technique")
set_font(run, size=14, bold=True, color="000000")
add_table(doc, TECHNICAL)

paragraph = doc.add_paragraph(style="Heading 1")
style_paragraph(paragraph, space_after=6)
run = paragraph.add_run("Points de vigilance")
set_font(run, size=14, bold=True, color="000000")

for text in [
    "Certificate n'est pas une table. Il s'agit d'un type de document de financement, enregistré avec les documents de demande de financement.",
    "La table notifications reste nécessaire aux alertes SaaS, banques, clients et concessionnaires, même si un écran de notifications est retiré du back office SaaS.",
    "Si seuls les concessionnaires doivent créer des produits, product et product_parameter_value peuvent devenir des tables historiques. Elles ne doivent être supprimées qu'après retrait de leurs services backend, de leurs API et de leurs références dans le chatbot et les demandes de financement.",
]:
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, space_after=4)
    run = p.add_run(text)
    set_font(run, size=10.3, color="334155")

doc.core_properties.title = "Tables et rôles de la base Matchia"
doc.core_properties.subject = "Référentiel des tables de la base de données Matchia"
doc.core_properties.author = "Matchia"
doc.save(OUTPUT)
print(OUTPUT)
